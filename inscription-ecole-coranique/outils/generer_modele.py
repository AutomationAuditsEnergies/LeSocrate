"""Génère modele-inscription.docx : réplique de la استمارة التسجيل avec des
balises {{...}} que le Google Apps Script remplace à chaque inscription.

Usage : python3 outils/generer_modele.py  (depuis le dossier inscription-ecole-coranique)
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
FONT = "Amiri"


def rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.insert(0, OxmlElement("w:bidi"))


def add_run(paragraph, text, size=13, bold=False, color=None, underline=False):
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    fonts.set(qn("w:cs"), FONT)
    rpr.append(OxmlElement("w:rtl"))
    sz = OxmlElement("w:szCs")
    sz.set(qn("w:val"), str(size * 2))
    rpr.append(sz)
    if bold:
        rpr.append(OxmlElement("w:bCs"))
    return run


# align=None : pas d'alignement explicite → le paragraphe RTL s'aligne à droite
# partout (Word, LibreOffice, Google Docs). Un "right" explicite est ambigu en RTL.
def para(container, text="", size=13, bold=False, align=None,
         space_after=4, color=None, underline=False):
    p = container.add_paragraph()
    rtl(p)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        add_run(p, text, size, bold, color, underline)
    return p


def shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def cell_borders(cell, style="none", color="000000", size=8):
    tcpr = cell._tc.get_or_add_tcPr()
    for old in tcpr.findall(qn("w:tcBorders")):
        tcpr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), style)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tcpr.append(borders)


def fixed_width(table, widths):
    """Largeurs fixes (sinon LibreOffice/Google Docs recalculent les colonnes)."""
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(int(w.twips) for w in widths)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for col, w in zip(table.columns, widths):
        col.width = w
        for cell in col.cells:
            cell.width = w


def section_bar(doc, text, fill):
    p = para(doc, space_after=6)
    shade(p, fill)
    p.paragraph_format.space_before = Pt(10)
    add_run(p, text, 14, bold=True, underline=True)


def field(doc, label, key, size=13):
    p = para(doc, space_after=6)
    add_run(p, f"{label}: ", size, bold=True)
    add_run(p, "{{" + key + "}}", size)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.4)

# --- En-tête : photo | nom de l'école | logo -------------------------------
head = doc.add_table(rows=1, cols=3)
head.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = (Cm(3.6), Cm(11), Cm(3.2))
fixed_width(head, widths)
for cell in head.rows[0].cells:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_borders(cell)
photo, title, logo = head.rows[0].cells
cell_borders(photo, "dashed", "000000", 6)
photo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
rtl(photo.paragraphs[0])
add_run(photo.paragraphs[0], "{{PHOTO}}", 11, bold=True)
head.rows[0].height = Cm(4.4)

title.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
rtl(title.paragraphs[0])
add_run(title.paragraphs[0], "المدرسة القرآنية محمد الشريف ڨاهر", 20, bold=True)
para(title, "التابعة لمسجد سلمان الفارسي رضي الله عنه", 14, bold=True,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para(title, "باب الزوار", 12, align=WD_ALIGN_PARAGRAPH.CENTER)

logo.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
logo.paragraphs[0].add_run().add_picture(str(ROOT / "logo.png"), width=Cm(3))

# --- Titre encadré ---------------------------------------------------------
para(doc, space_after=4)  # sépare les 2 tableaux (sinon ils fusionnent)
box = doc.add_table(rows=1, cols=1)
box.alignment = WD_TABLE_ALIGNMENT.CENTER
fixed_width(box, (Cm(11.5),))
bc = box.rows[0].cells[0]
cell_borders(bc, "single", "000000", 18)
tcpr = bc._tc.get_or_add_tcPr()
shd = OxmlElement("w:shd")
shd.set(qn("w:val"), "clear")
shd.set(qn("w:fill"), "D9D9D9")
tcpr.append(shd)
bc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
rtl(bc.paragraphs[0])
add_run(bc.paragraphs[0], "استمارة التسجيل الالكتروني", 20, bold=True)
para(doc, space_after=2)

para(doc, space_after=6)
p = field(doc, "تاريخ التسجيل", "DATE_INSCRIPTION")
add_run(p, "        رقم التسجيل: ", 13, bold=True)
add_run(p, "{{NUMERO}}", 13)

# --- Élève -----------------------------------------------------------------
section_bar(doc, "معلومات خاصة بالطالب", "BDD7EE")
field(doc, "إسم ولقب الطالب", "NOM_ELEVE")
field(doc, "تاريخ الميلاد", "DATE_NAISSANCE")
field(doc, "عنوان الإقامة", "ADRESSE")
field(doc, "المستوى التعليمي", "NIVEAU")
field(doc, "رقم هاتف الطالب", "TEL_ELEVE")
p = para(doc, space_after=6)
add_run(p, "هل يعاني الطالب من مشاكل صحيّة؟   ", 13, bold=True)
add_run(p, "نعم {{CASE_OUI}}      لا {{CASE_NON}}", 13)
field(doc, "حددها", "SANTE_DETAILS")

# --- Tuteur ----------------------------------------------------------------
section_bar(doc, "معلومات خاصة بولي أمر الطالب", "A9D08E")
field(doc, "إسم ولي أمر الطالب", "NOM_TUTEUR")
field(doc, "رقم بطاقة تعريف ولي أمر الطالب", "CNI_TUTEUR")
p = field(doc, "رقم الهاتف", "TEL_TUTEUR")
add_run(p, "        البريد الإلكتروني: ", 13, bold=True)
add_run(p, "{{EMAIL}}", 13)

# --- Engagement ------------------------------------------------------------
section_bar(doc, "تعهّد و إلتزام", "808080")
para(doc, "أنا ولي أمر الطالب (المذكور أعلاه)، أوافق على النظام الداخلي للمدرسة القرآنية "
          "وأتعهد بالالتزام به ومتعاوناً معكم في كل ما يخص إبني، دراسةً وسلوكاً وانتظاماً.",
     13, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
field(doc, "الموافقة على النظام الداخلي", "ACCORD")
para(doc, space_after=10)
para(doc, "تم إنشاء هذه الاستمارة تلقائياً عند التسجيل الإلكتروني بتاريخ {{HORODATAGE}}",
     9, align=WD_ALIGN_PARAGRAPH.CENTER, color="5A5A5A")

# Remet les propriétés dans l'ordre exigé par le schéma OOXML (Word / Google
# Docs ignorent sinon une partie de la mise en forme : largeurs, bordures...).
ORDER = {
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl",
            "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens",
            "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN",
            "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
            "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"],
    "rPr": ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
            "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish",
            "webHidden", "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight",
            "u", "effect", "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
            "eastAsianLayout", "specVanish", "oMath"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap",
             "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
              "shd", "tblLayout", "tblCellMar", "tblLook"],
}
for tag, order in ORDER.items():
    rank = {name: i for i, name in enumerate(order)}
    for el in doc.element.body.iter(qn(f"w:{tag}")):
        children = sorted(el, key=lambda c: rank.get(c.tag.split("}")[1], len(order)))
        for c in children:
            el.append(c)

out = ROOT / "modele-inscription.docx"
doc.save(out)
print("OK", out)
