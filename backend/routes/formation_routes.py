"""
Routes pipeline formation automatisé.

POST /api/formation/search-rncp           Recherche RNCP depuis nom TP
POST /api/formation/init                  Crée un job (tp_name, total_hours, rncp_code)
GET  /api/formation/<job_id>              Statut + données du job
POST /api/formation/<job_id>/fetch-reac   Télécharge + extrait le REAC
POST /api/formation/<job_id>/generate-global   Lance génération programme global
POST /api/formation/<job_id>/validate-global   Valide (et éventuellement édite) le programme global
POST /api/formation/<job_id>/split-daily       Lance le découpage en journées
POST /api/formation/<job_id>/validate-daily    Valide les programmes journée
POST /api/formation/<job_id>/launch-tts        Crée les dossiers et lance la génération TEXTE des cours
GET  /api/formation/<job_id>/content      Liste les dossiers cours (journées) + état génération texte
GET  /api/formation/<job_id>/content/<folder_id>/docx  Télécharge le document Word d'une journée
POST /api/formation/<job_id>/launch-audio Lance la synthèse TTS Fish Audio sur toutes les journées
GET  /api/formation/list                  Liste les jobs de la plateforme
"""

import os
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

from flask import Blueprint, jsonify, request, session, send_file
from io import BytesIO

from services.formation_pipeline_service import (
    search_rncp,
    download_reac_text_with_retry,
    download_rc_text,
    fetch_rome_data,
    launch_global_program_generation,
    launch_daily_split,
    run_daily_split,
    launch_tts_for_all_days,
    create_job,
    update_job,
    get_job,
    list_jobs,
    HOURS_PER_DAY,
    COURSE_AUDIO_SLOTS,
    _normalize_day_audio_slots,
    _format_slot_generation_source,
)
from services.knowledge_base_service import (
    launch_kb_building,
    list_kb,
    kb_stats,
)
from utils.logger import get_logger

logger = get_logger(__name__)

formation_bp = Blueprint("formation", __name__)

_PIPELINE_MODEL_ALIASES = {
    "sonnet": "claude-sonnet-4-20250514",
    "haiku": "claude-haiku-4-5-20251001",
    "flash": "deepseek-v4-flash",
    "pro": "deepseek-v4-pro",
}
_PIPELINE_MODEL_CHOICES = set(_PIPELINE_MODEL_ALIASES)


def _slides_folder_workers(default: int = 3) -> int:
    """Nombre de journées dont les decks slides sont générés en parallèle."""
    try:
        workers = int(os.getenv("FORMATION_SLIDES_FOLDER_WORKERS", str(default)))
    except (TypeError, ValueError):
        workers = default
    return max(1, min(8, workers))


def _resolve_pipeline_slide_model(api_model: str | None) -> str | None:
    """Modèle dédié à la curation slides.

    L'itération manuelle "Régénérer curation + slides" utilise DeepSeek Pro par
    défaut. On aligne l'auto-pilot dessus, tout en laissant un override env pour
    les environnements sans clé DeepSeek.
    """
    override = (os.getenv("FORMATION_SLIDES_MODEL") or "").strip()
    if override:
        return _PIPELINE_MODEL_ALIASES.get(override.lower(), override)
    return "deepseek-v4-pro"


def _formation_content_day_workers(default: int = 1) -> int:
    """Nombre de journees de contenu generees en parallele par l'auto-pilot API."""
    try:
        workers = int(os.getenv("FORMATION_CONTENT_DAY_WORKERS", str(default)))
    except (TypeError, ValueError):
        workers = default
    return max(1, min(52, workers))


def _normalize_pipeline_model_choice(raw, default=None):
    """Normalise le choix UI persistant de l'auto-pilot."""
    value = (raw or default or "").strip().lower()
    if value in _PIPELINE_MODEL_CHOICES:
        return value
    return default


def _resolve_pipeline_api_model(job: dict | None, requested_model=None):
    """Résout le modèle API à utiliser pour une action du job.

    Priorité :
      1. modèle explicite passé en argument
      2. modèle choisi au lancement (`auto_pilot_model`)
      3. fallback FORMATION_LLM_PROVIDER (env var) → deepseek-v4-pro / sonnet
      4. fallback DEEPSEEK_API_KEY (sans ANTHROPIC_API_KEY) → deepseek-v4-pro
      5. None (laisse les services retomber sur leur default_model())

    Garantit qu'un job lancé en DeepSeek reste en DeepSeek pour TOUTES les
    étapes, même si `auto_pilot_model` n'a pas été persisté côté DB (jobs
    historiques) — ce qui évite que les services de transition retombent
    silencieusement sur Anthropic.
    """
    model = requested_model or (job or {}).get("auto_pilot_model")
    if not model:
        provider = (
            os.environ.get("FORMATION_LLM_PROVIDER")
            or os.environ.get("LLM_PROVIDER")
            or ""
        ).strip().lower()
        if provider == "deepseek":
            model = "deepseek-v4-pro"
        elif provider == "anthropic":
            model = "sonnet"
        elif os.environ.get("DEEPSEEK_API_KEY") and not os.environ.get("ANTHROPIC_API_KEY"):
            model = "deepseek-v4-pro"
    if not model:
        return None
    model = str(model).strip()
    return _PIPELINE_MODEL_ALIASES.get(model.lower(), model)


def _get_platform_id():
    """Retourne platform_id depuis session ou header."""
    pid = request.headers.get("X-Platform-Id") or session.get("platform_id", 1)
    try:
        return int(pid)
    except (ValueError, TypeError):
        return 1


def _require_admin():
    """Retourne True si l'utilisateur est authentifié admin."""
    return session.get("is_admin", False)


# ─── Recherche RNCP ───────────────────────────────────────────────────────────

@formation_bp.route("/api/formation/search-rncp", methods=["POST"])
def search_rncp_route():
    """
    Recherche des titres RNCP correspondant à un nom de TP.
    Body: { "query": "TP CRCD" }
    Retourne: [{ rncp_code, title }]
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    data = request.get_json() or {}
    query = (data.get("query") or "").strip()
    if not query:
        return jsonify({"error": "Le champ 'query' est requis"}), 400

    try:
        results = search_rncp(query)
        return jsonify({"results": results})
    except Exception as e:
        logger.error(f"❌ search_rncp : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Helpers parsing DOCX/TXT pour le mode test ──────────────────────────────

# Sous-parties standard par défaut pour les jobs de test où on n'a pas de
# vrai daily split. Aligné sur la playlist cours : 7 créneaux × 3 passes = 21 segments.
_TEST_SUB_PARTS = [
    f"{slot['label']} — {slot['start']}-{slot['end']} — Créneau test {slot['index'] + 1}"
    for slot in COURSE_AUDIO_SLOTS
]


def _read_doc_text(file_storage) -> str:
    """Lit un FileStorage uploadé (.docx ou .txt) et retourne le texte brut."""
    filename = (file_storage.filename or "").lower()
    if filename.endswith(".txt"):
        raw = file_storage.read()
        try:
            return raw.decode("utf-8")
        except UnicodeDecodeError:
            return raw.decode("latin-1", errors="ignore")
    if filename.endswith(".docx"):
        from docx import Document
        doc = Document(file_storage)
        return "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
    raise ValueError(f"Format de fichier non supporté : {filename} (attendu .docx ou .txt)")


def _split_into_21_chunks(text: str) -> list:
    """Découpe un texte en 21 chunks à peu près équilibrés en paragraphes.

    21 = 7 créneaux cours × 3 passes (playlist audio réelle).
    Si le texte a moins de 21 paragraphes, on le pad par duplication. Le but
    n'est pas de produire du contenu pédagogique fin (c'est un mode test) mais
    d'avoir 21 segments avec du texte non vide pour que la review et l'audio
    aient quelque chose à mâcher.
    """
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if len(paragraphs) < 21:
        # Pad par duplication cyclique
        paragraphs = (paragraphs * (21 // max(len(paragraphs), 1) + 1))[: max(21, len(paragraphs))]

    chunks = []
    paras_per_chunk = max(1, len(paragraphs) // 21)
    for i in range(21):
        start = i * paras_per_chunk
        end = start + paras_per_chunk if i < 20 else len(paragraphs)
        chunks.append("\n\n".join(paragraphs[start:end]).strip() or paragraphs[i % len(paragraphs)])
    return chunks


# ─── Initialisation d'un job ──────────────────────────────────────────────────

@formation_bp.route("/api/formation/init", methods=["POST"])
def init_formation():
    """
    Crée un job pipeline formation + une nouvelle plateforme dédiée.
    Body: { "platform_name": "TP CRCD 2026", "tp_name": "TP CRCD", "total_hours": 70, "rncp_code": "35304" }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    data = request.get_json() or {}
    platform_name = (data.get("platform_name") or "").strip()
    tp_name = (data.get("tp_name") or "").strip()
    rncp_code = (data.get("rncp_code") or "").strip()
    total_hours = data.get("total_hours")
    model_choice = _normalize_pipeline_model_choice(data.get("model"))

    if not platform_name:
        return jsonify({"error": "Le champ 'platform_name' est requis"}), 400
    if not tp_name:
        return jsonify({"error": "Le champ 'tp_name' est requis"}), 400
    if not rncp_code:
        return jsonify({"error": "Le champ 'rncp_code' est requis"}), 400
    if not total_hours or int(total_hours) <= 0:
        return jsonify({"error": "Le champ 'total_hours' doit être > 0"}), 400

    total_hours = int(total_hours)
    if total_hours % HOURS_PER_DAY != 0:
        return jsonify({
            "error": f"Le champ 'total_hours' doit être un multiple de {HOURS_PER_DAY} (1 journée = {HOURS_PER_DAY}h)",
        }), 400
    nb_days = total_hours // HOURS_PER_DAY

    try:
        from database.db import get_db_connection
        from datetime import datetime
        from config import FRANCE_TZ

        conn = get_db_connection()
        cursor = conn.cursor()

        # Créer une nouvelle plateforme dédiée à cette formation
        now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute(
            "INSERT INTO platform_config (name, upload_locked, updated_at) VALUES (?, 1, ?)",
            (platform_name, now_str)
        )
        new_platform_id = cursor.lastrowid

        # Containers Azure auto-générés (à créer manuellement dans Azure si besoin)
        slug = platform_name.lower().replace(" ", "-").replace("_", "-")[:40]
        cursor.execute(
            """UPDATE platform_config
               SET audio_container=?, pdf_container=?, archive_container=?, slug=?, audio_base_url=?
               WHERE id=?""",
            (
                f"formationaudio-p{new_platform_id}",
                f"formationpdf-p{new_platform_id}",
                f"formationaudio-p{new_platform_id}-archives",
                slug,
                "",  # à configurer dans Azure ensuite
                new_platform_id,
            )
        )
        conn.commit()
        conn.close()

        logger.info(f"✅ Nouvelle plateforme créée : id={new_platform_id} '{platform_name}'")

        job_id = create_job(
            platform_id=new_platform_id,
            tp_name=tp_name,
            rncp_code=rncp_code,
            total_hours=total_hours,
            nb_days=nb_days,
        )
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE platform_config SET source_formation_id = ? WHERE id = ? AND source_formation_id IS NULL",
            (job_id, new_platform_id),
        )
        conn.commit()
        conn.close()
        if model_choice:
            update_job(job_id, auto_pilot_model=model_choice)
        logger.info(f"✅ Job formation créé : {job_id} ({tp_name}, {total_hours}h, {nb_days} jours, plateforme {new_platform_id})")
        return jsonify({
            "job_id": job_id,
            "platform_id": new_platform_id,
            "platform_name": platform_name,
            "tp_name": tp_name,
            "rncp_code": rncp_code,
            "total_hours": total_hours,
            "nb_days": nb_days,
            "model": model_choice,
            "status": "init",
        }), 201

    except Exception as e:
        logger.error(f"❌ init_formation : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Mode test : init avec DOCX pré-injectés (skip génération content) ───────

@formation_bp.route("/api/formation/init-test", methods=["POST"])
def init_test_pipeline():
    """Crée une plateforme + job + folders + segments en mode TEST.

    L'auto-pilot relancé dessus skippera naturellement KB/global/daily/content
    (car tous les artefacts sont déjà en DB) et ne tournera que finalize +
    review + audio + health-check. Permet de valider la pipeline en aval en
    ~5-10 min au lieu de 30-60.

    Multipart form-data :
      - platform_name (str, requis)
      - tp_name (str, requis)
      - rncp_code (str, requis)
      - total_hours (int, requis) — doit être un multiple de 7 (1 doc par 7h)
      - tts_mode (str, optionnel, défaut 'mock')
      - auto_pilot (bool, optionnel, défaut true)
      - docs (files[], requis) — N fichiers .docx ou .txt, où N = total_hours/7

    Retourne 202 avec { job_id, platform_id, nb_days, segments_inserted }.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    platform_name = (request.form.get("platform_name") or "").strip()
    tp_name = (request.form.get("tp_name") or "").strip()
    rncp_code = (request.form.get("rncp_code") or "").strip()
    total_hours_raw = request.form.get("total_hours") or "0"
    tts_mode = (request.form.get("tts_mode") or "mock").lower()
    auto_pilot = (request.form.get("auto_pilot") or "true").lower() == "true"

    if not platform_name or not tp_name or not rncp_code:
        return jsonify({"error": "platform_name, tp_name, rncp_code sont requis"}), 400
    try:
        total_hours = int(total_hours_raw)
    except ValueError:
        return jsonify({"error": "total_hours doit être un entier"}), 400
    if total_hours <= 0:
        return jsonify({"error": "total_hours doit être > 0"}), 400
    if total_hours % HOURS_PER_DAY != 0:
        return jsonify({
            "error": f"total_hours doit être un multiple de {HOURS_PER_DAY} (1 journée = {HOURS_PER_DAY}h)",
        }), 400

    nb_days = total_hours // HOURS_PER_DAY
    docs = request.files.getlist("docs")
    if len(docs) != nb_days:
        return jsonify({
            "error": f"Tu dois fournir exactement {nb_days} fichier(s) (1 par journée de 7h). Reçu : {len(docs)}",
        }), 400

    try:
        from database.db import get_db_connection
        from datetime import datetime
        from config import FRANCE_TZ
        import json

        # 1. Crée la plateforme (idem init_formation)
        conn = get_db_connection()
        cursor = conn.cursor()
        now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute(
            "INSERT INTO platform_config (name, upload_locked, updated_at) VALUES (?, 1, ?)",
            (platform_name, now_str),
        )
        platform_id = cursor.lastrowid
        slug = platform_name.lower().replace(" ", "-").replace("_", "-")[:40]
        cursor.execute(
            """UPDATE platform_config
               SET audio_container=?, pdf_container=?, archive_container=?, slug=?, audio_base_url=?
               WHERE id=?""",
            (
                f"formationaudio-p{platform_id}",
                f"formationpdf-p{platform_id}",
                f"formationaudio-p{platform_id}-archives",
                slug, "", platform_id,
            ),
        )
        conn.commit()
        conn.close()
        logger.info(f"🧪 [TEST] Plateforme créée : id={platform_id} '{platform_name}'")

        # 2. Crée le job pipeline avec stubs (REAC mock, global mock, daily mock)
        job_id = create_job(
            platform_id=platform_id, tp_name=tp_name, rncp_code=rncp_code,
            total_hours=total_hours, nb_days=nb_days,
        )
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            "UPDATE platform_config SET source_formation_id = ? WHERE id = ? AND source_formation_id IS NULL",
            (job_id, platform_id),
        )
        conn.commit()
        conn.close()

        daily_programs_stub = [
            _normalize_day_audio_slots({
                "day_number": i + 1,
                "title": f"Journée {i+1} (test)",
                "sub_parts": [
                    {
                        "index": idx,
                        "name": sp_name,
                        "module_content": f"Contenu test créneau cours {idx+1}",
                    }
                    for idx, sp_name in enumerate(_TEST_SUB_PARTS)
                ],
                "day_recap": "" if i == 0 else "Lors de la dernière séance, nous avons vu les fondamentaux.",
                "day_transition": "À la prochaine séance, nous aborderons la suite du programme.",
            })
            for i in range(nb_days)
        ]
        update_job(
            job_id,
            reac_text=f"[TEST STUB] REAC mock pour {tp_name} (RNCP {rncp_code})",
            global_program=f"[TEST STUB] Programme global mock pour {tp_name}",
            daily_programs=json.dumps(daily_programs_stub, ensure_ascii=False),
            global_program_validated=1,
            daily_programs_validated=1,
            status="daily_validated",
        )
        logger.info(f"🧪 [TEST] Job pipeline {job_id} créé avec stubs (KB/global/daily seront skippés par l'auto-pilot)")

        # 3. Crée N cours_folders + cg_jobs + 21 segments par cg_job
        conn = get_db_connection()
        cursor = conn.cursor()
        segments_inserted = 0
        for day_idx, doc_file in enumerate(docs):
            day_num = day_idx + 1
            folder_name = f"Jour {day_num} — Journée {day_num} (test)"

            cursor.execute(
                "SELECT COALESCE(MAX(position), -1) + 1 FROM cours_folders WHERE platform_id = ?",
                (platform_id,),
            )
            position = cursor.fetchone()[0]
            cursor.execute(
                """
                INSERT INTO cours_folders (platform_id, name, position, formation_job_id)
                VALUES (?, ?, ?, ?)
                """,
                (platform_id, folder_name, position, job_id),
            )
            folder_id = cursor.lastrowid

            cursor.execute(
                """
                INSERT INTO content_generation_jobs
                    (folder_id, platform_id, program_text, program_title,
                     sub_parts, from_scratch, module_contents,
                     status, current_sub_part, current_passe, total_words, error_message)
                VALUES (?, ?, ?, ?, ?, 1, ?, 'idle', 0, 1, 0, NULL)
                """,
                (
                    folder_id, platform_id,
                    f"[TEST] Programme journée {day_num}", tp_name,
                    json.dumps(_TEST_SUB_PARTS, ensure_ascii=False),
                    json.dumps({sp: f"Contenu test {sp}" for sp in _TEST_SUB_PARTS}, ensure_ascii=False),
                ),
            )
            cg_job_id = cursor.lastrowid

            # Parse le doc + split en 21 chunks
            try:
                full_text = _read_doc_text(doc_file)
            except Exception as e:
                conn.close()
                return jsonify({"error": f"Lecture fichier '{doc_file.filename}' : {e}"}), 400
            chunks_21 = _split_into_21_chunks(full_text)

            # Insère 21 segments (7 créneaux cours × 3 passes)
            for sub_idx in range(len(_TEST_SUB_PARTS)):
                for passe in range(1, 4):
                    seg_idx = sub_idx * 3 + (passe - 1)
                    text = chunks_21[seg_idx]
                    word_count = len(text.split())
                    cursor.execute(
                        """
                        INSERT INTO content_generation_segments
                            (job_id, sub_part_index, sub_part_name, passe, status,
                             text_content, word_count, dirty, reviewed, review_error)
                        VALUES (?, ?, ?, ?, 'completed', ?, ?, 1, 0, NULL)
                        """,
                        (
                            cg_job_id, sub_idx, _TEST_SUB_PARTS[sub_idx], passe,
                            text, word_count,
                        ),
                    )
                    segments_inserted += 1

            logger.info(f"🧪 [TEST] Folder {folder_id} (Jour {day_num}) : 21 segments injectés depuis '{doc_file.filename}'")

        conn.commit()
        conn.close()

        # 4. Si auto_pilot demandé, lancer l'auto-pilot en mode Claude Code.
        # Important : même si KB/global/daily/content sont skippés (segments déjà
        # en DB), la conformité locale en aval CONSOMME de l'IA et donc du crédit :
        #   - Review conformité : audit règles hors micro-éthique, par morceau
        # Sans use_claude_code=True, ces appels passent par l'API LLM configurée
        # (DeepSeek ou Anthropic). Avec True, ils utilisent le forfait Pro/Max
        # via subprocess `claude`.
        # L'audio est désormais découplé : l'auto-pilot prépare le texte/Word 2,
        # puis la synthèse se lance séparément par journée/semaine.
        # Modèle SONNET (pas Haiku) : la conformité locale demande du jugement
        # linguistique fin. Haiku produit des patches mécaniques et parfois cassés.
        # Le coût est sur le forfait CC, donc gratuit côté API.
        if auto_pilot:
            import eventlet
            update_job(job_id,
                       auto_pilot_enabled=1,
                       auto_pilot_model="sonnet",
                       auto_pilot_tts_mode=tts_mode,
                       auto_pilot_use_cc=1,
                       auto_pilot_skip_vs=1,   # volume safety skippée en mode TEST
                       auto_pilot_generate_audio=0,
                       auto_pilot_volume_done=0,
                       auto_pilot_post_review_docs_done=0,
                       auto_pilot_error=None)
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "pipeline_started",
                    step="start",
                    status="running",
                    model="sonnet",
                    message="Auto-pilot test lancé",
                    data={"tts_mode": tts_mode, "use_claude_code": True, "test_mode": True},
                )
            except Exception:
                pass
            eventlet.spawn(_tick_auto_pilot, job_id)
            logger.info(
                f"🧪 [TEST] Auto-pilot (DB state machine) spawné pour job {job_id} "
                f"(tts={tts_mode}, CC Sonnet, volume_safety skippée)"
            )

        return jsonify({
            "ok": True,
            "job_id": job_id,
            "platform_id": platform_id,
            "platform_name": platform_name,
            "nb_days": nb_days,
            "segments_inserted": segments_inserted,
            "auto_pilot_started": auto_pilot,
            "tts_mode": tts_mode,
            "test_mode": True,
        }), 202

    except Exception as e:
        logger.error(f"❌ init_test_pipeline : {e}")
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ─── Statut d'un job ──────────────────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>", methods=["GET"])
def get_formation_job(job_id):
    """Retourne l'état complet d'un job."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    # Ne pas renvoyer les textes bruts (trop lourds), seulement leurs tailles
    result = {k: v for k, v in job.items() if k not in ("reac_text", "rc_text", "rome_text")}
    result["reac_available"] = bool(job.get("reac_text"))
    result["reac_length"] = len(job.get("reac_text") or "")
    result["rc_length"] = len(job.get("rc_text") or "")
    result["rome_length"] = len(job.get("rome_text") or "")

    return jsonify(result)


# ─── Téléchargement REAC ──────────────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/fetch-reac", methods=["POST"])
def fetch_reac(job_id):
    """
    Télécharge le REAC PDF depuis France Compétences et en extrait le texte.
    Lance en background (peut prendre ~10s).
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if job["status"] not in ("init", "error", "reac_ready"):
        return jsonify({"error": f"Impossible de télécharger depuis le statut '{job['status']}'"}), 400

    def _fetch_thread():
        try:
            update_job(job_id, status="reac_fetching")
            rncp_code = job["rncp_code"]

            # Télécharger REAC + RC + ROME en parallèle
            results = {"reac": "", "rc": "", "rome": ""}
            errors = []

            def _log_reac_attempt(**payload):
                try:
                    from services.formation_observability_service import log_pipeline_event
                    status = payload.get("status") or "info"
                    attempt = payload.get("attempt")
                    total = payload.get("total")
                    wait_seconds = payload.get("wait_seconds") or 0
                    error = payload.get("error")
                    message = f"REAC tentative {attempt}/{total} : {status}"
                    if status == "retrying":
                        message += f" — nouvelle tentative dans {wait_seconds:.0f}s"
                    log_pipeline_event(
                        job_id,
                        "reac_download_attempt",
                        step="reac",
                        status="error" if status == "failed" else status,
                        message=message,
                        data={
                            "attempt": attempt,
                            "total": total,
                            "wait_seconds": wait_seconds,
                            "rncp_code": rncp_code,
                        },
                        error=error,
                    )
                except Exception:
                    pass

            def _dl_reac():
                try:
                    results["reac"] = download_reac_text_with_retry(
                        rncp_code,
                        attempts=3,
                        on_attempt=_log_reac_attempt,
                    )
                except Exception as e:
                    errors.append(f"REAC: {e}")

            def _dl_rc():
                try:
                    results["rc"] = download_rc_text(rncp_code)
                except Exception as e:
                    logger.warning(f"⚠️ RC optionnel non disponible : {e}")

            def _dl_rome():
                try:
                    results["rome"] = fetch_rome_data(rncp_code)
                except Exception as e:
                    logger.warning(f"⚠️ ROME optionnel non disponible : {e}")

            threads = [
                threading.Thread(target=_dl_reac, daemon=True),
                threading.Thread(target=_dl_rc, daemon=True),
                threading.Thread(target=_dl_rome, daemon=True),
            ]
            for t in threads:
                t.start()
            for t in threads:
                t.join()

            if errors:
                raise Exception("; ".join(errors))

            update_job(
                job_id,
                status="reac_ready",
                reac_text=results["reac"],
                rc_text=results["rc"] or None,
                rome_text=results["rome"] or None,
            )
            logger.info(
                f"✅ Job {job_id} : REAC={len(results['reac'])}c "
                f"RC={len(results['rc'])}c ROME={len(results['rome'])}c"
            )
        except Exception as e:
            logger.error(f"❌ Job {job_id} fetch-reac : {e}")
            update_job(job_id, status="error", error_message=str(e))

    threading.Thread(target=_fetch_thread, daemon=True).start()
    return jsonify({"message": "Téléchargement REAC lancé", "status": "reac_fetching"})


# ─── Couche 1 : Enrichissement REAC → Knowledge Base ─────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/enrich-reac", methods=["POST"])
def enrich_reac(job_id):
    """
    Lance la construction de la knowledge base enrichie à partir du REAC.
    Claude extrait les compétences puis les enrichit une par une (définition,
    études de cas, pièges, vocabulaire, contexte terrain, liens connexes).
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if not job.get("reac_text"):
        return jsonify({"error": "REAC non disponible. Lancez d'abord fetch-reac."}), 400

    if job["status"] not in ("reac_ready", "kb_ready", "error", "kb_building"):
        return jsonify({"error": f"Statut '{job['status']}' invalide pour cette action"}), 400

    data = request.get_json() or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))
    launch_kb_building(job_id, model=model)
    return jsonify({"message": "Construction knowledge base lancée", "status": "kb_building", "model": model})


@formation_bp.route("/api/formation/<int:job_id>/kb", methods=["GET"])
def get_kb(job_id):
    """Retourne les entrées KB + statistiques pour un job."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    entries = list_kb(job_id)
    stats = kb_stats(job_id)
    return jsonify({"entries": entries, "stats": stats})


# ─── Génération programme global ──────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/generate-global", methods=["POST"])
def generate_global(job_id):
    """Lance la génération du programme global depuis le REAC + KB enrichie."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if not job.get("reac_text"):
        return jsonify({"error": "REAC non disponible. Lancez d'abord fetch-reac."}), 400

    if job["status"] not in ("reac_ready", "kb_ready", "error", "global_ready"):
        return jsonify({"error": f"Statut '{job['status']}' invalide pour cette action"}), 400

    data = request.get_json() or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))
    launch_global_program_generation(job_id, model=model)
    return jsonify({"message": "Génération programme global lancée", "status": "global_generating", "model": model})


# ─── Validation programme global ──────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/validate-global", methods=["POST"])
def validate_global(job_id):
    """
    Valide (et éventuellement corrige) le programme global.
    Body (optionnel): { "program_text": "...texte édité par l'humain..." }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if job["status"] not in ("global_ready", "daily_ready"):
        return jsonify({"error": "Programme global pas encore généré"}), 400

    data = request.get_json() or {}
    edited_program = data.get("program_text")  # None = garder le texte généré

    update_kwargs = {"global_program_validated": 1, "status": "global_validated"}
    if edited_program:
        update_kwargs["global_program"] = edited_program

    update_job(job_id, **update_kwargs)
    logger.info(f"✅ Job {job_id} : programme global validé{' (édité)' if edited_program else ''}")
    return jsonify({"message": "Programme global validé", "status": "global_validated"})


# ─── Découpage en journées ────────────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/split-daily", methods=["POST"])
def split_daily(job_id):
    """Lance le découpage du programme global en N journées."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if not job.get("global_program"):
        return jsonify({"error": "Programme global non disponible"}), 400

    if job["status"] not in ("global_validated", "global_ready", "error", "daily_ready"):
        return jsonify({"error": f"Statut '{job['status']}' invalide pour cette action"}), 400

    data = request.get_json() or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))
    launch_daily_split(job_id, model=model)
    return jsonify({"message": "Découpage journées lancé", "status": "daily_splitting", "model": model})


# ─── Validation programmes journée ───────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/validate-daily", methods=["POST"])
def validate_daily(job_id):
    """
    Valide (et éventuellement corrige) les programmes journée.
    Body (optionnel): { "daily_programs": [...array JSON édité...] }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if job["status"] not in ("daily_ready", "daily_validated"):
        return jsonify({"error": "Programmes journée pas encore générés"}), 400

    data = request.get_json() or {}
    edited_programs = data.get("daily_programs")  # None = garder les programmes générés

    import json
    update_kwargs = {"daily_programs_validated": 1, "status": "daily_validated"}
    if edited_programs:
        update_kwargs["daily_programs"] = json.dumps(edited_programs, ensure_ascii=False)

    update_job(job_id, **update_kwargs)
    logger.info(f"✅ Job {job_id} : programmes journée validés{' (édités)' if edited_programs else ''}")
    return jsonify({"message": "Programmes journée validés", "status": "daily_validated"})


# ─── Lancement TTS ────────────────────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/launch-tts", methods=["POST"])
def launch_tts(job_id):
    """
    Crée les dossiers cours (un par journée) et lance la génération TTS from scratch.
    Les programmes journée doivent être validés.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    if not job.get("daily_programs_validated"):
        return jsonify({"error": "Les programmes journée doivent être validés avant de lancer le TTS"}), 400

    if job["status"] == "tts_launched":
        return jsonify({"error": "TTS déjà lancé pour ce job"}), 400

    platform_id = job["platform_id"]
    data = request.get_json() or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))

    try:
        folder_ids = launch_tts_for_all_days(job_id, platform_id, model=model)
        return jsonify({
            "message": f"Génération TTS lancée pour {len(folder_ids)} journées",
            "folder_ids": folder_ids,
            "model": model,
            "status": "tts_launched",
        })
    except Exception as e:
        logger.error(f"❌ Job {job_id} launch-tts : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Affinage IA (refine) ─────────────────────────────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/refine", methods=["POST"])
def refine(job_id):
    """
    Affine un contenu généré via une instruction en langage naturel.
    Body: { content_type: "global"|"daily", instruction: "...", current_content: "...", model: "..." }
    Retourne: { revised_content: "..." }
    Appel synchrone (l'utilisateur attend la réponse).
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    data = request.get_json() or {}
    content_type = data.get("content_type", "global")
    instruction = (data.get("instruction") or "").strip()
    current_content = (data.get("current_content") or "").strip()
    model = _resolve_pipeline_api_model(job, data.get("model"))

    if not instruction:
        return jsonify({"error": "Le champ 'instruction' est requis"}), 400
    if not current_content:
        return jsonify({"error": "Le champ 'current_content' est requis"}), 400

    try:
        from services.formation_pipeline_service import refine_content
        revised = refine_content(
            content_type=content_type,
            current_content=current_content,
            instruction=instruction,
            tp_name=job["tp_name"],
            model=model,
        )
        return jsonify({"revised_content": revised})
    except Exception as e:
        logger.error(f"❌ Job {job_id} refine : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Étape 6 : Contenu des journées (lecture + PDF) ──────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/content", methods=["GET"])
def list_content(job_id):
    """
    Retourne la liste des dossiers cours (journées) de ce job avec leur état
    de génération texte : total de mots par journée, nb segments completed,
    statut du content_generation_job associé.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    try:
        from services.formation_pipeline_service import repair_orphan_content_folders
        repair_orphan_content_folders(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Réparation cours_folders job {job_id} ignorée : {e}")

    import json as _json
    from services.formation_pipeline_service import get_expected_course_folders
    from repositories.pipeline_repository import list_content_completion_rows_for_folders
    from services.script_slide_generation_service import get_latest_script_slide_deck

    folder_state = get_expected_course_folders(job_id)
    folders = [
        (
            f["folder_id"],
            f["name"],
            f["position"],
            f["platform_id"],
            f["formation_job_id"],
        )
        for f in folder_state.get("folders", [])
    ]
    folder_ids = [int(folder[0]) for folder in folders]
    content_rows = {
        int(row["folder_id"]): row
        for row in list_content_completion_rows_for_folders(folder_ids)
    }

    daily_programs = _json.loads(job["daily_programs"] or "[]")
    result = []
    for idx, (fid, fname, fpos, f_platform_id, f_formation_job_id) in enumerate(folders):
        day_meta = daily_programs[idx] if idx < len(daily_programs) else {}
        content_row = content_rows.get(int(fid)) or {}
        cg_id = content_row.get("content_job_id")
        cg_status = content_row.get("status")
        cg_words = content_row.get("total_words") or 0
        cur_sub = content_row.get("current_sub_part") or 0
        cur_passe = content_row.get("current_passe") or 1
        cg_err = content_row.get("error_message")
        n_completed = content_row.get("completed_segments") or 0
        n_reviewed = content_row.get("reviewed_segments") or 0
        n_humanized = content_row.get("humanized_segments") or 0
        n_review_errors = content_row.get("review_error_segments") or 0
        n_dirty = content_row.get("dirty_segments") or 0
        slide_deck_id = None
        slide_count = 0
        slide_generation_mode = None
        if cg_id:
            try:
                deck = get_latest_script_slide_deck(int(fid), content_job_id=int(cg_id))
                if deck:
                    slide_deck_id = deck.get("deck_id")
                    slide_count = len(deck.get("slides") or [])
                    slide_generation_mode = (deck.get("stats") or {}).get("generation_mode")
            except Exception:
                slide_deck_id = None
                slide_count = 0
                slide_generation_mode = None

        day_sub_parts = day_meta.get("sub_parts")
        segment_total = (len(day_sub_parts) if day_sub_parts else 6) * 3
        result.append({
            "folder_id": fid,
            "folder_label": f"F{fid}",
            "folder_name": fname,
            "position": fpos,
            "platform_id": f_platform_id,
            "formation_job_id": f_formation_job_id,
            "content_job_id": cg_id,
            "day_number": day_meta.get("day_number", idx + 1),
            "day_title": day_meta.get("title", fname),
            "content_status": cg_status,
            "total_words": cg_words or 0,
            "segments_completed": n_completed,
            "segments_total": max(3, segment_total),
            "segments_humanized": n_humanized,
            "segments_reviewed": n_reviewed,
            "segments_review_errors": n_review_errors,
            "dirty_segments": n_dirty,
            "slide_deck_id": slide_deck_id,
            "slide_count": slide_count,
            "slide_generation_mode": slide_generation_mode,
            "current_sub_part": cur_sub,
            "current_passe": cur_passe,
            "error_message": cg_err,
        })

    return jsonify({
        "folders": result,
        "job_status": job["status"],
        "folder_resolution": {
            "expected_count": folder_state.get("expected_count", 0),
            "duplicates": folder_state.get("duplicates", []),
            "missing": folder_state.get("missing", []),
        },
    })


@formation_bp.route("/api/formation/<int:job_id>/content/<int:folder_id>/text", methods=["GET"])
def get_course_text(job_id, folder_id):
    """Retourne le texte complet d'une journée assemblé section par section (pour relecture UI)."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    try:
        from services.formation_docx_service import (
            _get_segments_for_folder, _strip_tts_tags,
        )
        segments = _get_segments_for_folder(folder_id)
        # Format lisible pour la modal : chaque section précédée de son titre
        sections = []
        for i, part in enumerate(segments):
            title = f"{i + 1}. {part['name']}"
            body = _strip_tts_tags(part['body'])
            sections.append(f"═══ {title} ═══\n\n{body}")
        return jsonify({"text": "\n\n\n".join(sections), "sections_count": len(segments)})
    except Exception as e:
        logger.error(f"❌ Job {job_id} get_course_text folder={folder_id} : {e}")
        return jsonify({"error": str(e)}), 500


@formation_bp.route("/api/formation/<int:job_id>/content/<int:folder_id>/artifact/<path:filename>", methods=["GET"])
def get_content_artifact(job_id, folder_id, filename):
    """Retourne un artefact JSON structuré d'une journée pour l'audit UI."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    from services.content_pipeline.artifacts import (
        CONTENT_ARTIFACT_BLOBS,
        load_content_artifact,
    )
    from repositories.pipeline_repository import get_content_generation_job_by_folder

    filename = os.path.basename(str(filename or ""))
    if filename not in CONTENT_ARTIFACT_BLOBS:
        return jsonify({"error": "Artefact non autorisé"}), 400

    folder_row = get_content_generation_job_by_folder(folder_id)
    if not folder_row or int(folder_row.get("formation_job_id") or 0) != int(job_id):
        return jsonify({"error": "Folder introuvable ou hors pipeline"}), 404

    platform_id = int(folder_row["platform_id"])
    folder_name = folder_row.get("name") or ""
    artifact = load_content_artifact(platform_id, folder_id, filename)
    if not artifact:
        return jsonify({"error": "Artefact indisponible", "filename": filename}), 404

    return jsonify({
        "artifact": artifact,
        "filename": filename,
        "folder_id": folder_id,
        "folder_name": folder_name,
    }), 200


@formation_bp.route("/api/formation/<int:job_id>/content/<int:folder_id>/docx", methods=["GET"])
def download_course_docx(job_id, folder_id):
    """Télécharge le document Word d'une journée de formation (programme officiel).

    Query param `version` :
    - "current" (défaut) : texte DB actuel (= post-révision si appliquée)
    - "pre_review" : snapshot pris au finalize content (= avant révision)
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    version = request.args.get("version", "current")
    if version not in ("current", "pre_review"):
        return jsonify({"error": f"version invalide : {version}"}), 400

    try:
        from services.formation_docx_service import build_course_docx
        docx_bytes, filename = build_course_docx(
            job_id=job_id, folder_id=folder_id, version=version,
        )
        return send_file(
            BytesIO(docx_bytes),
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            as_attachment=True,
            download_name=filename,
        )
    except Exception as e:
        logger.error(f"❌ Job {job_id} download_course_docx folder={folder_id} v={version} : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Rapport de révision conformité ──────────────────────────────────────────

def _short_review_excerpt(text: str, limit: int = 220) -> str:
    text = " ".join((text or "").split())
    if len(text) <= limit:
        return text
    return text[:limit].rstrip() + "..."


def _build_db_review_report(job_id: int, folder_id: int) -> dict | None:
    """Fallback persistant pour les reviews API: reconstruit un rapport depuis
    la DB quand aucun review_report.json Claude Code n'existe.

    Les anciennes reviews API ne stockaient pas l'historique exact des patches.
    On peut tout de même afficher un rapport exploitable: segments relus,
    erreurs éventuelles, et segments dont le texte courant diffère du snapshot
    pre-review.
    """
    from database.db import get_db_connection

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT cf.name, s.id, s.sub_part_index, s.passe,
               COALESCE(s.reviewed, 0), COALESCE(s.review_error, ''),
               COALESCE(s.text_content, ''), COALESCE(s.text_content_pre_review, ''),
               COALESCE(s.word_count, 0)
        FROM cours_folders cf
        JOIN content_generation_jobs cj ON cj.folder_id = cf.id
        JOIN content_generation_segments s ON s.job_id = cj.id
        WHERE cf.id = ? AND cf.formation_job_id = ? AND s.status = 'completed'
        ORDER BY s.sub_part_index ASC, s.passe ASC
        """,
        (folder_id, job_id),
    )
    rows = cursor.fetchall()
    conn.close()
    if not rows:
        return None

    folder_name = rows[0][0]
    by_segment = []
    changed_count = 0
    reviewed_count = 0
    failed_count = 0

    for _, seg_id, sub_idx, passe, reviewed, review_error, text, pre_text, word_count in rows:
        reviewed = int(reviewed or 0)
        if reviewed:
            reviewed_count += 1
        if review_error and not reviewed:
            failed_count += 1

        before = (pre_text or text or "").strip()
        after = (text or "").strip()
        changed = bool(reviewed and pre_text and before != after)
        patches_detail = []
        patches_applied = 0
        patches_rejected = 0

        if changed:
            changed_count += 1
            patches_applied = 1
            patches_detail.append({
                "rule": "#DB",
                "reason": "Texte modifié par la révision API; détail exact des patches non persisté sur les anciennes exécutions.",
                "original": _short_review_excerpt(before),
                "replacement": _short_review_excerpt(after),
                "status": "applied",
            })
        elif review_error:
            patches_rejected = 1
            patches_detail.append({
                "rule": "#ERR",
                "reason": review_error[:220],
                "original": "",
                "replacement": "",
                "status": "rejected",
                "reject_reason": "erreur reviewer",
            })

        by_segment.append({
            "sub_idx": sub_idx,
            "passe": passe,
            "segment_id_actual": seg_id,
            "word_count": word_count,
            "patches_applied": patches_applied,
            "patches_rejected": patches_rejected,
            "patches_detail": patches_detail,
        })

    by_rule = {}
    if changed_count:
        by_rule["#DB"] = {
            "proposed": changed_count,
            "applied": changed_count,
            "rejected": 0,
            "unknown": 0,
        }
    if failed_count:
        by_rule["#ERR"] = {
            "proposed": failed_count,
            "applied": 0,
            "rejected": failed_count,
            "unknown": 0,
        }

    return {
        "folder_id": folder_id,
        "folder_name": folder_name,
        "imported_at": None,
        "generated_via": "db_review_status",
        "is_db_fallback": True,
        "reconstruction_note": (
            "Aucun fichier review_report.json n'a été trouvé. Rapport reconstruit "
            "depuis la base: statut des segments, erreurs reviewer et comparaison "
            "texte courant / snapshot avant révision."
        ),
        "summary": {
            "segments_reviewed": reviewed_count,
            "patches_proposed": changed_count + failed_count,
            "patches_applied": changed_count,
            "patches_rejected": failed_count,
            "patches_unknown": 0,
            "segments_failed": failed_count,
        },
        "by_rule": by_rule,
        "by_segment": by_segment,
    }


def _write_api_review_report(job_id: int, folder_id: int, result: dict, model: str | None) -> dict | None:
    """Persiste un rapport conformité pour les reviews lancées via l'API.

    La DB est la source durable. Le fichier local est seulement un artefact de
    debug, car il peut disparaître ou être inaccessible en hébergement.
    """
    import json as _json
    import os
    from datetime import datetime
    from services.claude_code_mission_service import mission_dir
    from database.db import get_db_connection

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT name, position FROM cours_folders WHERE id = ? AND formation_job_id = ?",
        (folder_id, job_id),
    )
    row = cursor.fetchone()
    conn.close()
    if not row:
        return
    folder_name, position = row
    review_kind = result.get("review_kind") or "compliance"
    review_label = result.get("review_label") or "Révision conformité"

    by_rule = {}
    by_segment = []
    for detail in result.get("details") or []:
        applied = detail.get("applied") or []
        rejected = detail.get("rejected") or []
        patches_detail = []

        for status, patches in (("applied", applied), ("rejected", rejected)):
            for p in patches:
                rule = str(p.get("rule_violated") or p.get("rule") or "?")
                stat = by_rule.setdefault(rule, {"proposed": 0, "applied": 0, "rejected": 0, "unknown": 0})
                stat["proposed"] += 1
                stat[status] += 1
                patches_detail.append({
                    "rule": rule,
                    "reason": str(p.get("reason") or "")[:240],
                    "original": _short_review_excerpt(p.get("original") or ""),
                    "replacement": _short_review_excerpt(p.get("replacement") or ""),
                    "status": status,
                    "reject_reason": p.get("reject_reason"),
                })

        if detail.get("error"):
            stat = by_rule.setdefault("#ERR", {"proposed": 0, "applied": 0, "rejected": 0, "unknown": 0})
            stat["proposed"] += 1
            stat["rejected"] += 1
            patches_detail.append({
                "rule": "#ERR",
                "reason": str(detail.get("error") or "")[:240],
                "original": "",
                "replacement": "",
                "status": "rejected",
                "reject_reason": "erreur reviewer",
            })

        by_segment.append({
            "sub_idx": detail.get("sub_idx"),
            "passe": detail.get("passe"),
            "segment_id_actual": detail.get("segment_id"),
            "patches_proposed": detail.get("proposed", len(applied) + len(rejected)),
            "patches_applied": len(applied),
            "patches_rejected": len(rejected) + (1 if detail.get("error") else 0),
            "patches_detail": patches_detail,
        })

    report = {
        "folder_id": folder_id,
        "folder_name": folder_name,
        "imported_at": datetime.utcnow().isoformat() + "Z",
        "generated_via": model or "api",
        "review_kind": review_kind,
        "review_label": review_label,
        "review_signature": result.get("review_signature"),
        "force": bool(result.get("force")),
        "summary": {
            "segments_reviewed": result.get("segments_reviewed", 0),
            "segments_already_current": result.get("segments_already_current", 0),
            "segments_total_completed": result.get("segments_total_completed", 0),
            "patches_proposed": result.get(
                "patches_proposed",
                result.get("patches_applied", 0) + result.get("patches_rejected", 0),
            ),
            "patches_applied": result.get("patches_applied", 0),
            "patches_rejected": result.get("patches_rejected", 0),
            "segments_failed": result.get("segments_failed", 0),
        },
        "by_rule": by_rule,
        "by_segment": sorted(by_segment, key=lambda x: (x.get("sub_idx") or 0, x.get("passe") or 0)),
    }

    from services.formation_observability_service import persist_review_report
    report_id = persist_review_report(
        job_id,
        folder_id,
        report,
        source="api" if review_kind == "compliance" else f"{review_kind}_api",
        generated_via=model or "api",
    )
    report["persisted_report_id"] = report_id

    suffix = "review_api" if review_kind == "compliance" else f"review_{review_kind}_api"
    chunk_id = f"day_{int(position or 0) + 1}_{suffix}"
    chunk_dir = os.path.join(mission_dir(job_id, "review"), chunk_id)
    try:
        os.makedirs(chunk_dir, exist_ok=True)
        with open(os.path.join(chunk_dir, "review_report.json"), "w", encoding="utf-8") as f:
            _json.dump(report, f, ensure_ascii=False, indent=2)
    except Exception as e:
        logger.warning(
            f"⚠️ Rapport conformité fichier non écrit job={job_id} "
            f"folder={folder_id} : {e}"
        )
    return report


def _review_chunk_ids_for_position(position: int) -> list[str]:
    from services.claude_code_mission_service import _REVIEW_RULE_GROUPS

    day = int(position or 0) + 1
    return (
        [f"day_{day}_review_{g['id']}" for g in _REVIEW_RULE_GROUPS]
        + [
            f"day_{day}_review_api",
            f"day_{day}_review_local_compliance_api",
            f"day_{day}_review_humanization_api",
            f"day_{day}_review",
        ]
    )


def _humanization_chunk_ids_for_position(position: int) -> list[str]:
    from services.content_generation_service import _HUMANIZATION_REVIEW_RULE_GROUPS

    day = int(position or 0) + 1
    return [f"day_{day}_review_{g['id']}" for g in _HUMANIZATION_REVIEW_RULE_GROUPS]


def _delete_active_review_artifacts(job_id: int, position: int) -> int:
    """Supprime les rapports actifs de cette journée avant une relance aval.

    Les archives `_done` sont conservées, mais la route de lecture les filtre
    ensuite par date de relance pour ne pas afficher un rapport ancien.
    """
    import os
    import shutil
    from services.claude_code_mission_service import mission_dir

    deleted = 0
    for step_key, chunk_ids in (
        ("review", _review_chunk_ids_for_position(position)),
        ("humanization_review", _humanization_chunk_ids_for_position(position)),
    ):
        base_dir = mission_dir(job_id, step_key)
        for chunk_id in chunk_ids:
            path = os.path.join(base_dir, chunk_id)
            if os.path.isdir(path):
                shutil.rmtree(path)
                deleted += 1
    return deleted


def _parse_report_timestamp(value):
    if not value:
        return None
    from datetime import datetime, timezone

    raw = str(value).strip()
    if not raw:
        return None
    iso = raw.replace(" ", "T")
    if iso.endswith("Z"):
        iso = iso[:-1] + "+00:00"
    try:
        dt = datetime.fromisoformat(iso)
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.timestamp()


def _latest_continue_after_text_started_at(job_id: int, folder_id: int) -> str | None:
    from database.db import get_db_connection

    try:
        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute(
            """
            SELECT created_at
            FROM formation_pipeline_events
            WHERE job_id = ? AND folder_id = ?
              AND event_type = 'continue_after_text_started'
            ORDER BY created_at DESC, id DESC
            LIMIT 1
            """,
            (job_id, folder_id),
        )
        row = cursor.fetchone()
        conn.close()
        return row[0] if row else None
    except Exception:
        return None


def _report_is_after_cutoff(report: dict | None, cutoff: str | None) -> bool:
    if not cutoff:
        return True
    if not report:
        return False
    report_ts = (
        report.get("persisted_at")
        or report.get("imported_at")
        or report.get("created_at")
    )
    report_time = _parse_report_timestamp(report_ts)
    cutoff_time = _parse_report_timestamp(cutoff)
    if report_time is None or cutoff_time is None:
        return False
    return report_time >= cutoff_time


def _db_review_report_is_complete(report: dict | None) -> bool:
    if not report:
        return False
    segments = report.get("by_segment") or []
    if not segments:
        return False
    summary = report.get("summary") or {}
    processed = int(summary.get("segments_reviewed") or 0) + int(summary.get("segments_failed") or 0)
    return processed >= len(segments)


@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/review-report",
    methods=["GET"],
)
def get_review_report(job_id, folder_id):
    """Retourne le rapport JSON détaillé de la révision conformité pour 1
    journée. Lit `review_queue/job_X/step_review/day_N_review/review_report.json`
    (chunked) ou son équivalent archivé dans `_done/`.

    Format de retour :
    {
      "summary": {segments_reviewed, patches_proposed, patches_applied,
                  patches_rejected, segments_failed},
      "by_rule": {"#22": {proposed, applied, rejected}, ...},
      "by_segment": [{sub_idx, passe, patches_applied, patches_rejected,
                      patches_detail: [{rule, original, replacement,
                                        reason, status, reject_reason?}]}],
      "imported_at", "generated_via", "via_positional_fallback"
    }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    import os
    import json as _json
    from services.claude_code_mission_service import (
        mission_dir, _DONE_ROOT, _REVIEW_QUEUE_ROOT,
    )

    # Trouver la position du folder pour reconstruire le chunk_id (day_N_review)
    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    from database.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT position FROM cours_folders WHERE id = ? AND formation_job_id = ?",
        (folder_id, job_id),
    )
    row = cursor.fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "Folder introuvable ou hors pipeline"}), 404
    position = row[0]
    retry_cutoff = _latest_continue_after_text_started_at(job_id, folder_id)
    stale_report_ignored = False

    try:
        from services.formation_observability_service import get_latest_review_report
        persisted_report = get_latest_review_report(job_id, folder_id)
        if persisted_report and _report_is_after_cutoff(persisted_report, retry_cutoff):
            return jsonify({"report": persisted_report, "source": "db"}), 200
        if persisted_report:
            stale_report_ignored = True
    except Exception as e:
        logger.warning(
            f"⚠️ Lecture rapport conformité DB impossible job={job_id} "
            f"folder={folder_id} : {e}"
        )

    # Multi-agents : 1 chunk_dir par groupe de règles. On cherche le rapport
    # courant, mais après une relance aval on ignore tout rapport plus ancien.
    chunk_id_candidates = _review_chunk_ids_for_position(position)

    if os.path.isdir(_DONE_ROOT):
        archived = sorted(
            (d for d in os.listdir(_DONE_ROOT) if d.endswith(f"-job{job_id}-review")),
            reverse=True,
        )
    else:
        archived = []

    # Collecte tous les review_report.json existants (multi-chunks)
    sub_reports = []
    for cid in chunk_id_candidates:
        paths = [os.path.join(mission_dir(job_id, "review"), cid, "review_report.json")]
        for arch in archived:
            paths.append(os.path.join(_DONE_ROOT, arch, cid, "review_report.json"))
        for p in paths:
            if os.path.exists(p):
                try:
                    with open(p, "r", encoding="utf-8") as f:
                        loaded_report = _json.load(f)
                    if _report_is_after_cutoff(loaded_report, retry_cutoff):
                        sub_reports.append((cid, loaded_report))
                    else:
                        stale_report_ignored = True
                    break  # 1 seul par chunk_id (le plus récent)
                except Exception:
                    pass

    if sub_reports:
        # Agrégation des sub_reports en un rapport unique
        agg_summary = {
            "segments_reviewed": 0, "patches_proposed": 0,
            "patches_applied": 0, "patches_rejected": 0, "segments_failed": 0,
        }
        agg_by_rule = {}
        agg_by_segment = {}  # clé (sub_idx, passe) → segment agrégé
        any_positional = False
        latest_imported = None
        generated_via = None

        for cid, rep in sub_reports:
            s = rep.get("summary", {})
            for k in agg_summary:
                agg_summary[k] = max(agg_summary[k], s.get(k, 0)) if k == "segments_reviewed" else agg_summary[k] + s.get(k, 0)
            for rule, st in (rep.get("by_rule") or {}).items():
                agg = agg_by_rule.setdefault(rule, {"proposed": 0, "applied": 0, "rejected": 0, "unknown": 0})
                for kk, vv in st.items():
                    agg[kk] = agg.get(kk, 0) + vv
            for seg in (rep.get("by_segment") or []):
                key = (seg.get("sub_idx"), seg.get("passe"))
                if key not in agg_by_segment:
                    agg_by_segment[key] = {
                        "sub_idx": seg.get("sub_idx"),
                        "passe": seg.get("passe"),
                        "segment_id_actual": seg.get("segment_id_actual"),
                        "patches_applied": 0, "patches_rejected": 0,
                        "patches_detail": [],
                    }
                agg_by_segment[key]["patches_applied"] += seg.get("patches_applied", 0)
                agg_by_segment[key]["patches_rejected"] += seg.get("patches_rejected", 0)
                for d in seg.get("patches_detail") or []:
                    d_copy = dict(d)
                    d_copy["agent_group"] = cid.replace(f"day_{position + 1}_review_", "")
                    agg_by_segment[key]["patches_detail"].append(d_copy)
            if rep.get("via_positional_fallback"):
                any_positional = True
            if rep.get("imported_at") and (latest_imported is None or rep["imported_at"] > latest_imported):
                latest_imported = rep["imported_at"]
            if not generated_via:
                generated_via = rep.get("generated_via")

        # by_segment trié par (sub_idx, passe)
        agg_segments = sorted(agg_by_segment.values(), key=lambda x: (x["sub_idx"] or 0, x["passe"] or 0))

        report = {
            "folder_id": folder_id,
            "folder_name": (sub_reports[0][1].get("folder_name") if sub_reports else None),
            "imported_at": latest_imported,
            "generated_via": generated_via,
            "via_positional_fallback": any_positional,
            "n_agents": len(sub_reports),
            "agents_used": [cid.replace(f"day_{position + 1}_review_", "") for cid, _ in sub_reports],
            "summary": agg_summary,
            "by_rule": agg_by_rule,
            "by_segment": agg_segments,
        }
        try:
            from services.formation_observability_service import persist_review_report
            persist_review_report(
                job_id,
                folder_id,
                report,
                source="file_import",
                generated_via=generated_via,
            )
        except Exception as e:
            logger.warning(
                f"⚠️ Rapport conformité fichier non persisté job={job_id} "
                f"folder={folder_id} : {e}"
            )
        return jsonify({"report": report, "n_sub_reports": len(sub_reports)}), 200

    # Fallback : pas de review_report.json mais peut-être un output.md
    # existant d'un import antérieur (avant que je code l'écriture du
    # rapport). On reconstitue un rapport "best-effort" depuis output.md +
    # input.md + DB actuelle (heuristique : un patch est "appliqué" si son
    # replacement est dans le texte courant et son original n'y est plus).
    # Cherche dans tous les chunk_id candidats (multi-agents + legacy).
    output_md_paths = []
    for cid in chunk_id_candidates:
        output_md_paths.append(
            os.path.join(mission_dir(job_id, "review"), cid, "output.md")
        )
        for arch in archived:
            output_md_paths.append(
                os.path.join(_DONE_ROOT, arch, cid, "output.md")
            )

    chunk_dir_with_output = None
    if not retry_cutoff:
        for p in output_md_paths:
            if os.path.exists(p):
                chunk_dir_with_output = os.path.dirname(p)
                break

    if not chunk_dir_with_output:
        report = _build_db_review_report(job_id, folder_id)
        if report and (not retry_cutoff or _db_review_report_is_complete(report)):
            return jsonify({"report": report, "source": "db_fallback"}), 200
        if retry_cutoff:
            return jsonify({
                "error": "Nouveau rapport de révision pas encore disponible pour cette relance",
                "report": None,
                "retry_started_at": retry_cutoff,
                "stale_report_ignored": stale_report_ignored,
            }), 404
        return jsonify({
            "error": "Aucun rapport de révision trouvé pour ce dossier",
            "report": None,
        }), 404

    try:
        with open(os.path.join(chunk_dir_with_output, "output.md"), "r", encoding="utf-8") as f:
            output_text = f.read()
        from services.claude_code_mission_service import _extract_json
        parsed = _json.loads(_extract_json(output_text))
        reviews = parsed.get("reviews", [])
    except Exception as e:
        return jsonify({"error": f"output.md illisible : {e}"}), 500

    # Récupère input.md (segments d'origine) pour la résolution positionnelle
    input_md_path = os.path.join(chunk_dir_with_output, "input.md")
    input_segments = []
    if os.path.exists(input_md_path):
        try:
            with open(input_md_path, "r", encoding="utf-8") as f:
                input_segments = _json.loads(f.read())
        except Exception:
            input_segments = []

    # Récupère les textes actuels en DB (résolution par sub_idx, passe via folder)
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT s.id, s.sub_part_index, s.passe, s.text_content
        FROM content_generation_segments s
        JOIN content_generation_jobs cj ON cj.id = s.job_id
        WHERE cj.folder_id = ? AND s.status = 'completed'
        ORDER BY cj.id DESC, s.sub_part_index ASC, s.passe ASC
        """,
        (folder_id,),
    )
    db_rows = cursor.fetchall()
    conn.close()
    # Map (sub_idx, passe) → text actuel (le plus récent en cas de doublons)
    db_text_by_sp = {}
    for r in db_rows:
        key = (r[1], r[2])
        if key not in db_text_by_sp:
            db_text_by_sp[key] = (r[0], r[3])

    # Construction du rapport lite
    by_rule = {}
    by_segment = []
    applied_total = 0
    rejected_total = 0
    unknown_total = 0

    for i, rev in enumerate(reviews):
        patches = rev.get("patches", []) or []
        # Résolution positionnelle via input.md
        if i < len(input_segments):
            sub_idx = input_segments[i].get("sub_idx", -1)
            passe = input_segments[i].get("passe", -1)
        else:
            sub_idx = -1
            passe = -1
        seg_db = db_text_by_sp.get((sub_idx, passe))
        current_text = seg_db[1] if seg_db else ""
        actual_segment_id = seg_db[0] if seg_db else None

        seg_applied = 0
        seg_rejected = 0
        seg_unknown = 0
        seg_patches_detail = []
        for p in patches[:5]:
            original = (p.get("original") or "")[:1000]
            replacement = (p.get("replacement") or "")[:1000]
            rule = str(p.get("rule_violated", "?"))[:10]
            reason = str(p.get("reason", ""))[:200]
            if not original or not replacement:
                continue
            rule_stat = by_rule.setdefault(rule, {"proposed": 0, "applied": 0, "rejected": 0, "unknown": 0})
            rule_stat["proposed"] += 1
            # Heuristique
            if current_text:
                has_repl = replacement in current_text
                has_orig = original in current_text
                if has_repl and not has_orig:
                    status = "applied"
                    seg_applied += 1
                    rule_stat["applied"] += 1
                elif has_orig:
                    status = "rejected"
                    seg_rejected += 1
                    rule_stat["rejected"] += 1
                else:
                    status = "unknown"
                    seg_unknown += 1
                    rule_stat["unknown"] += 1
            else:
                status = "unknown"
                seg_unknown += 1
                rule_stat["unknown"] += 1
            seg_patches_detail.append({
                "rule": rule, "reason": reason,
                "original": original[:200], "replacement": replacement[:200],
                "status": status,
                "reject_reason": "détecté heuristiquement" if status == "rejected" else None,
            })

        applied_total += seg_applied
        rejected_total += seg_rejected
        unknown_total += seg_unknown

        by_segment.append({
            "sub_idx": sub_idx,
            "passe": passe,
            "segment_id_actual": actual_segment_id,
            "patches_applied": seg_applied,
            "patches_rejected": seg_rejected,
            "patches_unknown": seg_unknown,
            "patches_detail": seg_patches_detail,
        })

    report = {
        "folder_id": folder_id,
        "folder_name": None,  # pas dispo en mode reconstruit
        "imported_at": None,
        "generated_via": "reconstructed_from_output_md",
        "is_reconstructed": True,
        "reconstruction_note": (
            "Rapport reconstitué a posteriori depuis output.md + DB actuelle. "
            "Statuts applied/rejected déduits par heuristique (replacement "
            "présent dans le texte = appliqué, original encore présent = rejeté). "
            "Précision approximative."
        ),
        "summary": {
            "segments_reviewed": len(reviews),
            "patches_proposed": sum(len(r.get("patches", [])) for r in reviews),
            "patches_applied": applied_total,
            "patches_rejected": rejected_total,
            "patches_unknown": unknown_total,
            "segments_failed": 0,
        },
        "by_rule": by_rule,
        "by_segment": by_segment,
    }
    try:
        from services.formation_observability_service import persist_review_report
        persist_review_report(
            job_id,
            folder_id,
            report,
            source="output_md_reconstruction",
            generated_via=report["generated_via"],
        )
    except Exception as e:
        logger.warning(
            f"⚠️ Rapport conformité reconstruit non persisté job={job_id} "
            f"folder={folder_id} : {e}"
        )
    return jsonify({"report": report, "source_path": chunk_dir_with_output, "lite": True}), 200


@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/humanization-report",
    methods=["GET"],
)
def get_humanization_report(job_id, folder_id):
    """Retourne le rapport JSON de la passe humanisation (intros/transitions/rythme)."""
    from services.formation_observability_service import get_latest_review_report
    report = get_latest_review_report(job_id, folder_id, kind="humanization")
    if not report:
        return jsonify({"error": "Aucun rapport d'humanisation disponible pour cette journée"}), 404
    return jsonify({"report": report}), 200


# ─── Legacy — audit volume lisible, enrichissement append-only désactivé ─────

@formation_bp.route("/api/formation/<int:job_id>/volume-audit", methods=["GET"])
def volume_audit(job_id):
    """Retourne l'audit volume par-folder pour un job.

    Pour chaque folder du job, calcule :
      - total_words : mots parlés des segments completed, tags TTS exclus
      - deficit : mots manquants sous le budget minimal audio
      - shortest_segments : top N (5) des segments les plus courts (pour
        l'affichage UI et le ciblage de l'enrichissement)

    Pas d'effet de bord — purement lecture.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    from services.claude_code_mission_service import compute_volume_audit
    try:
        audit = compute_volume_audit(job_id)
        return jsonify(audit), 200
    except Exception as e:
        logger.error(f"❌ volume_audit job {job_id} : {e}")
        return jsonify({"error": str(e)}), 500


@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/volume-safety",
    methods=["POST"],
)
def launch_volume_safety(job_id, folder_id):
    """Ancienne sécurité volume append-only, désactivée.

    Le rattrapage de volume doit se faire pendant le calibrage budget texte,
    avant les reviews, avec le plan verrouillé et les sections comme contexte.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    return jsonify({
        "error": (
            "Sécurité volume désactivée : cette ancienne réparation append-only "
            "pouvait ajouter du texte après les conclusions/Q-R. Relance la "
            "génération structurée pour utiliser le calibrage budget texte."
        )
    }), 410


@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/volume-safety/status",
    methods=["GET"],
)
def volume_safety_status(job_id, folder_id):
    """Récupère l'état de la dernière exécution volume safety pour ce folder.

    Retourne {status: 'idle'|'running'|'done'|'error', model?, result?, error?}.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    from services.claude_code_mission_service import _EXECUTION_STATE
    state = _EXECUTION_STATE.get((job_id, f"volume_safety_{folder_id}"))
    if not state:
        return jsonify({"status": "idle"}), 200
    return jsonify(state), 200


# ─── Reprise de la génération texte (après crash / restart backend) ──────────

@formation_bp.route("/api/formation/<int:job_id>/resume-content", methods=["POST"])
def resume_content(job_id):
    """
    Relance la génération texte sur tous les dossiers du job qui ne sont pas
    encore 'completed'. Utilise le checkpointing de run_content_generation :
    les segments déjà persistés en DB sont skippés automatiquement. Aucun
    dossier n'est recréé, aucun segment n'est effacé.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    from database.db import get_db_connection
    from services.formation_pipeline_service import get_expected_course_folders
    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    if not folder_ids:
        return jsonify({"error": "Aucun cours_folder attendu pour ce job"}), 400
    placeholders = ",".join("?" * len(folder_ids))
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        f"""SELECT f.id, cg.status
           FROM cours_folders f
           LEFT JOIN content_generation_jobs cg ON cg.folder_id = f.id
           WHERE f.id IN ({placeholders}) ORDER BY f.position ASC, f.id ASC""",
        tuple(folder_ids),
    )
    rows = cursor.fetchall()
    conn.close()

    to_resume = [fid for fid, status in rows if status != "completed"]
    if not to_resume:
        return jsonify({"message": "Tous les dossiers sont déjà complets", "resumed": []})

    data = request.get_json(silent=True) or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))

    import eventlet
    from services.content_generation_service import run_content_generation

    def _resume_one(folder_id):
        try:
            logger.info(f"♻️ Job {job_id} folder {folder_id} : reprise génération texte")
            run_content_generation(folder_id, mode="normal", model=model)
            logger.info(f"✅ Job {job_id} folder {folder_id} : reprise terminée")
        except Exception as e:
            logger.error(f"❌ Job {job_id} folder {folder_id} : reprise échouée : {e}")

    for fid in to_resume:
        eventlet.spawn(_resume_one, fid)

    logger.info(f"♻️ Job {job_id} : reprise texte pour {len(to_resume)} dossier(s)")
    return jsonify({
        "message": f"Reprise lancée pour {len(to_resume)} dossier(s)",
        "resumed": to_resume,
    })


# ─── Étape 6bis : Révision conformité via reviewer API Claude ────────────────
# Phase 1 de memoire/03-decisions/pipeline-dual-api-et-claude-code.md
# Scope strict : API uniquement, pas de refonte UI double colonne ici.

@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/review",
    methods=["POST"],
)
def review_content(job_id, folder_id):
    """
    Lance la révision conformité du texte généré pour un dossier cours.
    Relit les segments non validés ou validés avec une ancienne signature de
    règles. `force=true` relance explicitement tous les segments completed.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job pipeline introuvable"}), 404

    # Vérifier que le folder appartient bien à ce job formation
    from database.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        "SELECT id FROM cours_folders WHERE id = ? AND formation_job_id = ?",
        (folder_id, job_id),
    )
    row = cursor.fetchone()
    conn.close()
    if not row:
        return jsonify({"error": "Folder inexistant ou hors pipeline"}), 404

    data = request.get_json(silent=True) or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))
    force = bool(data.get("force") or data.get("force_review"))

    import eventlet
    from services.content_generation_service import run_content_review

    def _run_review(_folder_id):
        import sys, traceback
        logger.info(f"🚀 SPAWN review greenlet job={job_id} folder={_folder_id} model={model} force={force}")
        sys.stdout.flush()
        started_at = time.time()
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "review_started",
                step="review",
                status="running",
                folder_id=_folder_id,
                model=str(model) if model else None,
                message="Révision conformité API démarrée",
                data={"force": force},
            )
        except Exception:
            pass
        try:
            result = run_content_review(_folder_id, model=model, force=force)
            duration_ms = int((time.time() - started_at) * 1000)
            try:
                _write_api_review_report(job_id, _folder_id, result, model)
            except Exception as report_error:
                logger.error(
                    f"❌ Rapport review API non persisté job={job_id} "
                    f"folder={_folder_id} : {report_error}"
                )
                raise
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "review_completed",
                    step="review",
                    status="completed",
                    folder_id=_folder_id,
                    model=str(model) if model else None,
                    duration_ms=duration_ms,
                    message="Révision conformité API terminée",
                    data={
                        "segments_reviewed": result.get("segments_reviewed", 0),
                        "segments_already_current": result.get("segments_already_current", 0),
                        "patches_proposed": result.get("patches_proposed", 0),
                        "patches_applied": result.get("patches_applied", 0),
                        "patches_rejected": result.get("patches_rejected", 0),
                        "segments_failed": result.get("segments_failed", 0),
                        "review_signature": result.get("review_signature"),
                        "force": result.get("force", force),
                    },
                )
            except Exception:
                pass
            logger.info(
                f"✅ Review job={job_id} folder={_folder_id} : "
                f"{result['segments_reviewed']} audités, "
                f"{result.get('patches_proposed', 0)} proposés, "
                f"{result['patches_applied']} appliqués, "
                f"{result['patches_rejected']} rejetés"
            )
            sys.stdout.flush()
        except Exception as e:
            duration_ms = int((time.time() - started_at) * 1000)
            logger.error(f"❌ Review job={job_id} folder={_folder_id} : échec : {e}")
            logger.error(traceback.format_exc())
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "review_failed",
                    step="review",
                    status="error",
                    folder_id=_folder_id,
                    model=str(model) if model else None,
                    duration_ms=duration_ms,
                    message="Révision conformité API échouée",
                    error=str(e)[:500],
                )
            except Exception:
                pass
            sys.stdout.flush()

    eventlet.spawn(_run_review, folder_id)

    return jsonify({
        "message": "Révision conformité lancée en arrière-plan",
        "folder_id": folder_id,
        "model": model,
        "force": force,
    }), 202


# ─── Missions Claude Code local (Phase 3) ────────────────────────────────────
# Export/import manuel de tâches à faire dans `claude --model haiku|sonnet`.
# Gating : nécessite LOCAL_DEV=true côté backend. Spec complète :
# memoire/03-decisions/pipeline-dual-api-et-claude-code.md

@formation_bp.route(
    "/api/formation/<int:job_id>/missions/<string:step_key>/export",
    methods=["POST"],
)
def export_claude_code_mission(job_id, step_key):
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    from services.claude_code_mission_service import (
        export_mission, local_dev_enabled,
    )
    if not local_dev_enabled():
        return jsonify({"error": "LOCAL_DEV non activé — fonctionnalité dev uniquement"}), 403
    data = request.get_json(silent=True) or {}
    model = data.get("model", "haiku")
    try:
        mission = export_mission(job_id, step_key, model)
    except PermissionError as e:
        return jsonify({"error": str(e)}), 403
    except (ValueError, FileNotFoundError) as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        logger.error(f"❌ export_claude_code_mission : {e}")
        return jsonify({"error": str(e)}), 500
    return jsonify({"mission": mission}), 201


@formation_bp.route(
    "/api/formation/<int:job_id>/missions/<string:step_key>/import",
    methods=["POST"],
)
def import_claude_code_mission(job_id, step_key):
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    from services.claude_code_mission_service import (
        import_mission_result, local_dev_enabled,
    )
    if not local_dev_enabled():
        return jsonify({"error": "LOCAL_DEV non activé — fonctionnalité dev uniquement"}), 403
    try:
        result = import_mission_result(job_id, step_key)
    except PermissionError as e:
        return jsonify({"error": str(e)}), 403
    except FileNotFoundError as e:
        return jsonify({"error": str(e)}), 400
    except NotImplementedError as e:
        # Étapes kb / content : parsers non implémentés en V1. On renvoie un
        # vrai 501 pour que le frontend affiche l'erreur et ne supprime pas
        # la mission de la file (cf. audit point #2).
        return jsonify({"error": str(e), "not_implemented": True}), 501
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception as e:
        logger.error(f"❌ import_claude_code_mission : {e}")
        return jsonify({"error": str(e)}), 500
    return jsonify(result), 200


@formation_bp.route(
    "/api/formation/<int:job_id>/missions/<string:step_key>/execute",
    methods=["POST"],
)
def execute_claude_code_mission(job_id, step_key):
    """Lance `claude -p` en subprocess (greenlet eventlet) + import auto à la fin.
    Dev local only (LOCAL_DEV=true côté backend ET `claude` dans le PATH)."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    from services.claude_code_mission_service import (
        execute_mission_locally, local_dev_enabled,
    )
    if not local_dev_enabled():
        return jsonify({"error": "LOCAL_DEV non activé — fonctionnalité dev uniquement"}), 403
    data = request.get_json(silent=True) or {}
    model = data.get("model", "haiku")

    # État partagé de l'exécution par (job_id, step_key) — mis à jour par le
    # greenlet et lu par le polling frontend via /missions/pending.
    import eventlet
    from services.claude_code_mission_service import _EXECUTION_STATE

    key = (job_id, step_key)
    if _EXECUTION_STATE.get(key, {}).get("status") == "running":
        return jsonify({"error": "Une exécution est déjà en cours pour cette étape"}), 409

    _EXECUTION_STATE[key] = {"status": "running", "model": model, "error": None, "result": None}

    def _run():
        import sys, traceback
        try:
            logger.info(f"🤖 Exec Claude Code : job={job_id} step={step_key} model={model}")
            sys.stdout.flush()
            result = execute_mission_locally(job_id, step_key, model)
            _EXECUTION_STATE[key] = {"status": "done", "model": model, "error": None, "result": result}
            logger.info(f"✅ Exec Claude Code terminée : job={job_id} step={step_key}")
            sys.stdout.flush()
        except Exception as e:
            logger.error(f"❌ Exec Claude Code échouée : job={job_id} step={step_key} : {e}")
            logger.error(traceback.format_exc())
            _EXECUTION_STATE[key] = {"status": "error", "model": model, "error": str(e)[:500], "result": None}
            sys.stdout.flush()

    eventlet.spawn(_run)
    return jsonify({"status": "running", "step_key": step_key, "model": model}), 202


@formation_bp.route(
    "/api/formation/<int:job_id>/missions/<string:step_key>/logs",
    methods=["GET"],
)
def get_claude_code_mission_logs(job_id, step_key):
    """Retourne les N dernières lignes du execution.log de la mission.
    Cherche dans review_queue/job_X/step_Y/ d'abord, puis dans _done/ en
    fallback pour les missions déjà importées et archivées."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    from services.claude_code_mission_service import (
        mission_dir, _DONE_ROOT, local_dev_enabled,
    )
    import os
    if not local_dev_enabled():
        return jsonify({"logs": "", "source": "disabled"}), 200

    tail_n = int(request.args.get("tail", 200))
    # Chemin actuel
    active = os.path.join(mission_dir(job_id, step_key), "execution.log")
    # Fallback : chercher le dernier _done correspondant
    archive_path = None
    if os.path.isdir(_DONE_ROOT):
        candidates = sorted(
            (d for d in os.listdir(_DONE_ROOT) if d.endswith(f"-job{job_id}-{step_key}")),
            reverse=True,
        )
        if candidates:
            archive_path = os.path.join(_DONE_ROOT, candidates[0], "execution.log")

    source = None
    path = None
    if os.path.exists(active):
        path, source = active, "active"
    elif archive_path and os.path.exists(archive_path):
        path, source = archive_path, "archived"

    if not path:
        return jsonify({"logs": "", "source": "not_found"}), 200

    try:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            lines = f.readlines()
        tail = "".join(lines[-tail_n:])
    except Exception as e:
        return jsonify({"error": str(e)}), 500

    return jsonify({
        "logs": tail,
        "source": source,
        "total_lines": len(lines),
        "returned_lines": min(tail_n, len(lines)),
    }), 200


@formation_bp.route("/api/formation/<int:job_id>/missions/pending", methods=["GET"])
def list_claude_code_pending_missions(job_id):
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    from services.claude_code_mission_service import (
        list_pending_missions, local_dev_enabled,
    )
    if not local_dev_enabled():
        return jsonify({"missions": {}}), 200
    return jsonify({"missions": list_pending_missions(job_id)}), 200


# ─── Pre-flight et health-check (audit pipeline) ─────────────────────────────

@formation_bp.route("/api/formation/<int:job_id>/preflight", methods=["POST"])
def preflight_pipeline(job_id):
    """Audit AVANT lancement : valide que la pipeline a les chances de tourner
    one-shot (config + connectivité externe). Body :
      { "use_claude_code": bool, "tts_mode": "fish_audio"|"gtts"|"mock" }
    Retourne :
      { "ok": bool, "blocking": [...], "warnings": [...], "checks": {...} }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    payload = request.get_json(silent=True) or {}
    use_cc = bool(payload.get("use_claude_code", False))
    tts_mode = (payload.get("tts_mode") or "gtts").lower()
    try:
        from services.formation_health_service import compute_preflight
        result = compute_preflight(job_id, use_claude_code=use_cc, tts_mode=tts_mode)
        return jsonify(result), 200 if result["ok"] else 422
    except Exception as e:
        logger.error(f"❌ preflight job {job_id} : {e}")
        return jsonify({"error": str(e)}), 500


@formation_bp.route("/api/formation/<int:job_id>/health", methods=["GET"])
def health_pipeline(job_id):
    """Audit APRÈS lancement : vérifie que la pipeline a produit des artefacts
    cohérents (segments, DOCX, snapshots, audio, module persistant).
    Retourne :
      { "ok": bool, "blocking": [...], "warnings": [...], "checks": {...} }
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    try:
        from services.formation_health_service import compute_health
        result = compute_health(job_id)
        return jsonify(result), 200
    except Exception as e:
        logger.error(f"❌ health-check job {job_id} : {e}")
        return jsonify({"error": str(e)}), 500


# ─── Étape 7 : Lancement de la synthèse TTS Fish Audio ───────────────────────

def _make_audio_progress_logger(job_id: int, folder_id: int, voice_type: str, schedule_session_id: int | None = None):
    """Callback branché sur generate_audio_from_script pour sortir de la boîte noire."""
    last_session_touch = {"at": 0.0}

    def _on_progress(step, total, message):
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "audio_progress",
                step="audio",
                status="running",
                folder_id=folder_id,
                message=message,
                data={
                    "step": step,
                    "total": total,
                    "voice_type": voice_type,
                    "tts_engine": "edge-tts" if voice_type == "gtts" else voice_type,
                },
            )
        except Exception:
            pass
        if schedule_session_id:
            now = time.time()
            if now - last_session_touch["at"] >= 60:
                last_session_touch["at"] = now
                try:
                    from datetime import datetime
                    from config import FRANCE_TZ
                    from database.db import get_db_connection as _get_conn

                    now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
                    _conn = _get_conn()
                    _cur = _conn.cursor()
                    _cur.execute(
                        """
                        UPDATE course_sessions
                        SET audio_generation_status = 'running',
                            updated_at = ?
                        WHERE id = ?
                          AND audio_generation_completed_at IS NULL
                        """,
                        (now_str, int(schedule_session_id)),
                    )
                    _conn.commit()
                    _conn.close()
                except Exception:
                    logger.warning(
                        "⚠️ Impossible de mettre à jour le heartbeat audio session=%s",
                        schedule_session_id,
                        exc_info=True,
                    )
    return _on_progress


def _finalize_audio_ready_state(job_id: int, voice_type: str) -> dict:
    """Marque la plateforme exploitable et garantit le module persistant.

    Idempotent : utilisé après une relance audio manuelle, notamment
    `continue_after_text`, où les MP3 sont déjà produits mais la finalisation
    historique de `launch_audio` n'était pas rejouée.
    """
    from datetime import datetime as _dt
    from config import FRANCE_TZ as _tz
    from database.db import get_db_connection

    job = get_job(job_id)
    if not job:
        raise ValueError(f"Job {job_id} introuvable pour finalisation audio")

    platform_id = job.get("platform_id")
    rncp = job.get("rncp_code") or ""
    tp_name = job.get("tp_name") or f"Job {job_id}"

    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT center_account_id FROM platform_config WHERE id = ?", (platform_id,))
        center_row = cursor.fetchone()
        center_account_id = center_row[0] if center_row else None
        cursor.execute(
            "UPDATE platform_config SET status = 'ready' WHERE id = ? AND status = 'pending'",
            (platform_id,),
        )
        platform_ready_updated = cursor.rowcount or 0

        cursor.execute(
            "SELECT id, version FROM formation_modules WHERE source_pipeline_job_id = ?",
            (job_id,),
        )
        existing = cursor.fetchone()

        if existing:
            module_id, version = existing[0], existing[1]
            cursor.execute(
                """UPDATE formation_modules
                   SET voice_type = ?, voice_updated_at = CURRENT_TIMESTAMP,
                       source_platform_id = COALESCE(source_platform_id, ?),
                       center_account_id = COALESCE(center_account_id, ?),
                       status = 'validated',
                       validated_at = COALESCE(validated_at, CURRENT_TIMESTAMP)
                   WHERE id = ?""",
                (voice_type, platform_id, center_account_id, module_id),
            )
            module_created = False
        else:
            if center_account_id is None:
                cursor.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id IS NULL",
                    (rncp,),
                )
            else:
                cursor.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id = ?",
                    (rncp, center_account_id),
                )
            n = cursor.fetchone()[0] + 1
            version = f"{_dt.now(_tz).year}-v{n}"
            cursor.execute(
                """
                INSERT INTO formation_modules
                (rncp_code, tp_name, version, status, source_pipeline_job_id,
                 source_platform_id, center_account_id, voice_type, voice_updated_at, validated_at)
                VALUES (?, ?, ?, 'validated', ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                """,
                (rncp, tp_name, version, job_id, platform_id, center_account_id, voice_type),
            )
            module_id = cursor.lastrowid
            module_created = True

        conn.commit()
        result = {
            "platform_id": platform_id,
            "platform_ready_updated": int(platform_ready_updated),
            "module_id": module_id,
            "module_created": module_created,
            "module_version": version,
            "voice_type": voice_type,
        }
        logger.info(
            "PIPELINE_AUDIO_FINALIZED formation_job_id=%s platform_id=%s "
            "module_id=%s module_created=%s voice_type=%s ready_updated=%s",
            job_id, platform_id, module_id, module_created, voice_type, platform_ready_updated,
        )
        return result
    finally:
        conn.close()


def _finalize_text_ready_state(job_id: int) -> dict:
    """Rend la plateforme consultable dès que les textes sont prêts.

    L'audio est désormais lancé séparément. Donc la fin de la pipeline texte
    doit déjà enlever l'overlay "Module en construction" et créer l'enveloppe
    module qui pointe vers les dossiers de cours, sans marquer le module comme
    validé audio.
    """
    from datetime import datetime as _dt
    from config import FRANCE_TZ as _tz
    from database.db import get_db_connection

    job = get_job(job_id)
    if not job:
        raise ValueError(f"Job {job_id} introuvable pour finalisation texte")

    platform_id = job.get("platform_id")
    rncp = job.get("rncp_code") or ""
    tp_name = job.get("tp_name") or f"Job {job_id}"

    conn = get_db_connection()
    cursor = conn.cursor()
    try:
        cursor.execute("SELECT center_account_id FROM platform_config WHERE id = ?", (platform_id,))
        center_row = cursor.fetchone()
        center_account_id = center_row[0] if center_row else None
        cursor.execute(
            "UPDATE platform_config SET status = 'ready' WHERE id = ? AND status = 'pending'",
            (platform_id,),
        )
        platform_ready_updated = cursor.rowcount or 0

        cursor.execute(
            "SELECT id, version, status FROM formation_modules WHERE source_pipeline_job_id = ?",
            (job_id,),
        )
        existing = cursor.fetchone()

        if existing:
            module_id, version, status = existing[0], existing[1], existing[2]
            cursor.execute(
                """UPDATE formation_modules
                   SET source_platform_id = COALESCE(source_platform_id, ?),
                       center_account_id = COALESCE(center_account_id, ?),
                       status = CASE WHEN status = 'validated' THEN status ELSE 'draft' END
                   WHERE id = ?""",
                (platform_id, center_account_id, module_id),
            )
            module_created = False
            module_status = status if status == "validated" else "draft"
        else:
            if center_account_id is None:
                cursor.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id IS NULL",
                    (rncp,),
                )
            else:
                cursor.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id = ?",
                    (rncp, center_account_id),
                )
            n = cursor.fetchone()[0] + 1
            version = f"{_dt.now(_tz).year}-v{n}"
            cursor.execute(
                """
                INSERT INTO formation_modules
                (rncp_code, tp_name, version, status, source_pipeline_job_id,
                 source_platform_id, center_account_id, voice_type, voice_updated_at, validated_at)
                VALUES (?, ?, ?, 'draft', ?, ?, ?, NULL, NULL, NULL)
                """,
                (rncp, tp_name, version, job_id, platform_id, center_account_id),
            )
            module_id = cursor.lastrowid
            module_created = True
            module_status = "draft"

        conn.commit()
        result = {
            "platform_id": platform_id,
            "platform_ready_updated": int(platform_ready_updated),
            "module_id": module_id,
            "module_created": module_created,
            "module_version": version,
            "module_status": module_status,
        }
        logger.info(
            "PIPELINE_TEXT_FINALIZED formation_job_id=%s platform_id=%s "
            "module_id=%s module_created=%s status=%s ready_updated=%s",
            job_id, platform_id, module_id, module_created, module_status, platform_ready_updated,
        )
        return result
    finally:
        conn.close()


def _count_dirty_segments_for_job(job_id: int) -> int:
    from repositories.pipeline_repository import count_dirty_completed_segments_for_folders
    from services.formation_pipeline_service import get_expected_course_folders

    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    if not folder_ids:
        return 0
    return count_dirty_completed_segments_for_folders(folder_ids)


def _folder_text_reviews_ready(job_id: int, folder_id: int) -> tuple[bool, dict]:
    from repositories.pipeline_repository import get_folder_text_review_readiness
    from services.content_generation_service import (
        _current_compliance_review_signature,
    )

    compliance_signature = _current_compliance_review_signature()
    detail = get_folder_text_review_readiness(
        job_id=job_id,
        folder_id=folder_id,
        review_signature=compliance_signature,
    )
    total = int(detail.get("segments_completed") or 0)
    return total > 0 and detail["reviewed_current"] >= total, detail


def start_folder_audio_generation(
    job_id,
    folder_id,
    payload=None,
    *,
    schedule_session_id=None,
    trigger_source="manual",
    stale_started_before=None,
):
    """Lance l'audio d'une seule journée.

    Utilisé par le bouton manuel d'une journée et par le timer 24h avant cours.
    Retourne (payload, http_status) pour rester réutilisable hors route Flask.
    """
    job = get_job(job_id)
    if not job:
        return {"error": "Job introuvable"}, 404

    try:
        folder_id, folder_resolution = _resolve_continue_after_text_folder(job_id, int(folder_id))
    except ValueError as e:
        return {"error": str(e)}, 400

    data = payload or {}
    reviews_ready, review_detail = _folder_text_reviews_ready(job_id, folder_id)
    if not reviews_ready and not bool(data.get("allow_unreviewed")):
        return {
            "error": "Texte pas prêt pour l'audio : la conformité locale par morceau doit être terminée.",
            "review_detail": review_detail,
        }, 400

    tts_mode = (data.get("tts_mode") or job.get("auto_pilot_tts_mode") or "gtts").lower()
    if tts_mode not in ("fish_audio", "gtts", "mock"):
        return {"error": "tts_mode invalide (fish_audio | gtts | mock)"}, 400
    mock = tts_mode == "mock"
    basic_tts = tts_mode == "gtts"
    voice_type = "mock" if mock else ("gtts" if basic_tts else "fish_audio")
    force_all = bool(data.get("force_all", True))
    sync_slides = bool(data.get("sync_slides", True))
    auto_generate_slides = bool(data.get("auto_generate_slides", True))
    preserve_existing = bool(data.get("preserve_existing", False))
    max_slides = int(data.get("max_slides") or 60)
    pace = data.get("pace") or "normal"
    model = _resolve_pipeline_api_model(job, data.get("model"))

    from services.formation_pipeline_service import get_expected_course_folders
    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    if folder_id not in folder_ids:
        return {"error": "Folder hors journées attendues"}, 400
    idx = folder_ids.index(folder_id)
    next_folder_id = folder_ids[idx + 1] if idx + 1 < len(folder_ids) else None

    import eventlet
    from datetime import datetime
    from config import FRANCE_TZ
    from services.content_generation_service import generate_audio_from_script

    if schedule_session_id:
        _conn = None
        try:
            from database.db import get_db_connection as _get_conn
            now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
            stale_retry_clause = ""
            params = [now_str, job_id, folder_id, now_str, schedule_session_id]
            if stale_started_before:
                stale_retry_clause = """
                      OR (
                          COALESCE(audio_generation_status, 'pending') IN ('running', 'processing')
                          AND audio_generation_completed_at IS NULL
                          AND COALESCE(updated_at, audio_generation_started_at) <= ?
                      )
                """
                params.append(stale_started_before)
            _conn = _get_conn()
            _cur = _conn.cursor()
            _cur.execute(
                f"""
                UPDATE course_sessions
                SET audio_generation_status = 'running',
                    audio_generation_started_at = ?,
                    audio_generation_completed_at = NULL,
                    audio_generation_error = NULL,
                    audio_job_id = ?,
                    audio_folder_id = ?,
                    updated_at = ?
                WHERE id = ?
                  AND (
                      audio_generation_started_at IS NULL
                      OR (
                          COALESCE(audio_generation_status, 'pending') = 'error'
                          AND audio_generation_completed_at IS NULL
                      )
                      {stale_retry_clause}
                  )
                """,
                params,
            )
            if _cur.rowcount == 0:
                _conn.close()
                return {"error": "Audio déjà lancé ou terminé pour cette séance"}, 409
            _conn.commit()
            _conn.close()
        except Exception as exc:
            try:
                if _conn:
                    _conn.close()
            except Exception:
                pass
            logger.warning("⚠️ Impossible de marquer la séance audio running", exc_info=True)
            return {"error": f"Impossible de verrouiller la séance audio: {str(exc)[:200]}"}, 500

    def _run_one():
        started_at = time.time()
        try:
            from services.formation_observability_service import log_pipeline_event
            update_job(job_id, status="audio_running", error_message=None)
            log_pipeline_event(
                job_id,
                "audio_folder_started",
                step="audio",
                status="running",
                folder_id=folder_id,
                model=str(model) if model else None,
                message="Synthèse audio journée démarrée",
                data={
                    "voice_type": voice_type,
                    "force_all": force_all,
                    "preserve_existing": preserve_existing,
                    "sync_slides": sync_slides,
                    "auto_generate_slides": auto_generate_slides,
                    "single_folder": True,
                    "trigger_source": trigger_source,
                    "schedule_session_id": schedule_session_id,
                },
            )
            result_audio = generate_audio_from_script(
                folder_id,
                on_progress=_make_audio_progress_logger(
                    job_id,
                    folder_id,
                    voice_type,
                    schedule_session_id=schedule_session_id,
                ),
                force_all=force_all,
                mock=mock,
                basic_tts=basic_tts,
                next_folder_id=next_folder_id,
                is_last_folder=next_folder_id is None,
                sync_slides=sync_slides,
                auto_generate_slides=auto_generate_slides,
                slide_max_slides=max_slides,
                slide_pace=pace,
                slide_model=model,
                llm_model=model,
                preserve_existing=preserve_existing,
            )
            publish_result = None
            try:
                from services.audio_publish_service import publish_playlist_audio_to_platform
                publish_result = publish_playlist_audio_to_platform(
                    job["platform_id"],
                    folder_id,
                    archive_existing=True,
                    archive_reason=f"{trigger_source}-folder-{folder_id}",
                )
            except Exception as publish_error:
                publish_result = {"published": [], "publish_errors": [{"error": str(publish_error)}]}
                logger.error(
                    "❌ Publication audio journée échouée job=%s platform=%s folder=%s: %s",
                    job_id,
                    job.get("platform_id"),
                    folder_id,
                    publish_error,
                    exc_info=True,
                )
            n_dirty = _count_dirty_segments_for_job(job_id)
            finalize_result = None
            # L'audio d'une journée est une action d'exploitation à la demande :
            # elle ne clôture pas la pipeline de fabrication du module.
            update_job(job_id, status="text_ready", error_message=None)
            log_pipeline_event(
                job_id,
                "audio_folder_completed",
                step="audio",
                status="completed",
                folder_id=folder_id,
                model=str(model) if model else None,
                duration_ms=int((time.time() - started_at) * 1000),
                message="Synthèse audio journée terminée",
                data={
                    "voice_type": voice_type,
                    "remaining_dirty_segments": n_dirty,
                    "finalized": False,
                    "finalize_result": finalize_result,
                    "generated": result_audio.get("generated") if isinstance(result_audio, dict) else None,
                    "skipped": result_audio.get("skipped") if isinstance(result_audio, dict) else None,
                    "publish": publish_result,
                    "single_folder": True,
                    "trigger_source": trigger_source,
                    "schedule_session_id": schedule_session_id,
                },
            )
            if schedule_session_id:
                try:
                    from database.db import get_db_connection as _get_conn
                    now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
                    _conn = _get_conn()
                    _cur = _conn.cursor()
                    _cur.execute(
                        """
                        UPDATE course_sessions
                        SET audio_generation_status = 'completed',
                            audio_generation_completed_at = ?,
                            audio_generation_error = NULL,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (now_str, now_str, schedule_session_id),
                    )
                    _conn.commit()
                    _conn.close()
                except Exception:
                    logger.warning("⚠️ Impossible de marquer la séance audio completed", exc_info=True)
        except Exception as e:
            update_job(job_id, status="text_ready", error_message=f"audio folder {folder_id}: {str(e)[:500]}")
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "audio_folder_failed",
                    step="audio",
                    status="error",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    duration_ms=int((time.time() - started_at) * 1000),
                    message="Synthèse audio journée échouée",
                    data={
                        "voice_type": voice_type,
                        "trigger_source": trigger_source,
                        "schedule_session_id": schedule_session_id,
                    },
                    error=str(e)[:500],
                )
            except Exception:
                pass
            if schedule_session_id:
                try:
                    from database.db import get_db_connection as _get_conn
                    now_str = datetime.now(FRANCE_TZ).strftime("%Y-%m-%d %H:%M:%S")
                    _conn = _get_conn()
                    _cur = _conn.cursor()
                    _cur.execute(
                        """
                        UPDATE course_sessions
                        SET audio_generation_status = 'error',
                            audio_generation_error = ?,
                            updated_at = ?
                        WHERE id = ?
                        """,
                        (str(e)[:500], now_str, schedule_session_id),
                    )
                    _conn.commit()
                    _conn.close()
                except Exception:
                    logger.warning("⚠️ Impossible de marquer la séance audio error", exc_info=True)
            logger.error("❌ Audio folder job=%s folder=%s : %s", job_id, folder_id, e, exc_info=True)

    eventlet.spawn(_run_one)
    return {
        "message": "Synthèse audio lancée pour cette journée",
        "job_id": job_id,
        "folder_id": folder_id,
        "resolved_folder_id": folder_id,
        "folder_resolution": folder_resolution,
        "next_folder_id": next_folder_id,
        "tts_mode": tts_mode,
        "voice_type": voice_type,
        "status": "audio_running",
        "schedule_session_id": schedule_session_id,
        "trigger_source": trigger_source,
    }, 202


@formation_bp.route("/api/formation/<int:job_id>/content/<int:folder_id>/generate-audio", methods=["POST"])
def generate_folder_audio(job_id, folder_id):
    """Génère l'audio d'une seule journée/semaine, après pipeline texte complète."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    payload, status = start_folder_audio_generation(
        job_id,
        folder_id,
        request.get_json(silent=True) or {},
        trigger_source="manual",
    )
    return jsonify(payload), status


@formation_bp.route("/api/formation/<int:job_id>/launch-audio", methods=["POST"])
def launch_audio(job_id):
    """
    Lance la synthèse audio Fish Audio S2-Pro pour toutes les journées du job.
    Pré-requis : chaque dossier cours doit avoir son content_generation_job en
    status 'completed' (textes générés par Claude). Boucle sur les dossiers et
    spawn un greenlet eventlet par journée qui appelle generate_audio_from_script.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    from database.db import get_db_connection
    from services.formation_pipeline_service import get_expected_course_folders
    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    if not folder_ids:
        return jsonify({"error": "Aucun cours_folder attendu pour ce job"}), 400
    placeholders = ",".join("?" * len(folder_ids))
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        f"""SELECT f.id FROM cours_folders f
           WHERE f.id IN ({placeholders}) ORDER BY f.position ASC, f.id ASC""",
        tuple(folder_ids),
    )
    folder_ids = [r[0] for r in cursor.fetchall()]

    # Vérifier que tous les dossiers ont leur texte généré
    missing = []
    for fid in folder_ids:
        cursor.execute(
            "SELECT status FROM content_generation_jobs WHERE folder_id = ?",
            (fid,),
        )
        row = cursor.fetchone()
        if not row or row[0] != "completed":
            missing.append(fid)
    conn.close()

    if missing:
        return jsonify({
            "error": f"Textes pas encore prêts pour {len(missing)} dossier(s)",
            "missing_folder_ids": missing,
        }), 400

    not_review_ready = []
    for fid in folder_ids:
        ok, detail = _folder_text_reviews_ready(job_id, fid)
        if not ok:
            not_review_ready.append({"folder_id": fid, **detail})
    if not_review_ready:
        return jsonify({
            "error": "Conformité locale incomplète : audio bloqué tant que les textes ne sont pas validés.",
            "folders": not_review_ready,
        }), 400

    data = request.get_json(silent=True) or {}
    # force_all=True par défaut au lancement initial : les segments fraîchement
    # générés ont dirty=0, donc sans force_all, generate_audio_from_script
    # skipperait tous les blocs ("non modifiés, conservés") et aucun MP3 ne
    # serait produit. force_all=True garantit la 1re synthèse complète. Les
    # régénérations partielles ultérieures (via édition segment) utilisent le
    # dirty flag naturellement.
    force_all = bool(data.get("force_all", True))
    preserve_existing = bool(data.get("preserve_existing", True))
    # 3 modes de synthèse audio (priorité décroissante) :
    # - mock=True      → MP3 silence 1s, test gratuit
    # - basic_tts=True → Edge TTS (voix basique gratuite ; identifiant DB historique "gtts")
    # - (défaut)       → Fish Audio S2-Pro (voix studio payante)
    mock = bool(data.get("mock", False))
    basic_tts = bool(data.get("basic_tts", False))
    sync_slides = bool(data.get("sync_slides", False))
    auto_generate_slides = bool(data.get("auto_generate_slides", False))
    if mock and basic_tts:
        return jsonify({"error": "mock et basic_tts sont mutuellement exclusifs"}), 400

    # Parallélisation inter-folders (1 greenlet par journée en simultané).
    # Réservé à Edge TTS / mock pour éviter le rate limit Fish Audio.
    # Côté backend on hard-cap à 1 si basic_tts=False (donc Fish) — défense
    # en profondeur, le frontend ne doit pas l'envoyer pour les boutons Fish.
    raw_parallel = data.get("parallel_folders")
    try:
        parallel_folders = int(raw_parallel) if raw_parallel is not None else int(os.getenv("AUDIO_PARALLEL_FOLDERS", "1"))
    except (TypeError, ValueError):
        parallel_folders = 1
    parallel_folders = max(1, parallel_folders)
    if not (basic_tts or mock):
        # Hard guard : Fish Audio paye à l'usage et a un rate limit strict.
        # On refuse la parallélisation côté serveur même si le client la
        # demande. Le frontend ne devrait pas l'envoyer pour les boutons
        # payants — c'est un filet de sécurité supplémentaire.
        if parallel_folders > 1:
            logger.warning(
                f"⚠️ parallel_folders={parallel_folders} ignoré : Fish Audio reste séquentiel "
                f"(coût + rate limit)"
            )
        parallel_folders = 1
    parallel_folders = min(parallel_folders, len(folder_ids))

    import eventlet
    from services.content_generation_service import generate_audio_from_script

    def _run_audio(folder_id, next_folder_id=None, is_last_folder=False):
        import sys, traceback
        mode_label = "[MOCK]" if mock else "[BASIC edge-tts]" if basic_tts else ""
        logger.info(f"🚀 SPAWN greenlet job={job_id} folder={folder_id} mock={mock} basic_tts={basic_tts} engine={'edge-tts' if basic_tts else 'fish_audio'} force_all={force_all}")
        sys.stdout.flush()
        started_at = time.time()
        voice_type = "mock" if mock else ("gtts" if basic_tts else "fish_audio")
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "audio_folder_started",
                step="audio",
                status="running",
                folder_id=folder_id,
                message="Synthèse audio journée démarrée",
                model=_resolve_pipeline_api_model(job),
                data={
                    "voice_type": voice_type,
                    "force_all": force_all,
                    "preserve_existing": preserve_existing,
                    "sync_slides": sync_slides,
                    "auto_generate_slides": auto_generate_slides,
                },
            )
        except Exception:
            pass
        try:
            logger.info(f"🎙️ {mode_label} Job {job_id} folder {folder_id} : synthèse audio démarrée")
            sys.stdout.flush()
            generate_audio_from_script(
                folder_id,
                on_progress=_make_audio_progress_logger(job_id, folder_id, voice_type),
                force_all=force_all,
                mock=mock,
                basic_tts=basic_tts,
                next_folder_id=next_folder_id,
                is_last_folder=is_last_folder,
                sync_slides=sync_slides,
                auto_generate_slides=auto_generate_slides,
                llm_model=_resolve_pipeline_api_model(job),
                preserve_existing=preserve_existing,
            )
            duration_ms = int((time.time() - started_at) * 1000)
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "audio_folder_completed",
                    step="audio",
                    status="completed",
                    folder_id=folder_id,
                    duration_ms=duration_ms,
                    message="Synthèse audio journée terminée",
                    model=_resolve_pipeline_api_model(job),
                    data={"voice_type": voice_type},
                )
            except Exception:
                pass
            logger.info(f"✅ {mode_label} Job {job_id} folder {folder_id} : synthèse audio terminée")
            sys.stdout.flush()
            return True
        except Exception as e:
            duration_ms = int((time.time() - started_at) * 1000)
            logger.error(f"❌ Job {job_id} folder {folder_id} : synthèse audio échouée : {e}")
            logger.error(traceback.format_exc())
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "audio_folder_failed",
                    step="audio",
                    status="error",
                    folder_id=folder_id,
                    duration_ms=duration_ms,
                    message="Synthèse audio journée échouée",
                    model=_resolve_pipeline_api_model(job),
                    data={"voice_type": voice_type},
                    error=str(e)[:500],
                )
            except Exception:
                pass
            sys.stdout.flush()
            try:
                update_job(job_id, status="audio_error", error_message=f"folder {folder_id}: {str(e)[:500]}")
            except Exception as ue:
                logger.error(f"❌ Impossible de marquer job {job_id} en audio_error : {ue}")
            return False

    # Synthèse audio : séquentielle par défaut, ou parallèle si parallel_folders>1.
    # - Séquentiel (Fish Audio + défaut) : 1 folder à la fois + cooldown 30s.
    #   Évite le rate limit côté Fish Audio (qui est payant et a un quota strict).
    # - Parallèle (Edge TTS uniquement, sur demande explicite frontend ou env) :
    #   tous les folders en simultané dans un GreenPool. Pas de cooldown.
    #   Pas de carryover inter-jours (next_folder_id=None) parce que Jour 2 démarre
    #   avant que Jour 1 ne sache s'il a du surplus à reporter.
    def _run_all_audios():
        run_started_at = time.time()
        failures = []

        if parallel_folders > 1 and len(folder_ids) > 1:
            import eventlet as _ev
            logger.info(
                f"🚀 PARALLEL audio: {len(folder_ids)} folder(s) sur pool de "
                f"{parallel_folders} (mode={'edge' if basic_tts else 'mock'})"
            )
            pool = _ev.GreenPool(size=parallel_folders)
            pile = _ev.GreenPile(pool)
            for fid in folder_ids:
                # En parallèle, pas de carryover inter-jours (cf. raison ci-dessus).
                pile.spawn(_run_audio, fid, None, True)
            for idx, ok in enumerate(pile):
                if not ok:
                    failures.append(folder_ids[idx])
        else:
            for idx, fid in enumerate(folder_ids):
                next_fid = folder_ids[idx + 1] if idx + 1 < len(folder_ids) else None
                ok = _run_audio(fid, next_folder_id=next_fid, is_last_folder=next_fid is None)
                if not ok:
                    failures.append(fid)
                # Petit cooldown entre folders pour aérer le rate limit côté
                # API tierce. Configurable, surtout utile pour Fish Audio.
                import eventlet as _ev
                cooldown = int(os.getenv("AUDIO_COOLDOWN_BETWEEN_FOLDERS_SEC", "30"))
                if cooldown > 0 and idx + 1 < len(folder_ids):
                    logger.info(f"⏸ Cooldown {cooldown}s avant folder suivant")
                    _ev.sleep(cooldown)
        try:
            from services.formation_observability_service import log_pipeline_event
            if failures:
                log_pipeline_event(
                    job_id,
                    "audio_failed",
                    step="audio",
                    status="error",
                    duration_ms=int((time.time() - run_started_at) * 1000),
                    message=f"Synthèse audio terminée avec {len(failures)} échec(s)",
                    data={"failed_folder_ids": failures, "folder_ids": folder_ids},
                    error=f"{len(failures)} dossier(s) audio en erreur",
                )
            else:
                log_pipeline_event(
                    job_id,
                    "audio_completed",
                    step="audio",
                    status="completed",
                    duration_ms=int((time.time() - run_started_at) * 1000),
                    message=f"Synthèse audio terminée pour {len(folder_ids)} journée(s)",
                    data={"folder_ids": folder_ids},
                )
        except Exception:
            pass
        if failures:
            update_job(
                job_id,
                status="audio_error",
                error_message=f"Synthèse audio incomplète : {len(failures)} dossier(s) en erreur",
            )
        else:
            update_job(job_id, status="audio_completed", error_message=None)

    update_job(job_id, status="audio_running", error_message=None)
    try:
        from services.formation_observability_service import log_pipeline_event
        log_pipeline_event(
            job_id,
            "audio_started",
            step="audio",
            status="running",
            message=f"Synthèse audio lancée pour {len(folder_ids)} journée(s)",
            data={
                "folder_ids": folder_ids,
                "voice_type": "mock" if mock else ("gtts" if basic_tts else "fish_audio"),
                "force_all": force_all,
                "preserve_existing": preserve_existing,
                "sync_slides": sync_slides,
                "auto_generate_slides": auto_generate_slides,
            },
        )
    except Exception:
        pass

    eventlet.spawn(_run_all_audios)

    # Marquer la plateforme comme 'ready' : le contenu est validé, la synthèse
    # audio tourne en background. Côté HR Dashboard, l'overlay "Module en
    # construction" disparaît dès ce moment — le module est exploitable même
    # si les derniers MP3 finissent de s'uploader (surtout en mode mock où
    # c'est instantané).
    try:
        from database.db import get_db_connection as _get_conn
        _c = _get_conn()
        _cur = _c.cursor()
        _cur.execute(
            "UPDATE platform_config SET status = 'ready' WHERE id = ? AND status = 'pending'",
            (job["platform_id"],),
        )
        _c.commit()
        _c.close()
        logger.info(f"✅ Plateforme {job['platform_id']} : status pending → ready")
    except Exception as e:
        logger.warning(f"⚠️ Impossible de marquer la plateforme ready : {e}")

    # Auto-création du module persistant (idempotent via UNIQUE constraint sur
    # source_pipeline_job_id). Principe "1 RNCP = 1 module durable" : ce module
    # devient sélectionnable dans la modale "Nouvelle plateforme" pour créer
    # les futures promos sans re-lancer la pipeline.
    #
    # voice_type trace la voix TTS qui a produit les MP3 du module. À chaque
    # relance de l'étape 7 avec une voix différente, le champ est UPDATE.
    # Les MP3 dans Azure (clé = platform_id/folder_id/filename) sont écrasés
    # par le nouveau run, donc le module pointe automatiquement vers les
    # nouveaux audios — voice_type reflète ce changement de manière persistante.
    voice_type = "mock" if mock else ("gtts" if basic_tts else "fish_audio")
    try:
        from database.db import get_db_connection as _gc
        from datetime import datetime as _dt
        from config import FRANCE_TZ as _tz
        _c2 = _gc()
        _cur2 = _c2.cursor()
        year = _dt.now(_tz).year
        rncp = job.get("rncp_code") or ""
        platform_id = job["platform_id"]
        _cur2.execute("SELECT center_account_id FROM platform_config WHERE id = ?", (platform_id,))
        center_row = _cur2.fetchone()
        center_account_id = center_row[0] if center_row else None
        if center_account_id is None:
            _cur2.execute(
                "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id IS NULL",
                (rncp,),
            )
        else:
            _cur2.execute(
                "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id = ?",
                (rncp, center_account_id),
            )
        n = _cur2.fetchone()[0] + 1
        version = f"{year}-v{n}"
        _cur2.execute("""
            INSERT OR IGNORE INTO formation_modules
            (rncp_code, tp_name, version, status, source_pipeline_job_id,
             source_platform_id, center_account_id, voice_type, voice_updated_at, validated_at)
            VALUES (?, ?, ?, 'validated', ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
        """, (rncp, job["tp_name"], version, job_id, platform_id, center_account_id, voice_type))
        if _cur2.rowcount > 0:
            logger.info(
                f"📦 Module créé : {job['tp_name']} {version} (job {job_id}) "
                f"voix={voice_type}"
            )
        else:
            # Module déjà existant — relance TTS avec une voix possiblement
            # différente. On UPDATE voice_type pour refléter les MP3 actuels.
            _cur2.execute(
                """UPDATE formation_modules
                   SET voice_type = ?,
                       voice_updated_at = CURRENT_TIMESTAMP,
                       source_platform_id = COALESCE(source_platform_id, ?),
                       center_account_id = COALESCE(center_account_id, ?),
                       status = 'validated',
                       validated_at = COALESCE(validated_at, CURRENT_TIMESTAMP)
                   WHERE source_pipeline_job_id = ?""",
                (voice_type, platform_id, center_account_id, job_id),
            )
            logger.info(
                f"♻️ Module mis à jour pour job {job_id} : voix={voice_type} "
                f"(les MP3 Azure sont écrasés en place)"
            )
        _c2.commit()
        _c2.close()
    except Exception as e:
        logger.warning(f"⚠️ Création/MAJ module échouée : {e}")

    if mock:
        mode_suffix = " (MOCK — silence 1s)"
    elif basic_tts:
        mode_suffix = " (Edge TTS — voix basique gratuite)"
    else:
        mode_suffix = ""
    if sync_slides:
        mode_suffix += " + slides synchronisées"
    logger.info(f"🚀 Job {job_id} : synthèse audio lancée pour {len(folder_ids)} dossiers{mode_suffix}")
    return jsonify({
        "message": f"Synthèse audio lancée pour {len(folder_ids)} journées{mode_suffix}",
        "folder_ids": folder_ids,
        "status": "audio_running",
        "mock": mock,
        "basic_tts": basic_tts,
        "sync_slides": sync_slides,
        "auto_generate_slides": auto_generate_slides,
    })


def _reset_folder_downstream_to_generated_text(job_id: int, folder_id: int) -> dict:
    """Conserve le texte initial et remet à zéro volume/review/slides/audio."""
    from database.db import get_db_connection

    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT cf.id, cf.name, cf.position, cj.id, cj.platform_id
        FROM cours_folders cf
        JOIN content_generation_jobs cj ON cj.folder_id = cf.id
        WHERE cf.id = ? AND cf.formation_job_id = ? AND cj.status = 'completed'
        """,
        (folder_id, job_id),
    )
    row = cursor.fetchone()
    if not row:
        conn.close()
        raise ValueError("Journée introuvable ou texte non généré")

    _, folder_name, position, content_job_id, platform_id = row
    cursor.execute(
        """
        SELECT id, COALESCE(text_content, ''), text_content_pre_review
        FROM content_generation_segments
        WHERE job_id = ? AND status = 'completed'
        ORDER BY sub_part_index ASC, passe ASC
        """,
        (content_job_id,),
    )
    segments = cursor.fetchall()
    if not segments:
        conn.close()
        raise ValueError("Aucun segment texte complété pour cette journée")

    restored = 0
    total_words = 0
    for seg_id, current_text, original_text in segments:
        base_text = (original_text if original_text is not None else current_text) or ""
        word_count = len(base_text.split())
        total_words += word_count
        if original_text is not None and original_text != current_text:
            restored += 1
        cursor.execute(
            """
            UPDATE content_generation_segments
            SET text_content = ?, word_count = ?, dirty = 1,
                humanized = 0, humanization_error = NULL, humanization_signature = NULL,
                reviewed = 0, review_error = NULL, review_signature = NULL
            WHERE id = ?
            """,
            (base_text, word_count, seg_id),
        )

    cursor.execute(
        """
        UPDATE content_generation_jobs
        SET total_words = ?, status = 'completed', error_message = NULL,
            updated_at = CURRENT_TIMESTAMP
        WHERE id = ?
        """,
        (total_words, content_job_id),
    )
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='content_review_reports'")
    if cursor.fetchone():
        cursor.execute(
            "DELETE FROM content_review_reports WHERE job_id = ? AND folder_id = ?",
            (job_id, folder_id),
        )
    deleted_review_artifacts = _delete_active_review_artifacts(job_id, position)
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='script_slide_decks'")
    if cursor.fetchone():
        cursor.execute(
            "DELETE FROM script_slide_decks WHERE folder_id = ? AND content_job_id = ?",
            (folder_id, content_job_id),
        )
        deleted_decks = cursor.rowcount
    else:
        deleted_decks = 0

    conn.commit()
    conn.close()

    update_job(
        job_id,
        status="tts_launched",
        auto_pilot_volume_done=0,
        auto_pilot_post_review_docs_done=0,
        auto_pilot_error=None,
    )
    return {
        "folder_id": folder_id,
        "folder_name": folder_name,
        "position": position,
        "content_job_id": content_job_id,
        "platform_id": platform_id,
        "segments": len(segments),
        "segments_restored": restored,
        "total_words": total_words,
        "deleted_review_artifacts": deleted_review_artifacts,
        "deleted_slide_decks": deleted_decks,
    }


def _completed_text_folder_candidates(job_id: int) -> list[dict]:
    """Liste les dossiers rattachés au job qui ont vraiment un texte complet."""
    from repositories.pipeline_repository import list_text_folder_states_for_folders
    from services.formation_pipeline_service import get_expected_course_folders

    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    if not folder_ids:
        return []
    return list_text_folder_states_for_folders(folder_ids, completed_only=True)


def _requested_text_folder_state(job_id: int, folder_id: int) -> dict | None:
    """Retourne l'état texte du folder demandé, même s'il n'est pas completed."""
    from repositories.pipeline_repository import get_text_folder_state

    return get_text_folder_state(folder_id)


def _claim_single_completed_orphan_folder(job_id: int, requested: dict | None) -> dict | None:
    """Rattache un unique folder completed orphelin quand l'ancien lien est cassé."""
    from database.db import get_db_connection

    job = get_job(job_id)
    if not job:
        return None
    conn = get_db_connection()
    cursor = conn.cursor()
    params = [job["platform_id"]]
    day_filter = ""
    if requested and requested.get("folder_name"):
        import re

        match = re.match(r"^\s*Jour\s+(\d+)\b", requested["folder_name"], flags=re.IGNORECASE)
        if match:
            day_filter = "AND cf.name LIKE ?"
            params.append(f"Jour {match.group(1)}%")
    cursor.execute(
        f"""
        SELECT
            cf.id,
            cf.name,
            cf.position,
            cf.platform_id,
            cgj.id,
            cgj.status,
            COALESCE(cgj.total_words, 0),
            SUM(CASE WHEN cgs.status = 'completed' THEN 1 ELSE 0 END)
        FROM cours_folders cf
        JOIN content_generation_jobs cgj ON cgj.folder_id = cf.id
        LEFT JOIN content_generation_segments cgs ON cgs.job_id = cgj.id
        WHERE cf.platform_id = ?
          AND cf.formation_job_id IS NULL
          {day_filter}
        GROUP BY cf.id, cf.name, cf.position, cf.platform_id, cgj.id, cgj.status, cgj.total_words
        HAVING cgj.status = 'completed'
           AND SUM(CASE WHEN cgs.status = 'completed' THEN 1 ELSE 0 END) > 0
        ORDER BY cf.created_at DESC, cf.id DESC
        """,
        tuple(params),
    )
    rows = cursor.fetchall()
    if len(rows) != 1:
        conn.close()
        return None

    row = rows[0]
    folder_id = row[0]
    cursor.execute(
        """
        UPDATE cours_folders
        SET formation_job_id = ?
        WHERE id = ? AND formation_job_id IS NULL
        """,
        (job_id, folder_id),
    )
    conn.commit()
    conn.close()
    logger.warning(
        "PIPELINE_FOLDER_REPAIR job=%s repaired=1 missing=0 folders=%s reason=continue_after_text_orphan",
        job_id,
        [{"folder_id": folder_id, "name": row[1]}],
    )
    return {
        "folder_id": row[0],
        "folder_name": row[1],
        "position": row[2],
        "platform_id": row[3],
        "content_job_id": row[4],
        "content_status": row[5],
        "total_words": row[6] or 0,
        "segments_completed": row[7] or 0,
    }


def _resolve_continue_after_text_folder(job_id: int, requested_folder_id: int) -> tuple[int, dict | None]:
    """Valide le folder de relance aval et corrige les anciens liens cassés."""
    try:
        from services.formation_pipeline_service import repair_orphan_content_folders
        repair_orphan_content_folders(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Réparation folders avant relance aval job {job_id} ignorée : {e}")

    requested = _requested_text_folder_state(job_id, requested_folder_id)
    if (
        requested
        and requested.get("formation_job_id") == job_id
        and requested.get("content_status") == "completed"
        and (requested.get("segments_completed") or 0) > 0
    ):
        return requested_folder_id, None

    candidates = _completed_text_folder_candidates(job_id)
    if requested and requested.get("position") is not None:
        same_position = [c for c in candidates if c.get("position") == requested.get("position")]
        if len(same_position) == 1:
            resolved = same_position[0]
            return int(resolved["folder_id"]), {
                "requested_folder_id": requested_folder_id,
                "resolved_folder_id": resolved["folder_id"],
                "reason": "same_position_completed_folder",
                "requested_state": requested,
            }
    if len(candidates) == 1:
        resolved = candidates[0]
        return int(resolved["folder_id"]), {
            "requested_folder_id": requested_folder_id,
            "resolved_folder_id": resolved["folder_id"],
            "reason": "single_completed_folder_for_job",
            "requested_state": requested,
        }

    orphan = _claim_single_completed_orphan_folder(job_id, requested)
    if orphan:
        return int(orphan["folder_id"]), {
            "requested_folder_id": requested_folder_id,
            "resolved_folder_id": orphan["folder_id"],
            "reason": "single_completed_orphan_folder_claimed",
            "requested_state": requested,
        }

    available = [
        f"folder {c['folder_id']} ({c['segments_completed']} segments, {c['total_words']} mots)"
        for c in candidates
    ]
    requested_detail = (
        "introuvable"
        if not requested
        else (
            f"status={requested.get('content_status') or 'aucun content_job'}, "
            f"segments={requested.get('segments_completed') or 0}, "
            f"formation_job_id={requested.get('formation_job_id')}"
        )
    )
    raise ValueError(
        f"Folder {requested_folder_id} non prêt pour relance aval ({requested_detail}). "
        + (
            "Folder(s) texte completed disponible(s) : " + "; ".join(available)
            if available
            else "Aucun folder texte completed disponible pour ce job."
        )
    )


def _next_folder_in_formation(job_id: int, folder_id: int) -> int | None:
    from services.formation_pipeline_service import get_expected_course_folders

    folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
    try:
        idx = folder_ids.index(folder_id)
    except ValueError:
        return None
    return folder_ids[idx + 1] if idx + 1 < len(folder_ids) else None


def _get_folder_info_for_resume(job_id: int, folder_id: int) -> dict:
    """Lit platform_id / content_job_id sans modifier l'état."""
    from database.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute(
        """
        SELECT cf.id, cf.name, cf.position, cj.id, cj.platform_id
        FROM cours_folders cf
        JOIN content_generation_jobs cj ON cj.folder_id = cf.id
        WHERE cf.id = ? AND cf.formation_job_id = ? AND cj.status = 'completed'
        """,
        (folder_id, job_id),
    )
    row = cursor.fetchone()
    conn.close()
    if not row:
        raise ValueError("Journée introuvable ou texte non généré")
    _, folder_name, position, content_job_id, platform_id = row
    return {
        "folder_id": folder_id,
        "folder_name": folder_name,
        "position": position,
        "content_job_id": content_job_id,
        "platform_id": platform_id,
    }


def _delete_slide_deck_for_resume(folder_id: int, content_job_id: int) -> int:
    """Supprime le deck slides existant pour forcer la régénération à l'étape suivante."""
    from database.db import get_db_connection
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name='script_slide_decks'")
    if not cursor.fetchone():
        conn.close()
        return 0
    cursor.execute(
        "DELETE FROM script_slide_decks WHERE folder_id = ? AND content_job_id = ?",
        (folder_id, content_job_id),
    )
    deleted = cursor.rowcount
    conn.commit()
    conn.close()
    return deleted


@formation_bp.route(
    "/api/formation/<int:job_id>/content/<int:folder_id>/continue-after-text",
    methods=["POST"],
)
def continue_after_text(job_id, folder_id):
    """Relance les étapes aval d'une journée sans régénérer le texte initial.

    Flux complet : reset aval → conformité locale API → Word 2 → slides → Edge TTS sync.
    Paramètre from_step : 'review' (défaut), 'slides', 'tts' — saute les étapes en amont.
    Compatibilité : l'ancien from_step='volume' est mappé vers 'review'.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    data = request.get_json(silent=True) or {}
    model = _resolve_pipeline_api_model(job, data.get("model"))
    max_slides = int(data.get("max_slides") or 60)
    pace = data.get("pace") or "normal"
    requested_folder_id = int(folder_id)
    _STEP_ORDER = ["review", "slides", "tts"]
    raw_from_step = data.get("from_step", "review")
    fast_iteration = bool(
        data.get("iteration_mode") == "fast"
        or data.get("fast_iteration")
        or data.get("skip_audio")
    )
    skip_audio = bool(data.get("skip_audio") or fast_iteration)
    skip_post_content_review = bool(data.get("skip_post_content_review") or fast_iteration)
    fast_tts_pipeline = bool(raw_from_step == "tts_fast" or data.get("fast_tts_pipeline"))
    restart_from_content = (raw_from_step == "content")
    from_step = "tts" if raw_from_step == "tts_fast" else raw_from_step
    if from_step == "volume":
        from_step = "review"
    from_step = from_step if from_step in _STEP_ORDER else "review"
    if restart_from_content and skip_post_content_review:
        from_step = "slides"
    if from_step != "tts":
        fast_tts_pipeline = False
    from_step_idx = _STEP_ORDER.index(from_step)

    logger.info(
        "PIPELINE_RESUME_REQUEST formation_job_id=%s requested_folder_id=%s from_step=%s "
        "raw_from_step=%s model=%s max_slides=%s pace=%s fast_tts_pipeline=%s "
        "fast_iteration=%s skip_audio=%s skip_post_content_review=%s prev_status=%s",
        job_id, requested_folder_id, from_step, raw_from_step, model, max_slides, pace,
        fast_tts_pipeline, fast_iteration, skip_audio, skip_post_content_review, job.get("status"),
    )

    # Persiste le choix de modèle utilisé pour cette relance, pour que les
    # futurs `continue_after_text` ou redémarrages auto-pilot retrouvent le
    # bon provider sans qu'il faille le repasser dans le payload.
    if model and model != job.get("auto_pilot_model"):
        try:
            update_job(job_id, auto_pilot_model=model)
            job["auto_pilot_model"] = model
            logger.info("PIPELINE_RESUME_MODEL_PERSISTED formation_job_id=%s model=%s", job_id, model)
        except Exception as e:
            logger.warning(f"⚠️ Persistance auto_pilot_model={model} échouée pour job {job_id}: {e}")

    try:
        folder_id, folder_resolution = _resolve_continue_after_text_folder(job_id, requested_folder_id)
    except ValueError as e:
        logger.warning(
            "PIPELINE_RESUME_FOLDER_INVALID formation_job_id=%s requested_folder_id=%s error=%s",
            job_id, requested_folder_id, str(e)[:500],
        )
        return jsonify({
            "error": str(e),
            "requested_folder_id": requested_folder_id,
        }), 400

    if folder_id != requested_folder_id:
        logger.warning(
            "PIPELINE_RESUME_FOLDER_RESOLVED formation_job_id=%s requested_folder_id=%s "
            "resolved_folder_id=%s reason=%s",
            job_id, requested_folder_id, folder_id,
            (folder_resolution or {}).get("reason"),
        )

    from services.claude_code_mission_service import _EXECUTION_STATE
    state_key = (job_id, f"continue_after_text_{folder_id}")
    prev_state = _EXECUTION_STATE.get(state_key, {})
    if prev_state.get("status") == "running":
        # Stale lock : si l'état mémoire dit "running" mais que rien ne tourne
        # vraiment (typique après un crash + redémarrage gunicorn), on libère.
        # Heuristique : on accepte de re-prendre la main si pas de heartbeat
        # récent (le greenlet aurait mis à jour _EXECUTION_STATE en cas de
        # vraie activité). En l'absence de timestamp dans _EXECUTION_STATE,
        # on log et on libère systématiquement — l'utilisateur a explicitement
        # demandé une nouvelle relance, c'est un signal fort.
        logger.warning(
            "PIPELINE_RESUME_STALE_LOCK_RELEASED formation_job_id=%s folder_id=%s "
            "prev_state=%s — l'utilisateur force une nouvelle relance",
            job_id, folder_id, {k: prev_state.get(k) for k in ("status", "model", "error")},
        )
        _EXECUTION_STATE.pop(state_key, None)

    update_job(
        job_id,
        auto_pilot_enabled=0,
        auto_pilot_step="manual_continue_after_text",
        auto_pilot_error=None,
        auto_pilot_locked_at=None,
        auto_pilot_lock_owner=None,
    )

    # Reset agressif du status job : si on était en `audio_running` /
    # `audio_error` d'un run précédent (qui peut avoir crashé), on le repose
    # à `tts_launched` pour signaler "TTS à refaire". Le _run le repassera à
    # `audio_running` au moment d'appeler generate_audio_from_script.
    prev_status = job.get("status")
    if prev_status in ("audio_running", "audio_error", "audio_completed"):
        update_job(job_id, status="tts_launched", error_message=None)
        logger.info(
            "PIPELINE_RESUME_STATUS_RESET formation_job_id=%s folder_id=%s "
            "prev_status=%s new_status=tts_launched",
            job_id, folder_id, prev_status,
        )

    try:
        from services.formation_observability_service import clear_pipeline_events
        cleared_events_count = clear_pipeline_events(job_id)
        logger.info(
            "PIPELINE_RESUME_EVENTS_CLEARED formation_job_id=%s folder_id=%s cleared=%s",
            job_id, folder_id, cleared_events_count,
        )
    except Exception as e:
        logger.error(f"❌ Nettoyage événements relance aval impossible job={job_id}: {e}")
        return jsonify({
            "error": "Impossible de nettoyer le journal de pipeline avant relance",
            "details": str(e)[:500],
        }), 500

    _EXECUTION_STATE[state_key] = {"status": "running", "model": str(model), "folder_id": folder_id}
    logger.info(
        "PIPELINE_RESUME_SPAWN formation_job_id=%s folder_id=%s from_step=%s model=%s fast_tts_pipeline=%s",
        job_id, folder_id, from_step, model, fast_tts_pipeline,
    )

    import eventlet

    def _run():
        started_at = time.time()
        try:
            from services.content_generation_service import (
                _assemble_and_upload,
                _update_job_db,
                assert_course_day_word_budget,
                generate_audio_from_script,
                run_content_generation,
                run_content_review,
            )
            from services.formation_observability_service import log_pipeline_event

            logger.info(
                "PIPELINE_RESUME_RUN_START formation_job_id=%s folder_id=%s from_step=%s "
                "from_step_idx=%s model=%s fast_tts_pipeline=%s greenlet_id=%s",
                job_id, folder_id, from_step, from_step_idx, model, fast_tts_pipeline, id(eventlet.getcurrent()),
            )

            log_pipeline_event(
                job_id,
                "continue_after_text_started",
                step="post_text_retry",
                status="running",
                folder_id=folder_id,
                model=str(model) if model else None,
                message=(
                    "Relance aval depuis l'étape 'tts' "
                    "(pipeline presque instantanée)"
                    if fast_tts_pipeline
                    else f"Relance aval depuis l'étape '{from_step}'"
                ),
                data={
                    "voice_type": "gtts",
                    "tts_engine": "edge-tts",
                    "sync_slides": True,
                    "max_slides": max_slides,
                    "pace": pace,
                    "cleared_previous_events": cleared_events_count,
                    "requested_folder_id": requested_folder_id,
                    "resolved_folder_id": folder_id,
                    "folder_resolution": folder_resolution,
                    "from_step": from_step,
                    "raw_from_step": raw_from_step,
                    "fast_tts_pipeline": fast_tts_pipeline,
                    "fast_iteration": fast_iteration,
                    "skip_audio": skip_audio,
                    "skip_post_content_review": skip_post_content_review,
                },
            )

            # ── Étape 0 : GÉNÉRATION TEXTE (si restart_from_content) ────────
            if restart_from_content:
                step_started = time.time()
                logger.info(
                    "PIPELINE_RESUME_STEP_CONTENT_START formation_job_id=%s folder_id=%s",
                    job_id, folder_id,
                )
                from database.db import get_db_connection
                conn = get_db_connection()
                cj = conn.execute(
                    "SELECT id FROM content_generation_jobs WHERE folder_id = ?", (folder_id,)
                ).fetchone()
                conn.close()
                if cj:
                    conn = get_db_connection()
                    conn.execute(
                        "DELETE FROM content_generation_segments WHERE job_id = ?", (cj[0],)
                    )
                    conn.execute(
                        "UPDATE content_generation_jobs SET status='pending', total_words=0, "
                        "current_sub_part=0, current_passe=1 WHERE id = ?",
                        (cj[0],),
                    )
                    conn.commit()
                    conn.close()
                run_content_generation(folder_id, model=model)
                logger.info(
                    "PIPELINE_RESUME_STEP_CONTENT_DONE formation_job_id=%s folder_id=%s duration_ms=%s",
                    job_id, folder_id, int((time.time() - step_started) * 1000),
                )

            # ── Étape 1 : RESET avant reprise complète des reviews ─────────
            if from_step_idx == 0:
                step_started = time.time()
                logger.info(
                    "PIPELINE_RESUME_STEP_RESET_START formation_job_id=%s folder_id=%s",
                    job_id, folder_id,
                )
                reset_info = _reset_folder_downstream_to_generated_text(job_id, folder_id)
                logger.info(
                    "PIPELINE_RESUME_STEP_RESET_DONE formation_job_id=%s folder_id=%s "
                    "content_job_id=%s platform_id=%s segments_restored=%s "
                    "deleted_review_artifacts=%s deleted_slide_decks=%s duration_ms=%s",
                    job_id, folder_id,
                    reset_info.get("content_job_id"),
                    reset_info.get("platform_id"),
                    reset_info.get("segments_restored"),
                    reset_info.get("deleted_review_artifacts"),
                    reset_info.get("deleted_slide_decks"),
                    int((time.time() - step_started) * 1000),
                )
                log_pipeline_event(
                    job_id,
                    "continue_after_text_reset",
                    step="post_text_retry",
                    status="completed",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    message="Étapes aval remises à zéro",
                    data=reset_info,
                )
            else:
                logger.info(
                    "PIPELINE_RESUME_STEP_RESET_SKIP formation_job_id=%s folder_id=%s "
                    "from_step=%s — pas de reset (l'utilisateur saute les reviews)",
                    job_id, folder_id, from_step,
                )
                reset_info = _get_folder_info_for_resume(job_id, folder_id)
                logger.info(
                    "PIPELINE_RESUME_FOLDER_INFO formation_job_id=%s folder_id=%s "
                    "content_job_id=%s platform_id=%s",
                    job_id, folder_id,
                    reset_info.get("content_job_id"),
                    reset_info.get("platform_id"),
                )

            # Sécurité volume append-only supprimée : le volume se corrige
            # uniquement dans le calibrage budget texte de la génération structurée.

            # ── Étape 3 : CONFORMITÉ LOCALE + Word 2
            # L'adhérence au plan est corrigée pendant la génération structurée,
            # juste après les sections et avant le calibrage budget.
            if from_step_idx <= 0:
                step_started = time.time()
                logger.info(
                    "PIPELINE_RESUME_STEP_REVIEW_START formation_job_id=%s folder_id=%s model=%s",
                    job_id, folder_id, model,
                )
                review_result = run_content_review(folder_id, model=model)
                logger.info(
                    "PIPELINE_RESUME_STEP_REVIEW_API_DONE formation_job_id=%s folder_id=%s "
                    "segments_reviewed=%s segments_failed=%s patches_applied=%s "
                    "patches_proposed=%s patches_rejected=%s",
                    job_id, folder_id,
                    review_result.get("segments_reviewed", 0),
                    review_result.get("segments_failed", 0),
                    review_result.get("patches_applied", 0),
                    review_result.get("patches_proposed", 0),
                    review_result.get("patches_rejected", 0),
                )
                _write_api_review_report(job_id, folder_id, review_result, model)
                if review_result.get("segments_failed", 0) > 0:
                    raise RuntimeError(
                        f"Révision conformité échouée sur {review_result['segments_failed']} segment(s)"
                    )
                logger.info(
                    "PIPELINE_RESUME_STEP_REVIEW_REPORT_WRITTEN formation_job_id=%s folder_id=%s",
                    job_id, folder_id,
                )

                budget_audit = assert_course_day_word_budget(
                    folder_id,
                    context="continue_after_text_before_word2",
                )
                log_pipeline_event(
                    job_id,
                    "continue_after_text_word_budget_verified",
                    step="word_budget_review",
                    status="completed",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    message="Budget mots journée vérifié avant Word 2",
                    data={
                        "spoken_words": budget_audit.get("spoken_words"),
                        "raw_words": budget_audit.get("raw_words"),
                        "deficit": budget_audit.get("deficit"),
                        "overflow": budget_audit.get("overflow"),
                        "budget": {
                            "target_words": budget_audit.get("budget", {}).get("target_words"),
                            "min_words": budget_audit.get("budget", {}).get("min_words"),
                            "max_words": budget_audit.get("budget", {}).get("max_words"),
                            "words_per_minute": budget_audit.get("budget", {}).get("words_per_minute"),
                            "course_seconds": budget_audit.get("budget", {}).get("course_seconds"),
                            "speakable_seconds": budget_audit.get("budget", {}).get("speakable_seconds"),
                        },
                    },
                )

                final_words, filename = _assemble_and_upload(
                    folder_id,
                    reset_info["platform_id"],
                    reset_info["content_job_id"],
                )
                logger.info(
                    "PIPELINE_RESUME_STEP_REVIEW_WORD2_BUILT formation_job_id=%s folder_id=%s "
                    "content_job_id=%s total_words=%s filename=%s duration_ms=%s",
                    job_id, folder_id,
                    reset_info["content_job_id"], final_words, filename,
                    int((time.time() - step_started) * 1000),
                )
                _update_job_db(reset_info["content_job_id"], total_words=final_words)
                log_pipeline_event(
                    job_id,
                    "continue_after_text_review_completed",
                    step="review",
                    status="completed",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    message="Conformité locale et Word 2 générés",
                    data={
                        "segments_reviewed": review_result.get("segments_reviewed", 0),
                        "segments_failed": review_result.get("segments_failed", 0),
                        "segments_already_current": review_result.get("segments_already_current", 0),
                        "patches_proposed": review_result.get("patches_proposed", 0),
                        "patches_applied": review_result.get("patches_applied", 0),
                        "patches_rejected": review_result.get("patches_rejected", 0),
                        "review_signature": review_result.get("review_signature"),
                        "doc_filename": filename,
                        "total_words": final_words,
                    },
                )
            else:
                logger.info(
                    "PIPELINE_RESUME_STEP_REVIEW_SKIP formation_job_id=%s folder_id=%s from_step=%s",
                    job_id, folder_id, from_step,
                )

            # ── Étape 4 : SLIDES (force régénération si from_step <= slides) ─
            # Slides et TTS sont jointes par la persistance du deck en DB.
            # Si on vient de "slides" ou avant, on supprime le deck pour forcer
            # une régénération propre. Si on vient de "tts", on conserve les
            # slides existantes et on relance uniquement le TTS sync dessus.
            if from_step_idx <= 1:
                deleted_decks = _delete_slide_deck_for_resume(folder_id, reset_info["content_job_id"])
                logger.info(
                    "PIPELINE_RESUME_STEP_SLIDES_RESET formation_job_id=%s folder_id=%s "
                    "content_job_id=%s deleted_decks=%s — slides seront régénérées avant TTS",
                    job_id, folder_id, reset_info["content_job_id"], deleted_decks,
                )
                log_pipeline_event(
                    job_id,
                    "continue_after_text_slides_reset",
                    step="slides",
                    status="completed",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    message="Deck slides supprimé — régénération automatique avant TTS",
                    data={"deleted_decks": deleted_decks},
                )
            else:
                logger.info(
                    "PIPELINE_RESUME_STEP_SLIDES_SKIP formation_job_id=%s folder_id=%s "
                    "from_step=%s — slides existantes conservées",
                    job_id, folder_id, from_step,
                )

            if skip_audio:
                from services.script_slide_generation_service import generate_slides_from_script

                slide_result = None
                if from_step_idx <= 1:
                    slides_started = time.time()
                    logger.info(
                        "PIPELINE_RESUME_STEP_SLIDES_GENERATE_START formation_job_id=%s folder_id=%s "
                        "content_job_id=%s max_slides=%s pace=%s fast_iteration=%s",
                        job_id, folder_id, reset_info["content_job_id"], max_slides, pace, fast_iteration,
                    )
                    log_pipeline_event(
                        job_id,
                        "slides_folder_started",
                        step="slides",
                        status="running",
                        folder_id=folder_id,
                        model=str(model) if model else None,
                        message="Génération slides anchor-first démarrée",
                        data={
                            "content_job_id": reset_info["content_job_id"],
                            "max_slides": max_slides,
                            "pace": pace,
                            "fast_iteration": fast_iteration,
                            "skip_audio": True,
                        },
                    )
                    slide_result = generate_slides_from_script(
                        folder_id=folder_id,
                        job_id=job_id,
                        platform_id=reset_info["platform_id"],
                        max_slides=max_slides,
                        pace=pace,
                        model=model,
                    )
                    logger.info(
                        "PIPELINE_RESUME_STEP_SLIDES_GENERATE_DONE formation_job_id=%s folder_id=%s "
                        "duration_ms=%s slides=%s source_alignment=%s",
                        job_id,
                        folder_id,
                        int((time.time() - slides_started) * 1000),
                        (slide_result.get("stats") or {}).get("slides_generated"),
                        (slide_result.get("stats") or {}).get("source_alignment"),
                    )
                    log_pipeline_event(
                        job_id,
                        "slides_folder_completed",
                        step="slides",
                        status="completed",
                        folder_id=folder_id,
                        model=str(model) if model else None,
                        duration_ms=int((time.time() - slides_started) * 1000),
                        message="Génération slides anchor-first terminée",
                        data={
                            "content_job_id": reset_info["content_job_id"],
                            "deck_id": (slide_result.get("stats") or {}).get("deck_id"),
                            "slides_generated": (slide_result.get("stats") or {}).get("slides_generated"),
                            "slide_anchors_found": (slide_result.get("stats") or {}).get("slide_anchors_found"),
                            "source_alignment": (slide_result.get("stats") or {}).get("source_alignment"),
                            "beat_aligned_segments": (slide_result.get("stats") or {}).get("beat_aligned_segments"),
                            "beat_aligned_anchors": (slide_result.get("stats") or {}).get("beat_aligned_anchors"),
                            "fast_iteration": fast_iteration,
                            "skip_audio": True,
                        },
                    )

                _finalize_text_ready_state(job_id)
                update_job(job_id, status="text_ready", error_message=None)
                logger.info(
                    "PIPELINE_RESUME_RUN_DONE_NO_AUDIO formation_job_id=%s folder_id=%s from_step=%s "
                    "total_duration_ms=%s fast_iteration=%s",
                    job_id, folder_id, from_step, int((time.time() - started_at) * 1000), fast_iteration,
                )
                log_pipeline_event(
                    job_id,
                    "continue_after_text_completed",
                    step="slides",
                    status="completed",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    duration_ms=int((time.time() - started_at) * 1000),
                    message=(
                        "Itération rapide terminée : texte beat-first et slides régénérés, audio non lancé"
                        if fast_iteration
                        else "Relance terminée sans audio"
                    ),
                    data={
                        "skip_audio": True,
                        "fast_iteration": fast_iteration,
                        "from_step": from_step,
                        "raw_from_step": raw_from_step,
                        "slides": (slide_result or {}).get("stats") or {},
                    },
                )
                _EXECUTION_STATE[state_key] = {
                    "status": "done",
                    "model": str(model),
                    "folder_id": folder_id,
                    "result": {
                        "skip_audio": True,
                        "fast_iteration": fast_iteration,
                        "slides": (slide_result or {}).get("stats") or {},
                    },
                }
                return

            # ── Étape 5 : TTS (toujours exécutée) ───────────────────────────
            tts_started = time.time()
            next_folder_id = _next_folder_in_formation(job_id, folder_id)
            logger.info(
                "PIPELINE_RESUME_STEP_TTS_START formation_job_id=%s folder_id=%s "
                "next_folder_id=%s is_last_folder=%s max_slides=%s pace=%s fast_tts_pipeline=%s",
                job_id, folder_id, next_folder_id, next_folder_id is None, max_slides, pace, fast_tts_pipeline,
            )
            update_job(job_id, status="audio_running", error_message=None)
            audio_result = generate_audio_from_script(
                folder_id,
                on_progress=_make_audio_progress_logger(job_id, folder_id, "gtts"),
                force_all=True,
                mock=False,
                basic_tts=True,
                next_folder_id=next_folder_id,
                is_last_folder=next_folder_id is None,
                sync_slides=True,
                auto_generate_slides=True,
                slide_max_slides=max_slides,
                slide_pace=pace,
                slide_model=model,
                llm_model=model,
                fast_tts_pipeline=fast_tts_pipeline,
            )
            logger.info(
                "PIPELINE_RESUME_STEP_TTS_DONE formation_job_id=%s folder_id=%s duration_ms=%s",
                job_id, folder_id, int((time.time() - tts_started) * 1000),
            )
            finalize_result = _finalize_audio_ready_state(job_id, "gtts")
            update_job(job_id, status="audio_completed", error_message=None)
            logger.info(
                "PIPELINE_RESUME_RUN_DONE formation_job_id=%s folder_id=%s from_step=%s "
                "total_duration_ms=%s",
                job_id, folder_id, from_step, int((time.time() - started_at) * 1000),
            )
            log_pipeline_event(
                job_id,
                "continue_after_text_completed",
                step="audio",
                status="completed",
                folder_id=folder_id,
                model=str(model) if model else None,
                duration_ms=int((time.time() - started_at) * 1000),
                message=(
                    "Relance aval terminée avec Edge TTS rapide test et slides synchronisées"
                    if fast_tts_pipeline
                    else "Relance aval terminée avec Edge TTS et slides synchronisées"
                ),
                data={
                    **(audio_result or {}),
                    "finalize": finalize_result,
                    "tts_engine": "edge-tts",
                    "fast_tts_pipeline": fast_tts_pipeline,
                },
            )
            _EXECUTION_STATE[state_key] = {
                "status": "done",
                "model": str(model),
                "folder_id": folder_id,
                "result": {
                    **(audio_result or {}),
                    "finalize": finalize_result,
                    "tts_engine": "edge-tts",
                    "fast_tts_pipeline": fast_tts_pipeline,
                },
            }
        except Exception as e:
            err = str(e)[:500]
            logger.exception(
                "PIPELINE_RESUME_RUN_FAILED formation_job_id=%s folder_id=%s from_step=%s "
                "duration_ms=%s error=%s",
                job_id, folder_id, from_step,
                int((time.time() - started_at) * 1000), err,
            )
            try:
                update_job(job_id, status="audio_error", error_message=f"folder {folder_id}: {err}")
            except Exception:
                pass
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "continue_after_text_failed",
                    step="post_text_retry",
                    status="error",
                    folder_id=folder_id,
                    model=str(model) if model else None,
                    duration_ms=int((time.time() - started_at) * 1000),
                    message="Relance aval échouée",
                    error=err,
                )
            except Exception:
                pass
            _EXECUTION_STATE[state_key] = {
                "status": "error",
                "model": str(model),
                "folder_id": folder_id,
                "error": err,
            }

    eventlet.spawn(_run)
    return jsonify({
        "ok": True,
        "status": "running",
        "folder_id": folder_id,
        "requested_folder_id": requested_folder_id,
        "folder_resolution": folder_resolution,
        "model": str(model),
        "tts_mode": "gtts",
        "tts_engine": "edge-tts",
        "sync_slides": True,
        "auto_generate_slides": True,
        "fast_tts_pipeline": fast_tts_pipeline,
    }), 202


# ─── Liste des jobs ───────────────────────────────────────────────────────────

@formation_bp.route("/api/formation/list", methods=["GET"])
def list_formations():
    """Liste tous les jobs formation (toutes plateformes)."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    jobs = list_jobs()
    return jsonify({"jobs": jobs})


# ─── Auto-pilot pipeline — state machine persistée en DB ─────────────────────
#
# Architecture : chaque tick exécute UNE seule étape, écrit l'état en DB,
# puis se respawn pour la suivante. Résiste aux restarts Azure App Service.
#
# DB columns : auto_pilot_enabled, auto_pilot_step, auto_pilot_model,
#              auto_pilot_tts_mode, auto_pilot_use_cc, auto_pilot_skip_vs,
#              auto_pilot_generate_audio,
#              auto_pilot_volume_done, auto_pilot_post_review_docs_done,
#              auto_pilot_error,
#              auto_pilot_locked_at, auto_pilot_lock_owner
#
# Lock TTL = 5 min. Au boot, _resume_interrupted_auto_pilots() reprend les
# jobs dont le lock est absent ou périmé.

_AP_LOCK_TTL = 300  # secondes — lock périmé si > 5 min sans mise à jour
_AP_WATCHDOG_INTERVAL = int(os.getenv("AUTO_PILOT_WATCHDOG_INTERVAL_SEC", "60"))
_AP_WATCHDOG_STARTED = False


def _acquire_ap_lock(job_id: int) -> bool:
    """Pose un lock optimiste sur l'auto-pilot. Retourne True si acquis."""
    from repositories.pipeline_repository import acquire_auto_pilot_lock
    owner = str(os.getpid())
    return acquire_auto_pilot_lock(job_id, owner=owner, ttl_seconds=_AP_LOCK_TTL)


def _release_ap_lock(job_id: int) -> None:
    from repositories.pipeline_repository import release_auto_pilot_lock
    release_auto_pilot_lock(job_id)


def _refresh_ap_lock(job_id: int) -> None:
    """Rafraîchit le timestamp du lock pour éviter l'expiration pendant une étape longue."""
    from repositories.pipeline_repository import refresh_auto_pilot_lock
    owner = str(os.getpid())
    refresh_auto_pilot_lock(job_id, owner=owner)


def _ap_lock_age_seconds(job: dict | None) -> float | None:
    if not job or not job.get("auto_pilot_locked_at"):
        return None
    try:
        import datetime
        locked_at = job.get("auto_pilot_locked_at")
        if isinstance(locked_at, str):
            dt = datetime.datetime.fromisoformat(locked_at)
        else:
            dt = locked_at
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=datetime.timezone.utc)
        return (datetime.datetime.now(datetime.timezone.utc) - dt).total_seconds()
    except Exception:
        return None


def _determine_next_ap_step(job_id: int) -> str | None:
    """Détermine la prochaine étape à exécuter (checks idempotents). None = terminé."""
    import json as _json
    from database.db import get_db_connection
    j = get_job(job_id)
    if not j:
        return None

    # Reset du statut erreur avant de relancer (même logique que l'ancien auto-pilot)
    if j["status"] in ("error", "audio_error"):
        if j.get("daily_programs_validated"):
            fallback = "daily_validated"
        elif j.get("global_program_validated"):
            fallback = "global_validated"
        elif j.get("global_program"):
            fallback = "global_ready"
        elif j.get("reac_text"):
            fallback = "reac_ready"
        else:
            fallback = "init"
        update_job(job_id, status=fallback, error_message=None)

    # 1. REAC
    if not j.get("reac_text"):
        return "reac"

    # 2. KB
    kb_done = j.get("status") in (
        "global_ready", "global_validated", "daily_ready",
        "daily_validated", "text_ready", "tts_launched", "audio_running",
        "audio_completed", "audio_launched",
    )
    if not kb_done:
        try:
            from services.knowledge_base_service import kb_stats as _kb_stats
            kb_done = _kb_stats(job_id).get("completed", 0) > 0
        except Exception:
            pass
    if not kb_done:
        return "kb"

    # 3. Programme global
    if not j.get("global_program"):
        return "global"

    # 4. Programmes journée
    daily = j.get("daily_programs") or "[]"
    if not daily or daily in ("[]", '"[]"'):
        return "daily"

    # 5. Génération contenu.
    # Ancien garde-fou : comparer les segments terminés à `sub_parts × 3`.
    # Ce n'est plus fiable depuis la génération structurée : un dossier peut être
    # complet avec un nombre de segments différent de l'ancien pipeline 3 passes.
    # La source de vérité devient donc le job contenu par dossier :
    # tous les dossiers attendus doivent exister, avoir un content job `completed`
    # et contenir au moins un segment finalisé.
    from services.formation_pipeline_service import get_expected_course_folders
    folder_state = get_expected_course_folders(job_id)
    folder_ids = folder_state.get("folder_ids") or []
    expected_folder_count = int(j.get("nb_days") or 0)
    if len(folder_ids) < expected_folder_count:
        return "content"
    from repositories.pipeline_repository import (
        count_dirty_completed_segments_for_folders,
        count_segments_pending_review_for_folders,
        get_latest_script_slide_deck_row,
        list_completed_content_jobs_for_folders,
        list_content_completion_rows_for_folders,
    )
    content_rows = list_content_completion_rows_for_folders(folder_ids)
    completed_folder_ids = {
        row["folder_id"]
        for row in content_rows
        if row.get("status") == "completed"
        and (int(row.get("total_words") or 0) > 0 or int(row.get("completed_segments") or 0) > 0)
    }
    if len(completed_folder_ids) < len(folder_ids):
        return "content"

    # Volume safety append-only retirée du flux auto-pilot : elle ajoutait parfois
    # du développement après les conclusions/Q-R. Le rattrapage de volume se fait
    # désormais dans le calibrage budget texte, avec le plan verrouillé comme contexte.

    from services.content_generation_service import _current_compliance_review_signature
    compliance_signature = _current_compliance_review_signature()

    # 6. Conformité locale par segment, après adhérence au plan, calibrage
    # budget et micro-review éthique intégrés à la génération structurée.
    not_reviewed = count_segments_pending_review_for_folders(folder_ids, compliance_signature)
    if not_reviewed > 0:
        return "review"

    # 7. Document post-révision : ré-assemble le texte courant validé avant
    # d'autoriser slides et synthèse audio.
    if not j.get("auto_pilot_post_review_docs_done"):
        return "post_review_docs"

    # 8. Slides anchor-first : elles sont générées explicitement avant la fin
    # texte, pour ne plus rester cachées dans l'étape TTS synchronisée.
    slide_rows = list_completed_content_jobs_for_folders(folder_ids)
    missing_slide_decks = []
    for row in slide_rows:
        fid = int(row["folder_id"])
        deck_row = get_latest_script_slide_deck_row(
            folder_id=fid,
            content_job_id=int(row["content_job_id"]),
        )
        deck_has_slides = False
        if deck_row:
            try:
                deck_has_slides = len(_json.loads(deck_row.get("slides_json") or "[]")) > 0
            except Exception:
                deck_has_slides = False
        if not deck_has_slides:
            missing_slide_decks.append(fid)
    if missing_slide_decks:
        return "slides"

    # 10. Audio TTS optionnel. Par défaut l'auto-pilot s'arrête texte prêt :
    # les audios se génèrent ensuite à la demande, journée/semaine par journée/semaine.
    if not j.get("auto_pilot_generate_audio"):
        try:
            _finalize_text_ready_state(job_id)
        except Exception as e:
            logger.warning(f"⚠️ Finalisation texte job {job_id} ignorée : {e}")
        if j.get("status") != "text_ready":
            update_job(job_id, status="text_ready", error_message=None)
        return None

    # Si l'audio auto est explicitement demandé, on vérifie via dirty=0 sur tous les segments (pas le status qui
    # est positionné au début du loop audio, donc non fiable en cas de restart)
    dirty_count = count_dirty_completed_segments_for_folders(folder_ids)
    if dirty_count > 0 or j.get("status") not in ("audio_completed", "audio_launched"):
        return "audio"

    return None  # tout est fait



def _execute_ap_step(job_id: int, step: str, job: dict) -> None:
    """Exécute UNE étape de l'auto-pilot (synchrone — bloque le tick courant)."""
    import eventlet
    from database.db import get_db_connection
    from services.claude_code_mission_service import execute_mission_locally

    model = _normalize_pipeline_model_choice(job.get("auto_pilot_model"), default="pro")
    tts_mode = job.get("auto_pilot_tts_mode") or "gtts"
    use_cc = bool(job.get("auto_pilot_use_cc"))
    platform_id = job["platform_id"]

    api_model = _PIPELINE_MODEL_ALIASES[model]
    cc_model = "haiku" if model == "haiku" else "sonnet"

    def _assert_cc_result_complete(step_label: str, result: dict | None) -> None:
        if isinstance(result, dict) and int(result.get("chunks_failed") or 0) > 0:
            errors = result.get("errors") or []
            first_error = ""
            if errors:
                first = errors[0]
                first_error = f" — {first.get('chunk')}: {first.get('error')}"
            raise RuntimeError(
                f"{step_label} Claude Code incomplète : "
                f"{result.get('chunks_failed')} chunk(s) en erreur{first_error}"
            )

    def _wait_status(targets, max_wait=3600):
        elapsed = 0
        while elapsed < max_wait:
            j = get_job(job_id)
            if not j:
                raise RuntimeError("Job introuvable")
            s = j["status"]
            if s in targets:
                return j
            if s in ("error", "audio_error"):
                raise RuntimeError(f"Pipeline erreur : {j.get('error_message') or s}")
            eventlet.sleep(3)
            elapsed += 3
        raise TimeoutError(f"Timeout {max_wait}s en attendant {targets}")

    if step == "reac":
        from services.formation_health_service import compute_preflight
        preflight = compute_preflight(job_id, use_claude_code=use_cc, tts_mode=tts_mode)
        if not preflight["ok"]:
            raise RuntimeError(f"Pre-flight bloqué : {', '.join(preflight['blocking'])}")
        logger.info(f"🛂 Pre-flight OK job {job_id}")
        update_job(job_id, status="reac_fetching")

        def _log_reac_attempt(**payload):
            try:
                from services.formation_observability_service import log_pipeline_event
                status = payload.get("status") or "info"
                attempt = payload.get("attempt")
                total = payload.get("total")
                wait_seconds = payload.get("wait_seconds") or 0
                error = payload.get("error")
                message = f"REAC tentative {attempt}/{total} : {status}"
                if status == "retrying":
                    message += f" — nouvelle tentative dans {wait_seconds:.0f}s"
                log_pipeline_event(
                    job_id,
                    "reac_download_attempt",
                    step="reac",
                    status="error" if status == "failed" else status,
                    message=message,
                    model=api_model,
                    data={
                        "attempt": attempt,
                        "total": total,
                        "wait_seconds": wait_seconds,
                        "rncp_code": job["rncp_code"],
                    },
                    error=error,
                )
            except Exception:
                pass

        reac = download_reac_text_with_retry(
            job["rncp_code"],
            attempts=3,
            on_attempt=_log_reac_attempt,
        )
        rc_text, rome_text = None, None
        try:
            rc_text = download_rc_text(job["rncp_code"]) or None
        except Exception:
            pass
        try:
            rome_text = fetch_rome_data(job["rncp_code"]) or None
        except Exception:
            pass
        update_job(job_id, status="reac_ready", reac_text=reac,
                   rc_text=rc_text, rome_text=rome_text)
        logger.info(f"🤖 ✓ REAC téléchargé job {job_id}")

    elif step == "kb":
        if use_cc:
            execute_mission_locally(job_id, "kb", cc_model)
        else:
            launch_kb_building(job_id, model=api_model)
            _wait_status({"kb_ready"}, max_wait=3600)
        logger.info(f"🤖 ✓ KB construite job {job_id}")

    elif step == "global":
        if use_cc:
            execute_mission_locally(job_id, "global", cc_model)
        else:
            launch_global_program_generation(job_id, model=api_model)
            _wait_status({"global_ready"}, max_wait=600)
        update_job(job_id, global_program_validated=1, status="global_validated")
        logger.info(f"🤖 ✓ Programme global validé job {job_id}")

    elif step == "daily":
        if use_cc:
            execute_mission_locally(job_id, "daily", cc_model)
        else:
            run_daily_split(job_id, model=api_model)
        update_job(job_id, daily_programs_validated=1, status="daily_validated")
        logger.info(f"🤖 ✓ Programmes journée validés job {job_id}")

    elif step == "content":
        if use_cc:
            execute_mission_locally(job_id, "content", cc_model)
        else:
            # Mode API : génération synchrone par folder, avec concurrence bornée.
            # Pas de thread background durable : run_content_generation reste appelé
            # dans le greenlet auto-pilot — résiste aux restarts Azure car :
            #   - folder existant mais job running/idle → run reprend les segments manquants
            #   - folder existant + job completed → skip via done_set
            #   - aucun thread mort possible (pas de thread du tout)
            import json as _json
            import eventlet as _eventlet
            from services.content_generation_service import run_content_generation, get_job_from_db
            from repositories.pipeline_repository import reset_and_upsert_content_generation_job
            from services.formation_pipeline_service import (
                _format_day_program_text,
                expected_course_folder_name,
                get_expected_course_folders,
            )

            daily_programs = _json.loads(job.get("daily_programs") or "[]")
            update_job(
                job_id,
                status="tts_launched",
                auto_pilot_volume_done=0,
                auto_pilot_post_review_docs_done=0,
                error_message=None,
            )
            folder_state = get_expected_course_folders(
                job_id,
                create_missing=True,
                platform_id=platform_id,
            )
            folders_by_name = {
                f["expected_name"]: f
                for f in folder_state.get("folders", [])
            }
            if folder_state.get("duplicates"):
                logger.warning(
                    "PIPELINE_CONTENT_DUPLICATE_FOLDERS job=%s duplicates=%s",
                    job_id,
                    [
                        {
                            "folder_id": d["folder_id"],
                            "name": d["name"],
                            "duplicate_of": d.get("duplicate_of"),
                        }
                        for d in folder_state["duplicates"]
                    ],
                )

            day_tasks = []
            for idx, day_data in enumerate(daily_programs):
                day_data = _normalize_day_audio_slots(day_data)
                folder_name = expected_course_folder_name(day_data, idx + 1)
                folder_info = folders_by_name.get(folder_name)
                if not folder_info:
                    raise RuntimeError(f"Folder attendu introuvable : {folder_name}")
                folder_id = folder_info["folder_id"]
                day_num = day_data.get("day_number", idx + 1)

                # 2. Créer le content_generation_job s'il n'existe pas (sans thread)
                cg_job = get_job_from_db(folder_id)
                if cg_job and cg_job.get("status") == "completed":
                    logger.info(f"🤖   ⏭ Jour {day_num} déjà complété (folder {folder_id}), skip")
                    continue

                if not cg_job:
                    sub_parts = [sp["name"] for sp in day_data.get("sub_parts", [])]
                    module_contents = {}
                    for sp in day_data.get("sub_parts", []):
                        module_contents[sp["name"]] = _format_slot_generation_source(sp)
                    reset_and_upsert_content_generation_job(
                        folder_id=folder_id,
                        platform_id=platform_id,
                        program_text=_format_day_program_text(day_data, job["tp_name"]),
                        program_title=job["tp_name"],
                        sub_parts_json=_json.dumps(sub_parts, ensure_ascii=False),
                        from_scratch=True,
                        module_contents_json=_json.dumps(module_contents, ensure_ascii=False),
                    )

                day_tasks.append({"day_num": day_num, "folder_id": folder_id})

            day_workers = min(_formation_content_day_workers(), max(1, len(day_tasks) or 1))
            logger.info(
                "PIPELINE_CONTENT_DAY_PARALLEL_CONFIG job=%s workers=%s days=%s",
                job_id,
                day_workers,
                len(day_tasks),
            )

            def _run_content_day(task: dict) -> dict:
                day_num = task["day_num"]
                folder_id = task["folder_id"]
                try:
                    # run_content_generation lit les segments déjà complétés :
                    # idempotent, reprise naturelle si un folder avait déjà démarré.
                    logger.info(f"🤖   Génération Jour {day_num} (folder {folder_id})…")
                    run_content_generation(folder_id, model=api_model)
                    logger.info(f"🤖   ✓ Jour {day_num} terminé (folder {folder_id})")
                    return {"ok": True, "day_num": day_num, "folder_id": folder_id}
                except Exception as exc:
                    logger.exception(
                        "PIPELINE_CONTENT_DAY_FAILED job=%s day=%s folder=%s",
                        job_id,
                        day_num,
                        folder_id,
                    )
                    return {
                        "ok": False,
                        "day_num": day_num,
                        "folder_id": folder_id,
                        "error": str(exc)[:500],
                    }

            if day_workers <= 1:
                day_results = [_run_content_day(task) for task in day_tasks]
            else:
                pool = _eventlet.GreenPool(size=day_workers)
                greenlets = [pool.spawn(_run_content_day, task) for task in day_tasks]
                day_results = [greenlet.wait() for greenlet in greenlets]

            failed_days = [result for result in day_results if not result.get("ok")]
            if failed_days:
                raise RuntimeError(
                    "Génération contenu échouée sur journées : "
                    + ", ".join(
                        f"J{result['day_num']} folder={result['folder_id']} ({result.get('error')})"
                        for result in failed_days
                    )
                )

        logger.info(f"🤖 ✓ Contenu généré job {job_id}")

    elif step == "volume_safety":
        logger.warning(
            "🤖 Auto-pilot job %s : étape volume_safety ignorée, réparation append-only désactivée",
            job_id,
        )
        update_job(job_id, auto_pilot_volume_done=1, auto_pilot_post_review_docs_done=0)

    elif step == "humanization_review":
        logger.info(
            "🤖 Auto-pilot job %s : étape humanization_review ignorée, "
            "l'oralité est portée par le prompt initial",
            job_id,
        )
        update_job(job_id, auto_pilot_post_review_docs_done=0)

    elif step == "review":
        from services.formation_pipeline_service import get_expected_course_folders
        folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
        if use_cc:
            result = execute_mission_locally(job_id, "review", cc_model)
            _assert_cc_result_complete("Révision conformité", result)
        else:
            from services.content_generation_service import run_content_review
            failed = []
            reports_written = 0
            for fid in folder_ids:
                try:
                    result = run_content_review(fid, model=api_model)
                    _write_api_review_report(job_id, fid, result, api_model)
                    reports_written += 1
                    if result.get("segments_failed", 0) > 0:
                        error_samples = []
                        for detail in result.get("details") or []:
                            error = str((detail or {}).get("error") or "").strip()
                            if not error:
                                continue
                            error = " ".join(error.split())[:180]
                            if error not in error_samples:
                                error_samples.append(error)
                            if len(error_samples) >= 2:
                                break
                        suffix = f" : {' ; '.join(error_samples)}" if error_samples else ""
                        failed.append(f"{fid}({result['segments_failed']} segments échoués{suffix})")
                except Exception as e:
                    logger.warning(f"⚠️ Review folder {fid} : {e}")
                    failed.append(str(fid))
            if failed:
                raise RuntimeError(f"Review échouée sur folders : {', '.join(failed)}")
            logger.info(
                "🤖 ✓ Rapports conformité persistés job %s : %s/%s",
                job_id,
                reports_written,
                len(folder_ids),
            )
        update_job(job_id, auto_pilot_post_review_docs_done=0)
        logger.info(f"🤖 ✓ Conformité locale terminée job {job_id}")

    elif step == "post_review_docs":
        from services.content_generation_service import (
            _assemble_and_upload,
            _update_job_db,
            assert_course_day_word_budget,
        )
        from services.formation_pipeline_service import get_expected_course_folders
        from repositories.pipeline_repository import list_completed_content_jobs_for_folders
        folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
        if not folder_ids:
            raise RuntimeError("Aucun texte complété à assembler après révision")
        rows = [
            (int(row["folder_id"]), int(row["content_job_id"]))
            for row in list_completed_content_jobs_for_folders(folder_ids)
        ]
        if not rows:
            raise RuntimeError("Aucun texte complété à assembler après révision")

        for fid, cg_job_id in rows:
            budget_audit = assert_course_day_word_budget(
                fid,
                context="auto_pilot_post_review_docs",
            )
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "day_word_budget_verified",
                    step="word_budget_review",
                    status="completed",
                    folder_id=fid,
                    model=api_model,
                    message="Budget mots journée vérifié avant Word 2",
                    data={
                        "spoken_words": budget_audit.get("spoken_words"),
                        "raw_words": budget_audit.get("raw_words"),
                        "deficit": budget_audit.get("deficit"),
                        "overflow": budget_audit.get("overflow"),
                        "budget": {
                            "target_words": budget_audit.get("budget", {}).get("target_words"),
                            "min_words": budget_audit.get("budget", {}).get("min_words"),
                            "max_words": budget_audit.get("budget", {}).get("max_words"),
                            "words_per_minute": budget_audit.get("budget", {}).get("words_per_minute"),
                            "course_seconds": budget_audit.get("budget", {}).get("course_seconds"),
                            "speakable_seconds": budget_audit.get("budget", {}).get("speakable_seconds"),
                        },
                    },
                )
            except Exception:
                pass
            final_words, filename = _assemble_and_upload(fid, platform_id, cg_job_id)
            _update_job_db(cg_job_id, total_words=final_words)
            try:
                deleted_decks = _delete_slide_deck_for_resume(fid, cg_job_id)
            except Exception as e:
                deleted_decks = 0
                logger.warning(
                    "⚠️ Nettoyage deck slides post-Word2 impossible job=%s folder=%s: %s",
                    job_id,
                    fid,
                    str(e)[:300],
                )
            logger.info(
                f"🤖   ✓ Document post-révision folder {fid} : "
                f"{final_words} mots, {filename}, decks slides supprimés={deleted_decks}"
            )
        next_status = "tts_launched" if job.get("auto_pilot_generate_audio") else "text_ready"
        if next_status == "text_ready":
            _finalize_text_ready_state(job_id)
        update_job(job_id, status=next_status, auto_pilot_post_review_docs_done=1)
        logger.info(f"🤖 ✓ Documents post-révision générés job {job_id} (status={next_status})")

    elif step == "slides":
        from services.formation_pipeline_service import get_expected_course_folders
        from services.script_slide_generation_service import (
            generate_slides_from_script,
            get_latest_script_slide_deck,
        )
        from services.formation_observability_service import log_pipeline_event
        from repositories.pipeline_repository import list_completed_content_jobs_for_folders

        folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
        if not folder_ids:
            raise RuntimeError("Aucun cours_folder trouvé pour générer les slides")
        rows = [
            (int(row["folder_id"]), int(row["content_job_id"]))
            for row in list_completed_content_jobs_for_folders(folder_ids)
        ]
        if not rows:
            raise RuntimeError("Aucun texte complété disponible pour générer les slides")

        generated = 0
        skipped = 0
        failures = []
        slide_api_model = _resolve_pipeline_slide_model(api_model)
        slide_workers = min(_slides_folder_workers(), max(1, len(rows)))
        pending_rows = []
        for fid, cg_job_id in rows:
            existing = get_latest_script_slide_deck(fid, content_job_id=cg_job_id)
            if existing and (existing.get("slides") or []):
                skipped += 1
                continue
            pending_rows.append((fid, cg_job_id))

        logger.info(
            "🤖 Slides job %s : folders=%s pending=%s skipped=%s workers=%s model=%s",
            job_id,
            len(rows),
            len(pending_rows),
            skipped,
            min(slide_workers, max(1, len(pending_rows))) if pending_rows else 0,
            slide_api_model,
        )

        def _generate_slide_folder(fid: int, cg_job_id: int) -> dict:
            started = time.time()
            try:
                log_pipeline_event(
                    job_id,
                    "slides_folder_started",
                    step="slides",
                    status="running",
                    folder_id=fid,
                    model=slide_api_model,
                    message="Génération slides anchor-first démarrée",
                    data={
                        "content_job_id": cg_job_id,
                        "max_slides": 60,
                        "pace": "normal",
                        "parallel_folders": min(slide_workers, max(1, len(pending_rows))),
                    },
                )
                result = generate_slides_from_script(
                    folder_id=fid,
                    job_id=job_id,
                    platform_id=platform_id,
                    max_slides=60,
                    pace="normal",
                    model=slide_api_model,
                )
                log_pipeline_event(
                    job_id,
                    "slides_folder_completed",
                    step="slides",
                    status="completed",
                    folder_id=fid,
                    model=slide_api_model,
                    duration_ms=int((time.time() - started) * 1000),
                    message="Génération slides anchor-first terminée",
                    data={
                        "content_job_id": cg_job_id,
                        "deck_id": (result.get("stats") or {}).get("deck_id"),
                        "slides_generated": (result.get("stats") or {}).get("slides_generated"),
                        "slide_anchors_found": (result.get("stats") or {}).get("slide_anchors_found"),
                        "generation_mode": (result.get("stats") or {}).get("generation_mode"),
                        "slide_batch_workers": (result.get("stats") or {}).get("slide_batch_workers"),
                    },
                )
                return {"status": "generated", "folder_id": fid}
            except Exception as e:
                err = str(e)[:500]
                try:
                    log_pipeline_event(
                        job_id,
                        "slides_folder_failed",
                        step="slides",
                        status="error",
                        folder_id=fid,
                        model=slide_api_model,
                        duration_ms=int((time.time() - started) * 1000),
                        message="Génération slides anchor-first échouée",
                        error=err,
                    )
                except Exception:
                    pass
                return {"status": "failed", "folder_id": fid, "error": err}

        if pending_rows:
            active_workers = min(slide_workers, len(pending_rows))
            if active_workers > 1:
                with ThreadPoolExecutor(max_workers=active_workers) as pool:
                    future_map = {
                        pool.submit(_generate_slide_folder, fid, cg_job_id): (fid, cg_job_id)
                        for fid, cg_job_id in pending_rows
                    }
                    for future in as_completed(future_map):
                        result = future.result()
                        if result.get("status") == "generated":
                            generated += 1
                        else:
                            failures.append(f"{result.get('folder_id')}({result.get('error')})")
            else:
                for fid, cg_job_id in pending_rows:
                    result = _generate_slide_folder(fid, cg_job_id)
                    if result.get("status") == "generated":
                        generated += 1
                    else:
                        failures.append(f"{result.get('folder_id')}({result.get('error')})")
        if failures:
            raise RuntimeError("Slides échouées sur folders : " + ", ".join(failures))
        logger.info(
            "🤖 ✓ Slides générées job %s : generated=%s skipped=%s workers=%s model=%s",
            job_id,
            generated,
            skipped,
            min(slide_workers, max(1, len(pending_rows))) if pending_rows else 0,
            slide_api_model,
        )

    elif step == "audio":
        from services.content_generation_service import generate_audio_from_script
        from services.formation_pipeline_service import get_expected_course_folders
        folder_ids = get_expected_course_folders(job_id).get("folder_ids") or []
        if not folder_ids:
            raise RuntimeError("Aucun cours_folder trouvé pour la plateforme")

        mock = (tts_mode == "mock")
        basic_tts = (tts_mode == "gtts")
        voice_type = "mock" if mock else ("gtts" if basic_tts else "fish_audio")
        update_job(job_id, status="audio_running", error_message=None)
        for idx, fid in enumerate(folder_ids):
            next_fid = folder_ids[idx + 1] if idx + 1 < len(folder_ids) else None
            folder_started_at = time.time()
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "audio_folder_started",
                    step="audio",
                    status="running",
                    folder_id=fid,
                    message="Synthèse audio journée démarrée",
                    model=_resolve_pipeline_api_model(job),
                    data={"voice_type": voice_type, "auto_pilot": True},
                )
            except Exception:
                pass
            try:
                generate_audio_from_script(
                    fid,
                    on_progress=_make_audio_progress_logger(job_id, fid, voice_type),
                    force_all=False,
                    mock=mock,
                    basic_tts=basic_tts,
                    next_folder_id=next_fid,
                    is_last_folder=next_fid is None,
                    sync_slides=True,
                    auto_generate_slides=True,
                    slide_max_slides=60,
                    slide_pace="normal",
                    slide_model=_resolve_pipeline_api_model(job),
                    llm_model=_resolve_pipeline_api_model(job),
                )
            except Exception as e:
                try:
                    from services.formation_observability_service import log_pipeline_event
                    log_pipeline_event(
                        job_id,
                        "audio_folder_failed",
                        step="audio",
                        status="error",
                        folder_id=fid,
                        duration_ms=int((time.time() - folder_started_at) * 1000),
                        message="Synthèse audio journée échouée",
                        model=_resolve_pipeline_api_model(job),
                        data={"voice_type": voice_type, "auto_pilot": True},
                        error=str(e)[:500],
                    )
                except Exception:
                    pass
                raise
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "audio_folder_completed",
                    step="audio",
                    status="completed",
                    folder_id=fid,
                    duration_ms=int((time.time() - folder_started_at) * 1000),
                    message="Synthèse audio journée terminée",
                    model=_resolve_pipeline_api_model(job),
                    data={"voice_type": voice_type, "auto_pilot": True},
                )
            except Exception:
                pass
            eventlet.sleep(5)
        # Status posé APRÈS le loop — si Azure redémarre en cours de route,
        # dirty=1 sur les folders non traités permettra à _determine_next_ap_step
        # de détecter que l'audio est incomplet et de relancer l'étape.
        update_job(job_id, status="audio_completed", error_message=None)

        # Module persistant
        try:
            from datetime import datetime as _dt
            from config import FRANCE_TZ as _tz
            j2 = get_job(job_id)
            conn = get_db_connection()
            cur = conn.cursor()
            rncp = j2.get("rncp_code") or ""
            cur.execute("SELECT center_account_id FROM platform_config WHERE id = ?", (platform_id,))
            center_row = cur.fetchone()
            center_account_id = center_row[0] if center_row else None
            if center_account_id is None:
                cur.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id IS NULL",
                    (rncp,),
                )
            else:
                cur.execute(
                    "SELECT COUNT(*) FROM formation_modules WHERE rncp_code = ? AND center_account_id = ?",
                    (rncp, center_account_id),
                )
            n = cur.fetchone()[0] + 1
            version = f"{_dt.now(_tz).year}-v{n}"
            cur.execute("""
                INSERT OR IGNORE INTO formation_modules
                (rncp_code, tp_name, version, status, source_pipeline_job_id,
                 source_platform_id, center_account_id, voice_type, voice_updated_at, validated_at)
                VALUES (?, ?, ?, 'validated', ?, ?, ?, ?, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
            """, (rncp, j2["tp_name"], version, job_id, platform_id, center_account_id, voice_type))
            if cur.rowcount == 0:
                cur.execute(
                    "UPDATE formation_modules SET voice_type=?, voice_updated_at=CURRENT_TIMESTAMP, "
                    "center_account_id=COALESCE(center_account_id, ?) "
                    "WHERE source_pipeline_job_id=?",
                    (voice_type, center_account_id, job_id),
                )
            conn.commit()
            conn.close()
        except Exception as e:
            logger.warning(f"⚠️ Module persistant : {e}")

        # Health-check final — bloque sur les incohérences bloquantes.
        # Un warning silencieux permettrait à l'auto-pilot de finir avec une
        # formation incomplète (segments manquants, audios dirty, etc.).
        from services.formation_health_service import compute_health
        health = compute_health(job_id)
        if health["ok"]:
            logger.info(f"💚 Health-check OK job {job_id}")
        else:
            blocking = health.get("blocking", [])
            raise RuntimeError(
                f"Health-check : {len(blocking)} incohérence(s) bloquante(s) : {blocking}"
            )
        logger.info(f"🤖 ✓ Audio TTS terminé job {job_id}")


def _tick_auto_pilot(job_id: int) -> None:
    """Runner auto-pilot : 1 tick = 1 étape. Se respawn pour la suivante.

    Persisté en DB — résiste aux restarts Azure App Service.
    Lock optimiste + heartbeat toutes les 60 s pour les étapes longues (content/audio).
    """
    import eventlet

    if not _acquire_ap_lock(job_id):
        logger.info(f"🤖 Auto-pilot job {job_id} : lock tenu par un autre worker, skip")
        return

    # Heartbeat : rafraîchit le lock toutes les 60 s pour les étapes qui durent
    # plus que le TTL (content = 2-4h, audio = 30-60 min).
    _hb_stop = [False]
    def _heartbeat():
        while not _hb_stop[0]:
            eventlet.sleep(60)
            if not _hb_stop[0]:
                _refresh_ap_lock(job_id)
                logger.info(
                    "PIPELINE_AUTOPILOT_HEARTBEAT job=%s step=%s lock_refreshed=1",
                    job_id,
                    current_step,
                )
    hb = eventlet.spawn(_heartbeat)

    should_respawn = False
    current_step = "?"
    started_at = None
    try:
        j = get_job(job_id)
        if not j or not j.get("auto_pilot_enabled"):
            return

        step = _determine_next_ap_step(job_id)
        if step is None:
            update_job(job_id, auto_pilot_step="done", auto_pilot_error=None)
            try:
                from services.formation_observability_service import log_pipeline_event
                log_pipeline_event(
                    job_id,
                    "pipeline_completed",
                    step="done",
                    status="completed",
                    model=j.get("auto_pilot_model"),
	                    message=(
	                        "Auto-pilot texte et slides terminé"
	                        if not j.get("auto_pilot_generate_audio")
	                        else "Auto-pilot terminé"
	                    ),
                    data={"generate_audio": bool(j.get("auto_pilot_generate_audio"))},
                )
            except Exception:
                pass
            logger.info(f"🤖 ✅ Auto-pilot TERMINÉ job {job_id}")
            return

        current_step = step
        update_job(job_id, auto_pilot_step=step, auto_pilot_error=None)
        logger.info(f"🤖 Auto-pilot job {job_id} → step={step}")
        started_at = time.time()
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "step_started",
                step=step,
                status="running",
                model=j.get("auto_pilot_model"),
                message=f"Étape auto-pilot démarrée : {step}",
                data={
                    "tts_mode": j.get("auto_pilot_tts_mode"),
                    "use_claude_code": bool(j.get("auto_pilot_use_cc")),
                    "generate_audio": bool(j.get("auto_pilot_generate_audio")),
                },
            )
        except Exception:
            pass

        _execute_ap_step(job_id, step, j)
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "step_completed",
                step=step,
                status="completed",
                model=j.get("auto_pilot_model"),
                duration_ms=int((time.time() - started_at) * 1000) if started_at else None,
                message=f"Étape auto-pilot terminée : {step}",
            )
        except Exception:
            pass
        should_respawn = True

    except Exception as e:
        err = str(e)[:500]
        logger.error(f"❌ Auto-pilot job {job_id} step={current_step} : {err}")
        update_job(job_id, auto_pilot_error=err)
        try:
            from services.formation_observability_service import log_pipeline_event
            log_pipeline_event(
                job_id,
                "step_failed",
                step=current_step,
                status="error",
                duration_ms=int((time.time() - started_at) * 1000) if started_at else None,
                message=f"Étape auto-pilot échouée : {current_step}",
                error=err,
                model=_resolve_pipeline_api_model(job),
            )
        except Exception:
            pass

    finally:
        _hb_stop[0] = True
        hb.kill()
        _release_ap_lock(job_id)

    if should_respawn:
        eventlet.spawn(_tick_auto_pilot, job_id)


def resume_interrupted_auto_pilots() -> None:
    """Appelé au boot : reprend les auto-pilots interrompus par un restart Azure."""
    import eventlet
    from services.formation_pipeline_service import get_auto_pilot_jobs_to_resume
    eventlet.sleep(5)  # laisse le temps à l'app de démarrer
    try:
        job_ids = get_auto_pilot_jobs_to_resume()
        if job_ids:
            logger.info(f"🤖 Boot recovery : {len(job_ids)} auto-pilot(s) à reprendre : {job_ids}")
            for jid in job_ids:
                eventlet.spawn(_tick_auto_pilot, jid)
        else:
            logger.info("🤖 Boot recovery : aucun auto-pilot interrompu")
    except Exception as e:
        logger.warning(f"⚠️ Boot recovery auto-pilot : {e}")


def _auto_pilot_watchdog_loop(interval_sec: int) -> None:
    """Surveille périodiquement les locks auto-pilot absents ou périmés."""
    import eventlet
    from services.formation_pipeline_service import get_auto_pilot_jobs_to_resume

    eventlet.sleep(5)  # laisse le temps à l'app de démarrer
    while True:
        try:
            job_ids = get_auto_pilot_jobs_to_resume()
            if job_ids:
                logger.warning(
                    f"🤖 Watchdog auto-pilot : reprise de {len(job_ids)} job(s) : {job_ids}"
                )
                for jid in job_ids:
                    eventlet.spawn(_tick_auto_pilot, jid)
            else:
                logger.debug("🤖 Watchdog auto-pilot : aucun job à reprendre")
        except Exception as e:
            logger.warning(f"⚠️ Watchdog auto-pilot : {e}")
        eventlet.sleep(interval_sec)


def start_auto_pilot_watchdog() -> None:
    """Démarre un watchdog par process pour reprendre les locks zombies."""
    global _AP_WATCHDOG_STARTED
    if _AP_WATCHDOG_STARTED:
        return
    _AP_WATCHDOG_STARTED = True

    import eventlet
    logger.info(
        f"🤖 Watchdog auto-pilot démarré (intervalle {_AP_WATCHDOG_INTERVAL}s)"
    )
    eventlet.spawn(_auto_pilot_watchdog_loop, _AP_WATCHDOG_INTERVAL)


@formation_bp.route("/api/formation/<int:job_id>/run-auto", methods=["POST"])
def run_auto_pilot(job_id):
    """Lance l'auto-pilot : REAC → KB → global → daily → content
    (génération sections → adhérence plan → budget → micro-éthique)
    → conformité locale par morceau → Word 2 → slides → audio optionnel.

    Body (optionnel) :
      - tts_mode : 'fish_audio' | 'gtts' | 'mock' (défaut 'gtts')
      - model : 'sonnet' | 'haiku' | 'flash' | 'pro' (défaut 'pro')
      - use_claude_code : bool (défaut false)
      - generate_audio : bool (défaut false) — legacy, enchaîne aussi l'audio

    Retourne 202 immédiatement. Suivre via GET /run-auto/status.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    payload = request.get_json(silent=True) or {}
    tts_mode = (payload.get("tts_mode") or "gtts").lower()
    if tts_mode not in ("fish_audio", "gtts", "mock"):
        return jsonify({"error": "tts_mode invalide (fish_audio | gtts | mock)"}), 400
    model = _normalize_pipeline_model_choice(payload.get("model"), default=job.get("auto_pilot_model") or "pro")
    if model not in _PIPELINE_MODEL_CHOICES:
        return jsonify({"error": "model invalide (sonnet | haiku | flash | pro)"}), 400
    use_cc = bool(payload.get("use_claude_code", False))
    if use_cc and model in ("flash", "pro"):
        return jsonify({
            "error": "use_claude_code=true est incompatible avec DeepSeek. Lance l'auto-pilot en mode API."
        }), 400
    generate_audio = bool(payload.get("generate_audio") or payload.get("include_audio"))

    # Vérifie qu'un tick n'est pas déjà en cours (lock actif non-périmé)
    if job.get("auto_pilot_step") and job.get("auto_pilot_step") != "done":
        locked_at = job.get("auto_pilot_locked_at")
        if locked_at:
            try:
                import datetime, pytz
                from config import FRANCE_TZ
                if isinstance(locked_at, str):
                    dt = datetime.datetime.fromisoformat(locked_at)
                    if dt.tzinfo is None:
                        dt = dt.replace(tzinfo=pytz.utc)
                    age = (datetime.datetime.now(pytz.utc) - dt).total_seconds()
                    if age < _AP_LOCK_TTL:
                        return jsonify({"error": "Auto-pilot déjà en cours pour ce job"}), 409
            except Exception:
                pass

    update_job(job_id,
               auto_pilot_enabled=1,
               auto_pilot_model=model,
               auto_pilot_tts_mode=tts_mode,
               auto_pilot_use_cc=int(use_cc),
               auto_pilot_skip_vs=0,
               auto_pilot_generate_audio=int(generate_audio),
               auto_pilot_volume_done=0,
               auto_pilot_post_review_docs_done=0,
               auto_pilot_error=None)
    try:
        from services.formation_observability_service import log_pipeline_event
        log_pipeline_event(
            job_id,
            "pipeline_started",
            step="start",
            status="running",
            model=model,
            message="Auto-pilot lancé",
            data={"tts_mode": tts_mode, "use_claude_code": use_cc, "generate_audio": generate_audio},
        )
    except Exception:
        pass

    import eventlet
    eventlet.spawn(_tick_auto_pilot, job_id)
    return jsonify({
        "ok": True, "status": "auto_pilot_started",
        "tts_mode": tts_mode, "model": model,
        "use_claude_code": use_cc,
        "generate_audio": generate_audio,
    }), 202


@formation_bp.route("/api/formation/<int:job_id>/run-auto/resume", methods=["POST"])
def resume_auto_pilot(job_id):
    """Reprend l'auto-pilot sans réinitialiser les flags déjà validés."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    lock_age = _ap_lock_age_seconds(job)
    lock_active = lock_age is not None and lock_age < _AP_LOCK_TTL
    force = bool((request.get_json(silent=True) or {}).get("force"))
    if lock_active and not force:
        return jsonify({
            "error": "Auto-pilot déjà en cours pour ce job",
            "lock_age_seconds": lock_age,
        }), 409

    try:
        next_step = _determine_next_ap_step(job_id)
    except Exception as e:
        return jsonify({"error": f"Impossible de calculer la prochaine étape : {str(e)[:300]}"}), 500

    update_job(
        job_id,
        auto_pilot_enabled=1,
        auto_pilot_error=None,
        auto_pilot_locked_at=None,
        auto_pilot_lock_owner=None,
        auto_pilot_step=next_step or "done",
    )

    if next_step is None:
        return jsonify({
            "ok": True,
            "status": "done",
            "step": "done",
            "next_step": None,
        }), 200

    try:
        from services.formation_observability_service import log_pipeline_event
        log_pipeline_event(
            job_id,
            "pipeline_resume_requested",
            step=next_step,
            status="running",
            model=job.get("auto_pilot_model"),
            message=f"Reprise auto-pilot demandée : {next_step}",
            data={"previous_step": job.get("auto_pilot_step"), "lock_age_seconds": lock_age},
        )
    except Exception:
        pass

    import eventlet
    eventlet.spawn(_tick_auto_pilot, job_id)
    return jsonify({
        "ok": True,
        "status": "auto_pilot_resumed",
        "step": next_step,
        "next_step": next_step,
        "model": job.get("auto_pilot_model"),
        "tts_mode": job.get("auto_pilot_tts_mode"),
        "generate_audio": bool(job.get("auto_pilot_generate_audio")),
    }), 202


@formation_bp.route("/api/formation/<int:job_id>/run-auto/stop", methods=["POST"])
def stop_auto_pilot(job_id):
    """Stoppe l'auto-pilot persistant pour un job.

    Arrêt coopératif : désactive la reprise/watchdog et libère le lock. Une
    étape déjà en cours peut finir son appel courant, mais ne respawnera pas.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403

    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    update_job(
        job_id,
        auto_pilot_enabled=0,
        auto_pilot_step="stopped",
        auto_pilot_error=None,
        auto_pilot_locked_at=None,
        auto_pilot_lock_owner=None,
    )
    try:
        from services.formation_observability_service import log_pipeline_event
        log_pipeline_event(
            job_id,
            "pipeline_stopped",
            step=job.get("auto_pilot_step") or "unknown",
            status="stopped",
            model=job.get("auto_pilot_model"),
            message="Auto-pilot stoppé manuellement",
        )
    except Exception:
        pass
    logger.warning("🤖 Auto-pilot job %s stoppé manuellement", job_id)
    return jsonify({"ok": True, "status": "stopped"}), 200


@formation_bp.route("/api/formation/<int:job_id>/run-auto/status", methods=["GET"])
def auto_pilot_status(job_id):
    """État de l'auto-pilot lu depuis la DB (résiste aux restarts)."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404
    if not job.get("auto_pilot_enabled"):
        if job.get("auto_pilot_step") == "stopped":
            try:
                next_step = _determine_next_ap_step(job_id)
            except Exception:
                next_step = None
            return jsonify({
                "status": "stopped",
                "step": "stopped",
                "next_step": next_step,
                "model": job.get("auto_pilot_model"),
                "tts_mode": job.get("auto_pilot_tts_mode"),
                "generate_audio": bool(job.get("auto_pilot_generate_audio")),
            }), 200
        return jsonify({"status": "idle"}), 200
    step = job.get("auto_pilot_step")
    error = job.get("auto_pilot_error")
    model = job.get("auto_pilot_model")
    tts_mode = job.get("auto_pilot_tts_mode")
    generate_audio = bool(job.get("auto_pilot_generate_audio"))
    lock_age = _ap_lock_age_seconds(job)
    lock_stale = lock_age is not None and lock_age >= _AP_LOCK_TTL
    try:
        next_step = _determine_next_ap_step(job_id)
    except Exception:
        next_step = None
    if step == "done":
        return jsonify({"status": "done", "step": "done", "next_step": next_step, "model": model, "tts_mode": tts_mode, "generate_audio": generate_audio}), 200
    if error:
        return jsonify({"status": "error", "step": step, "next_step": next_step, "error": error, "model": model, "tts_mode": tts_mode, "generate_audio": generate_audio, "lock_stale": lock_stale, "lock_age_seconds": lock_age}), 200
    if step:
        return jsonify({"status": "running", "step": step, "next_step": next_step, "model": model, "tts_mode": tts_mode, "generate_audio": generate_audio, "lock_stale": lock_stale, "lock_age_seconds": lock_age}), 200
    return jsonify({"status": "starting", "next_step": next_step, "model": model, "tts_mode": tts_mode, "generate_audio": generate_audio}), 200


@formation_bp.route("/api/formation/<int:job_id>/events", methods=["GET"])
def formation_pipeline_events(job_id):
    """Retourne les événements structurés de pipeline pour diagnostic/dashboard."""
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404
    try:
        limit = int(request.args.get("limit", 200))
    except (TypeError, ValueError):
        limit = 200
    from services.formation_observability_service import list_pipeline_events
    return jsonify({"events": list_pipeline_events(job_id, limit=limit)}), 200


@formation_bp.route("/api/formation/<int:job_id>/diagnostic", methods=["GET"])
def formation_pipeline_diagnostic(job_id):
    """Snapshot exploitable par l'UI : état job + health + événements récents.

    Objectif : ne plus diagnostiquer une pipeline depuis un simple status global.
    Cet endpoint agrège les signaux de contrôle sans relancer d'étape coûteuse.
    """
    if not _require_admin():
        return jsonify({"error": "Non autorisé"}), 403
    job = get_job(job_id)
    if not job:
        return jsonify({"error": "Job introuvable"}), 404

    try:
        from services.formation_pipeline_service import repair_orphan_content_folders
        repair_orphan_content_folders(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic repair folders job {job_id} : {e}")

    try:
        events_limit = int(request.args.get("events_limit", 80))
    except (TypeError, ValueError):
        events_limit = 80

    try:
        from services.formation_health_service import compute_health
        health = compute_health(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic health job {job_id} : {e}")
        health = {
            "ok": False,
            "blocking": ["health_error"],
            "warnings": [],
            "checks": {"health_error": {"ok": False, "detail": str(e)[:500]}},
        }

    try:
        from services.claude_code_mission_service import compute_volume_audit
        volume_audit = compute_volume_audit(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic volume job {job_id} : {e}")
        volume_audit = None

    from services.formation_observability_service import list_pipeline_events
    events = list_pipeline_events(job_id, limit=events_limit)

    try:
        from services.formation_pipeline_service import get_expected_course_folders
        folder_state = get_expected_course_folders(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic résolution folders job {job_id} : {e}")
        folder_state = {"expected_count": 0, "folder_ids": [], "duplicates": [], "missing": []}

    try:
        next_auto_step = _determine_next_ap_step(job_id)
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic next auto step job {job_id} : {e}")
        next_auto_step = None

    folders = []
    try:
        from database.db import get_db_connection
        conn = get_db_connection()
        cursor = conn.cursor()
        folder_ids = folder_state.get("folder_ids") or []
        if folder_ids:
            placeholders = ",".join("?" * len(folder_ids))
            cursor.execute(
                f"""
            SELECT
                cf.id,
                cf.name,
                cf.position,
                cf.platform_id,
                cf.formation_job_id,
                cgj.id,
                cgj.status,
                COALESCE(cgj.total_words, 0),
                COUNT(cgs.id),
                SUM(CASE WHEN cgs.status = 'completed' THEN 1 ELSE 0 END),
                SUM(CASE WHEN COALESCE(cgs.reviewed, 0) = 1 THEN 1 ELSE 0 END),
                SUM(CASE WHEN cgs.review_error IS NOT NULL THEN 1 ELSE 0 END),
                SUM(CASE WHEN COALESCE(cgs.dirty, 0) = 1 THEN 1 ELSE 0 END)
            FROM cours_folders cf
            LEFT JOIN content_generation_jobs cgj ON cgj.folder_id = cf.id
            LEFT JOIN content_generation_segments cgs ON cgs.job_id = cgj.id
            WHERE cf.id IN ({placeholders})
            GROUP BY cf.id, cf.name, cf.position, cf.platform_id, cf.formation_job_id,
                     cgj.id, cgj.status, cgj.total_words
            ORDER BY cf.position ASC, cf.id ASC
            """,
                tuple(folder_ids),
            )
            for row in cursor.fetchall():
                (
                    folder_id,
                    name,
                    position,
                    platform_id,
                    formation_job_id,
                    content_job_id,
                    content_status,
                    total_words,
                    segments_total,
                    segments_completed,
                    reviewed_segments,
                    review_errors,
                    dirty_segments,
                ) = row
                folders.append({
                    "folder_id": folder_id,
                    "folder_label": f"F{folder_id}",
                    "name": name,
                    "position": position,
                    "platform_id": platform_id,
                    "formation_job_id": formation_job_id,
                    "content_job_id": content_job_id,
                    "content_status": content_status,
                    "total_words": total_words or 0,
                    "segments_total": segments_total or 0,
                    "segments_completed": segments_completed or 0,
                    "reviewed_segments": reviewed_segments or 0,
                    "review_errors": review_errors or 0,
                    "dirty_segments": dirty_segments or 0,
                })
        conn.close()
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic folders job {job_id} : {e}")

    finalize_result = None
    try:
        audio_done = job.get("status") in ("audio_completed", "audio_launched")
        audio_clean = bool(folders) and sum(int(f.get("dirty_segments") or 0) for f in folders) == 0
        if audio_done and audio_clean:
            voice_type = job.get("auto_pilot_tts_mode") or "gtts"
            finalize_result = _finalize_audio_ready_state(job_id, voice_type)
            job = get_job(job_id) or job
            try:
                from services.formation_health_service import compute_health
                health = compute_health(job_id)
            except Exception as health_e:
                logger.warning(f"⚠️ Diagnostic health recompute job {job_id} : {health_e}")
    except Exception as e:
        logger.warning(f"⚠️ Diagnostic finalisation audio job {job_id} : {e}")

    public_job = {
        key: job.get(key)
        for key in (
            "id",
            "job_label",
            "status",
            "platform_id",
            "platform_label",
            "platform_name",
            "tp_name",
            "rncp_code",
            "nb_days",
            "auto_pilot_enabled",
            "auto_pilot_step",
            "auto_pilot_error",
            "auto_pilot_model",
            "auto_pilot_tts_mode",
            "error_message",
        )
        if key in job
    }

    return jsonify({
        "job": public_job,
        "health": health,
        "volume_audit": volume_audit,
        "folders": folders,
        "folder_resolution": {
            "expected_count": folder_state.get("expected_count", 0),
            "duplicates": folder_state.get("duplicates", []),
            "missing": folder_state.get("missing", []),
        },
        "next_auto_step": next_auto_step,
        "events": events,
        "finalize": finalize_result,
    }), 200
